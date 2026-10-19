"""
tests/test_concrete_workflow.py
===============================
Checks for design/concrete/workflow.py (the store between the beam and
column steps) and the column results file. No ETABS is needed.
"""

import os
import sys
import time
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from design.concrete.columns import report as col_report  # noqa: E402
from design.concrete.workflow import (  # noqa: E402
    DesignStore,
    _gravity_options,
    deflection_table,
    numeric_like,
    write_deflection_xlsx,
)


def test_number_like_text_becomes_numbers_but_names_stay_text():
    table = pd.DataFrame({"UniqueName": ["17", "2GX-1"], "Station": ["0", "500.5"],
                          "Cover": ["None", "40"], "Note": ["a", "1"]})
    out = numeric_like(table)
    assert out["UniqueName"].tolist() == ["17", "2GX-1"]
    assert out["Station"].tolist() == [0.0, 500.5]
    assert pd.isna(out["Cover"][0]) and out["Cover"][1] == 40.0
    assert out["Note"].tolist() == ["a", "1"]


def test_the_store_is_saved_next_to_the_model_and_found_again(tmp_path):
    model = tmp_path / "TOWER.EDB"
    model.write_bytes(b"model")
    store = DesignStore(str(model), os.path.getmtime(model),
                        {"FRAME DATA": pd.DataFrame({"A": [1]})}, {"smrf": True})
    store.beam_results = pd.DataFrame({"UniqueName": ["B1"]})
    path = store.save()
    assert path == str(tmp_path / "TOWER - design data.pkl")
    again = DesignStore.load(str(model))
    assert again.inputs == {"smrf": True}
    assert again.beam_results["UniqueName"].tolist() == ["B1"]
    assert not again.is_stale()
    later = time.time() + 10
    os.utime(model, (later, later))
    assert again.is_stale()
    assert DesignStore.load(str(tmp_path / "OTHER.EDB")) is None


def test_gravity_combinations_leave_out_seismic_and_wind():
    combos = ["ULS 100 1.4 DL", "ULS 105 1.2 DL + f LL + 0.5 Lr + 1.0 WX",
              "ULS 107 (1.2 + Ev) DL + f LL + 1.0 RSA1", "ULS 101 1.2 DL + 1.6 LL + 0.5 Lr"]
    assert _gravity_options(combos) == ["ULS 100 1.4 DL", "ULS 101 1.2 DL + 1.6 LL + 0.5 Lr"]


def test_column_results_are_saved_with_group_headers(tmp_path):
    from openpyxl import load_workbook

    report = pd.DataFrame({
        "Column_Label": ["C1", "C1"], "UniqueName": ["2-C1", "2-C1"], "Story": ["2F", "2F"],
        "End": ["I", "J"], "Combo": ["ULS 100", "ULS 100"],
        "Flexure_Check": ["PASS", "FAIL"], "Column_Design_Status": ["PASS", "FAIL"],
    })
    groups = [("COLUMN LABEL / LEVEL", ["Column_Label", "UniqueName", "Story", "End", "Combo"]),
              ("FLEXURE / AXIAL", ["Flexure_Check"]),
              ("DESIGN STATUS", ["Column_Design_Status"])]
    path = col_report.write_column_results_xlsx(report, groups, str(tmp_path / "columns.xlsx"))
    sheet = load_workbook(path).active
    assert sheet["A2"].value == "COLUMN LABEL / LEVEL"
    assert sheet["B3"].value == col_report.COLUMN_REPORT_LABELS.get("UniqueName", "UniqueName")
    assert sheet["D4"].value == "Bottom (I)" and sheet["D5"].value == "Top (J)"
    assert sheet["F5"].font.bold  # FAIL in bold red


def test_the_deflection_table_has_one_row_per_beam(tmp_path):
    results = pd.DataFrame({
        "Story": ["2F", "2F"], "UniqueName": ["2GX-1", "2GX-1"], "Face": ["TOP", "BOTTOM"],
        "Width": [300, 300], "Depth": [500, 500], "Defl_live_mm": [12.0, 12.0],
        "Defl_ratio": [1.2, 1.2], "Deflection_Check": ["FAIL", "FAIL"],
    })
    table = deflection_table(results)
    assert len(table) == 1
    assert {"Δ live (mm)", "Deflection check"} <= set(table.columns)
    from openpyxl import load_workbook

    path = write_deflection_xlsx(table, str(tmp_path / "d.xlsx"))
    sheet = load_workbook(path).active
    assert sheet.cell(row=2, column=table.columns.get_loc("Deflection check") + 1).font.bold


def test_a_design_stored_before_the_folders_is_still_read(tmp_path):
    """BEHAVIOUR: design data saved by v0.3.8 or earlier names its classes by the
    modules of that time (design.concrete_workflow, design.column_slenderness: a
    column report keeps the slenderness geometry of its columns). It must open."""
    import pickle

    from design.concrete import workflow
    from design.concrete.columns.slenderness import AxisGeometry

    model = str(tmp_path / "model.EDB")
    store = workflow.DesignStore(model, 1.0, {}, {"smrf": True})
    report = pd.DataFrame({"UniqueName": ["GF-C1"]})
    report.attrs["slenderness_geometry"] = {"GF-C1": {"2": AxisGeometry(3000.0)}}
    store.column_report = report
    # protocol 2 writes a class by its module and name as text, each ending a line:
    # the module names of that time can be put back
    old = pickle.dumps(store, protocol=2)
    for now, then in (("design.concrete.workflow", "design.concrete_workflow"),
                      ("design.concrete.columns.slenderness", "design.column_slenderness")):
        now, then = (now + chr(10)).encode(), (then + chr(10)).encode()
        assert now in old
        old = old.replace(now, then)
    assert b"design.concrete_workflow" in old and b"design.column_slenderness" in old
    with open(workflow.DesignStore.path_for(model), "wb") as handle:
        handle.write(old)

    loaded = workflow.DesignStore.load(model)

    assert loaded is not None and loaded.inputs == {"smrf": True}
    geometry = loaded.column_report.attrs["slenderness_geometry"]["GF-C1"]["2"]
    assert isinstance(geometry, AxisGeometry) and geometry.lu == 3000.0


# --------------------------------------------------------------------------- #
# The beam bar spacing question (v0.7.0: the code spacing is the default)
# --------------------------------------------------------------------------- #
def _default_of(monkeypatch, last):
    from design.concrete import workflow
    from utilities import _gui_helpers

    seen = {}

    def fake(title, message, options, default_index=0, **_):
        seen["default"] = options[default_index]
        return options[default_index]

    monkeypatch.setattr(_gui_helpers, "select_option", fake)
    return workflow.ask_bar_spacing("Beam Design", last), seen["default"]


def test_bar_spacing_defaults_to_the_code_spacing(monkeypatch):
    from design.concrete import workflow

    assert _default_of(monkeypatch, None) == (False, workflow.SPACING_CODE)
    assert _default_of(monkeypatch, {"office_bar_spacing": False})[0] is False
    # the office rule stays the default for a user who chose it last time
    assert _default_of(monkeypatch, {"office_bar_spacing": True}) == (True, workflow.SPACING_OFFICE)


# --------------------------------------------------------------------------- #
# sdt schedules (v0.8.1): the schedules again from a stored design, no ETABS
# --------------------------------------------------------------------------- #
def test_schedules_are_drawn_again_from_the_design_data(monkeypatch, tmp_path):
    import pickle

    import pandas as pd

    from design.concrete import workflow
    from utilities import _gui_helpers

    store = workflow.DesignStore("m.EDB", 0.0, {"CONNECTIVITY": pd.DataFrame()},
                                 {"smrf": True, "column_bars": {"dmain": 25, "cover": 40}},
                                 pd.DataFrame({"UniqueName": ["B1"]}),
                                 pd.DataFrame({"UniqueName": ["C1"]}))
    data = tmp_path / "m - design data.pkl"
    data.write_bytes(pickle.dumps(store))
    asked = []
    monkeypatch.setattr(_gui_helpers, "select_open_path", lambda *a: str(data))
    monkeypatch.setattr(_gui_helpers, "select_output_directory", lambda *a: str(tmp_path))

    def option(title, message, options, default_index=0, **_):
        asked.append(options)
        return {"Which schedules should be drawn?": "Columns only"}.get(
            message, options[default_index])

    monkeypatch.setattr(_gui_helpers, "select_option", option)
    monkeypatch.setattr(workflow, "_last", lambda: {"section_scale": 25.0})
    monkeypatch.setattr(workflow, "_remember", lambda values: None)
    calls = []
    monkeypatch.setattr(workflow, "write_schedules",
                        lambda *a: calls.append(a) or ["Column_Schedule.dxf"])
    monkeypatch.setattr("utilities.run_summary.RunSummary.show", lambda *a, **k: "")
    assert workflow.run_schedules() == ["Column_Schedule.dxf"]
    _, folder, beams, columns, style, scale = calls[0]
    assert (folder, beams, columns, style, scale) == (str(tmp_path), False, True,
                                                      "crossties", 25.0)
    assert asked[0] == ["Beams and columns", "Beams and girders only", "Columns only"]
