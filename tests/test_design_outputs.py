"""Checks for the result files, the optional checks and the closing summaries.

Rounded numbers and stated blanks in the workbooks, the levels where the
capacity-design checks may be left out, the 75 mm cover of beams against
earth, the deflection table of the beam report, and the summary every
command prints. None of these tests needs ETABS.
"""

import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from design import beam_designer_aci318 as beam  # noqa: E402
from design import column_designer_aci318 as col  # noqa: E402
from design import concrete_workflow as cw  # noqa: E402
from etabs_api.workflows import analysis_forces as af  # noqa: E402
from etabs_api.workflows import design_loop as dl  # noqa: E402
from utilities._xlsx_values import cell_value, is_blank  # noqa: E402
from utilities.run_summary import RunSummary, listed  # noqa: E402


# --------------------------------------------------------------------------
# WORKBOOK CELLS
# --------------------------------------------------------------------------
@pytest.mark.parametrize("value, shown", [
    (0.006002320578210737, 0.01), (112.80456, 112.8), (600.0, 600), (12, 12),
    (0.4, 0.4), ("PASS", "PASS"), (True, True), (np.float64(3.14159), 3.14),
])
def test_numbers_are_rounded_to_two_decimals(value, shown):
    assert cell_value(value) == shown


def test_small_ratios_keep_four_decimals_when_asked():
    assert cell_value(0.006002320578210737, 4) == 0.006


@pytest.mark.parametrize("blank", [None, float("nan"), "", "  ", pd.NA, pd.NaT])
def test_a_blank_is_written_with_its_reason(blank):
    assert is_blank(blank)
    assert cell_value(blank, blank="N/A - circular column") == "N/A - circular column"


def test_an_infinite_ratio_is_written_as_text():
    assert cell_value(float("inf")) == "Infinite"


def test_column_blanks_say_why_they_do_not_apply():
    rectangular = {"Shape": "Rectangular", "Confinement_Check": "PASS"}
    circular = {"Shape": "Circular", "Confinement_Check": "PASS"}
    assert col.column_blank_reason("Diameter_mm", rectangular) == "N/A - rectangular column"
    assert col.column_blank_reason("Width_mm", circular) == "N/A - circular column"
    assert "rectangular hoops" in col.column_blank_reason("Confinement_18_7_5_d", rectangular)
    assert "0.3 Ag" in col.column_blank_reason("Confinement_18_7_5_c", rectangular)
    assert "spiral column" in col.column_blank_reason("Confinement_18_7_5_a", circular)
    assert "0.3 Ag" in col.column_blank_reason("Kf", rectangular)
    off = {"Shape": "Rectangular", "Confinement_Check": "N/A - non-SMRF detailing"}
    assert col.column_blank_reason("Kn", off) == "N/A - seismic design is off"


def _column_report(**extra) -> tuple[pd.DataFrame, list]:
    result = {"UniqueName": "C1", "Story": "2F", "Shape": "Rectangular", "Width_mm": 500.0,
              "Depth_mm": 500.0, "Diameter_mm": np.nan, "Longitudinal_Bars": 12,
              "Longitudinal_Bar_Layout": "12 single bars", "Design_Status": "PASS",
              "Reinforcement_Ratio": 0.023561944, "Confinement_18_7_5_b": 0.006002320578,
              "Confinement_18_7_5_d": np.nan, "Confinement_Check": "PASS", **extra}
    checks = [{"UniqueName": "C1", "Combo": "ULS 1", "End": end, "Pu_kN": 112.80456,
               "Mu2_kNm": 1.0, "Mu3_kNm": 66.1649, "phi_Mn_kNm": 420.0,
               "Flexure_Utilization": 0.157535, "Axial_Check": "PASS", "Flexure_Check": "PASS",
               "Strength_Check": "PASS"} for end in ("I", "J")]
    return col._build_consolidated_column_report(
        pd.DataFrame([result]), pd.DataFrame(checks), pd.DataFrame(), pd.DataFrame(),
        joint_na_reason="N/A - no beam frames into this end")


def test_column_workbook_has_no_empty_cell_and_no_long_decimal(tmp_path):
    from openpyxl import load_workbook

    report, groups = _column_report()
    path = col.write_column_results_xlsx(report, groups, str(tmp_path / "columns.xlsx"))
    sheet = load_workbook(path).active
    headers = [cell.value for cell in sheet[3]]
    row = dict(zip(headers, [cell.value for cell in sheet[4]]))
    assert all(value not in (None, "") for value in row.values())
    assert row["Pᵤ (kN)"] == 112.8
    assert row["ρ Longitudinal"] == 0.0236                        # four decimals for a ratio
    assert row["ρₛ,req from 18.7.5(b)"] == 0.006
    assert row["Diameter (mm)"] == "N/A - rectangular column"
    assert row["ρₛ,req from 18.7.5(d)"].startswith("N/A - rectangular hoops")
    assert row["Slenderness Check"] == col.SLENDERNESS_NOT_RUN     # stated, not blank


def test_beam_workbook_drops_the_unused_etabs_fields_and_states_its_blanks(tmp_path):
    from openpyxl import load_workbook

    results = pd.DataFrame({
        "Story": ["2F", "2F"], "UniqueName": ["2GX-1", "2GX-1"], "Face": ["TOP", "BOTTOM"],
        "Mu_left": [123.456789, 61.5], "Diameter": [np.nan, np.nan],
        "NumCBars2": [None, None], "Defl_roof_mm": [None, None],
        "Defl_ratio": [0.4884, 0.4884], "Deflection_Check": ["PASS", "PASS"],
        "Design_Status": ["OK", "OK"],
    })
    path = beam.write_beam_results_xlsx(results, str(tmp_path / "beams.xlsx"))
    sheet = load_workbook(path).active
    headers = [cell.value for cell in sheet[1]]
    assert "Diameter" not in headers and "NumCBars2" not in headers
    row = dict(zip(headers, [cell.value for cell in sheet[2]]))
    assert row["Mᵤ, left"] == 123.46
    assert row["Δ roof live (mm)"] == "N/A - not a roof beam"
    assert row["Δ / limit (governing)"] == 0.488


# --------------------------------------------------------------------------
# LEVELS WHERE THE CAPACITY-DESIGN CHECKS MAY BE LEFT OUT
# --------------------------------------------------------------------------
def _joint_rows() -> pd.DataFrame:
    row = {"Load_Combo": "ULS 1", "Column_End_Members": "C1:J:X", "Column_Beam_Ratio": 0.9,
           "Strong_Column_Check": "FAIL", "Joint_Shear_Utilization": 1.3,
           "Joint_Shear_Check": "FAIL", "Joint_Check_Reason": "", "BCC_Exempt": False,
           "Sum_Column_Mn_kNm": 90.0, "Sum_Beam_Mn_kNm": 100.0}
    return pd.DataFrame([{**row, "Joint_Point": "top"}, {**row, "Joint_Point": "mid"}])


def test_joints_the_user_left_out_are_marked_and_no_longer_fail():
    joints = col._skip_joint_checks(_joint_rows(), {"top": col.TOP_LEVEL_SKIPPED})
    top, mid = joints.iloc[0], joints.iloc[1]
    assert top["Strong_Column_Check"] == col.TOP_LEVEL_SKIPPED
    assert top["Joint_Shear_Check"] == col.TOP_LEVEL_SKIPPED
    assert bool(top["BCC_Exempt"]) and pd.isna(top["Column_Beam_Ratio"])
    assert pd.isna(top["Joint_Shear_Utilization"])
    assert mid["Strong_Column_Check"] == "FAIL" and mid["Column_Beam_Ratio"] == 0.9


def test_nothing_changes_when_every_level_is_checked():
    joints = _joint_rows()
    pd.testing.assert_frame_equal(col._skip_joint_checks(joints, {}), joints)


def test_the_report_states_the_level_that_was_left_out():
    joint = {"Joint_Point": "top", "Load_Combo": "ULS 1", "Column_End_Members": "C1:J:X",
             "Column_Beam_Ratio": np.nan, "Strong_Column_Check": col.TOP_LEVEL_SKIPPED,
             "Joint_Shear_Utilization": np.nan, "Joint_Shear_Check": col.TOP_LEVEL_SKIPPED,
             "Joint_Check_Reason": col.TOP_LEVEL_SKIPPED}
    result = {"UniqueName": "C1", "Story": "RD", "Width_mm": 500.0, "Depth_mm": 500.0,
              "Longitudinal_Bars": 12, "Design_Status": "PASS"}
    checks = [{"UniqueName": "C1", "Combo": "ULS 1", "End": end, "Pu_kN": 100.0,
               "Mu2_kNm": 1.0, "Mu3_kNm": 2.0, "phi_Mn_kNm": 50.0, "Flexure_Utilization": 0.1,
               "Axial_Check": "PASS", "Flexure_Check": "PASS", "Strength_Check": "PASS"}
              for end in ("I", "J")]
    report, _ = col._build_consolidated_column_report(
        pd.DataFrame([result]), pd.DataFrame(checks), pd.DataFrame(), pd.DataFrame([joint]))
    top = report[report["End"] == "J"].iloc[0]
    assert top["BCC_Ratio_X"] == col.TOP_LEVEL_SKIPPED
    assert top["BCC_Status"] == col.TOP_LEVEL_SKIPPED
    assert top["Joint_Shear_Status"] == col.TOP_LEVEL_SKIPPED
    assert top["Column_Design_Status"] == "PASS"


def test_a_combination_without_a_joint_check_says_where_to_look():
    """Only the governing combinations are evaluated at a joint."""
    joint = {"Joint_Point": "j", "Load_Combo": "ULS 2", "Column_End_Members": "C1:J:X",
             "Column_Beam_Ratio": 1.5, "Strong_Column_Check": "PASS",
             "Joint_Shear_Utilization": 0.5, "Joint_Shear_Check": "PASS"}
    result = {"UniqueName": "C1", "Story": "2F", "Design_Status": "PASS"}
    checks = [{"UniqueName": "C1", "Combo": combo, "End": end, "Pu_kN": 1.0, "Mu2_kNm": 1.0,
               "Mu3_kNm": 1.0, "phi_Mn_kNm": 5.0, "Flexure_Utilization": 0.2,
               "Axial_Check": "PASS", "Flexure_Check": "PASS", "Strength_Check": "PASS"}
              for combo in ("ULS 1", "ULS 2") for end in ("I", "J")]
    report, _ = col._build_consolidated_column_report(
        pd.DataFrame([result]), pd.DataFrame(checks), pd.DataFrame(), pd.DataFrame([joint]),
        joint_na_reason="N/A - no beam frames into this end")
    top = report[report["End"] == "J"].set_index("Combo")
    assert top.loc["ULS 2", "BCC_Status"] == "PASS"
    assert top.loc["ULS 1", "BCC_Status"] == col.JOINT_OTHER_COMBINATION
    bottom = report[report["End"] == "I"].set_index("Combo")
    assert bottom.loc["ULS 1", "BCC_Status"] == "N/A - no beam frames into this end"


def test_no_seismic_design_asks_nothing_about_the_levels():
    assert col.ask_capacity_check_levels(False) == (True, True)


def test_a_slenderness_failure_asks_the_loop_for_a_larger_column():
    rows = pd.DataFrame({"Flexure_Check": ["PASS"], "Axial_Check": ["PASS"],
                         "Shear_Check": ["PASS"], "SMRF_Dimension_Check": ["PASS"],
                         "Slenderness_Check": ["FAIL: SLENDERNESS - axis 3: magnifier 1.6"]})
    assert dl.column_needs(rows) == {"square"}
    rows["Slenderness_Check"] = "PASS - not slender"
    assert dl.column_needs(rows) == set()


# --------------------------------------------------------------------------
# BEAMS: 75 mm COVER AGAINST EARTH, DEFLECTION TABLE
# --------------------------------------------------------------------------
def _beam_tables():
    frame = pd.DataFrame({
        "Story": ["GF", "2F", "2F"], "UniqueName": ["GFTB-1", "2GX-1", "2-C1"],
        "SectProp": ["FTB", "G", "C"], "DesignType": ["Beam", "Beam", "Column"],
        "Width": [300, 300, 400], "Depth": [500, 500, 400],
    })
    conn = pd.DataFrame({
        "UniqueName": ["GFTB-1", "2GX-1", "2-C1"], "DesignType": ["Beam", "Beam", "Column"],
        "UniquePtI": ["a", "b", "z"], "UniquePtJ": ["b", "c", "a"],
    })
    return frame, conn


def test_beams_of_the_chosen_levels_get_the_earth_contact_cover():
    frame, conn = _beam_tables()
    bars = {"dm": 20, "ds": 10, "dw": 12, "fyw": 275, "cc": 40}
    table = beam.prepare_beam_table(frame, conn, bars, earth_cover_stories=["GF"])
    cover = dict(zip(table["UniqueName"], table["cc"]))
    assert cover == {"GFTB-1": 75.0, "2GX-1": 40.0}
    plain = beam.prepare_beam_table(frame, conn, bars)
    assert set(plain["cc"]) == {40.0}


def test_a_typed_cover_above_75_is_kept():
    frame, conn = _beam_tables()
    bars = {"dm": 20, "ds": 10, "dw": 12, "fyw": 275, "cc": 90}
    table = beam.prepare_beam_table(frame, conn, bars, earth_cover_stories=["GF"])
    assert set(table["cc"]) == {90.0}
    assert beam.ask_earth_cover_stories(["GF", "2F"], 90.0) == []   # nothing to ask


def test_the_beam_report_has_a_deflection_table():
    top = pd.Series({
        "Defl_Ie_left": 3.125e9, "Defl_Ie_mid": 2.039e9, "Defl_Ie_right": 3.125e9,
        "Defl_lambda": 1.454, "Defl_live_mm": 2.05, "Defl_live_limit_mm": 27.78,
        "Defl_roof_mm": None, "Defl_roof_limit_mm": None, "Defl_long_mm": 10.16,
        "Defl_long_limit_mm": 20.83, "Defl_ratio": 0.488, "Deflection_Check": "PASS"})
    table = beam._beam_deflection_table(top)
    assert table.title == "Deflection (ACI 24.2)"
    live, roof, long_term = table.rows
    assert live[1:] == ["2.05", "27.78", "0.07"]
    assert roof[1] == "N/A - not a roof beam"
    assert long_term[1:] == ["10.16", "20.83", "0.49"]
    assert "Deflection check: PASS" in table.note and "1.454" in table.note


def test_a_beam_without_deflection_data_says_why_and_old_results_get_no_table():
    none = beam._beam_deflection_table(pd.Series({"Deflection_Check": "NO DEF COMBOS",
                                                  "Defl_ratio": None}))
    assert none.rows == [["Deflection check", "NO DEF COMBOS"]]
    assert beam._beam_deflection_table(pd.Series({"Design_Status": "OK"})) is None


def test_the_column_plot_caption_reports_bcc_without_drawing_it():
    rows = pd.DataFrame({"BCC_Ratio_X": [1.50, 1.31, "N/A"], "BCC_Ratio_Y": [2.2, 2.0, "N/A"],
                         "BCC_Status": ["PASS", "PASS", "N/A"]})
    text = col._bcc_caption(rows)
    assert "along X 1.31, along Y 2.00" in text and "PASS" in text and "not drawn" in text
    rows["BCC_Status"] = ["PASS", "FAIL", "N/A"]
    assert "FAIL" in col._bcc_caption(rows)
    skipped = pd.DataFrame({"BCC_Ratio_X": [col.TOP_LEVEL_SKIPPED],
                            "BCC_Status": [col.TOP_LEVEL_SKIPPED]})
    assert col.TOP_LEVEL_SKIPPED in col._bcc_caption(skipped)


# --------------------------------------------------------------------------
# FORCES: THE DEAD LOAD SHARE AND THE SPECTRUM FLAG
# --------------------------------------------------------------------------
def _column_table(cases: dict[str, np.ndarray]) -> pd.DataFrame:
    rows = []
    for case, values in cases.items():
        for station, forces in zip((0.0, 3000.0), values):
            row = {"Story": "2F", "Column": "C1", "UniqueName": "C1", "OutputCase": case,
                   "CaseType": "LinRespSpec" if case.startswith("RSA") else "LinStatic",
                   "StepType": "Max" if case.startswith("RSA") else None, "StepNumber": None,
                   "Station": station, "Element": "C1", "ElemStation": station}
            row.update(dict(zip(af.FORCES, forces)))
            rows.append(row)
    return pd.DataFrame(rows)


def test_columns_carry_the_sustained_axial_force_and_the_spectrum_flag():
    dead = np.tile([-1000e3, 0.0, 0.0, 0.0, 0.0, 50e6], (2, 1))
    live = np.tile([-400e3, 0.0, 0.0, 0.0, 0.0, 20e6], (2, 1))
    rsa = np.tile([100e3, 0.0, 0.0, 0.0, 0.0, 80e6], (2, 1))
    forces = af.case_forces(_column_table({"DEAD": dead, "LIVE": live, "RSAX": rsa}), "Column")
    combos = [af.ComboTerms("G", {"DEAD": 1.2, "LIVE": 1.6}),
              af.ComboTerms("E", {"DEAD": 1.2, "LIVE": 0.5}, {"RSAX": 1.0})]
    out = af.combine(forces, combos, "column", sustained_cases=["DEAD"])
    gravity = out[out["Combo"] == "G"]
    np.testing.assert_allclose(gravity["P"], -1.2 * 1000e3 - 1.6 * 400e3)
    np.testing.assert_allclose(gravity["P_sustained"], -1.2 * 1000e3)
    assert not gravity["Spectral"].any()
    seismic = out[out["Combo"] == "E"]
    assert seismic["Spectral"].all() and seismic["Permutation"].nunique() == 8
    np.testing.assert_allclose(seismic["P_sustained"], -1.2 * 1000e3)


def test_beams_keep_the_same_columns_as_before():
    dead = np.tile([0.0, 10e3, 0.0, 0.0, 0.0, 50e6], (2, 1))
    table = _column_table({"DEAD": dead}).rename(columns={"Column": "Beam"})
    out = af.combine(af.case_forces(table, "Beam"), [af.ComboTerms("G", {"DEAD": 1.2})], "beam",
                     sustained_cases=["DEAD"])
    assert "P_sustained" not in out.columns and "Spectral" not in out.columns


def test_the_sustained_force_follows_the_compression_positive_convention():
    table = pd.DataFrame({"P": [-900.0], "P_sustained": [-600.0]})
    converted = col._to_compression_positive(table)
    assert converted["P"].iloc[0] == 900.0 and converted["P_sustained"].iloc[0] == 600.0


# --------------------------------------------------------------------------
# SUMMARIES
# --------------------------------------------------------------------------
def test_a_summary_lists_counts_failures_and_files(tmp_path, capsys):
    summary = RunSummary("sdt columns", r"C:\models\rev01.EDB")
    summary.add("Columns designed", 47).add("Failing", 2)
    summary.fail("Column shear: GF-C1, GF-C2").note("The model was not changed.")
    summary.file("Results", "rev01 - Column Design.xlsx")
    saved = tmp_path / "summary.txt"
    text = summary.show(str(saved))
    printed = capsys.readouterr().out
    for expected in ("SUMMARY - sdt columns", "Model: rev01.EDB", "Columns designed  47",
                     "Needs attention (1):", "Column shear: GF-C1, GF-C2",
                     "Note: The model was not changed.", "rev01 - Column Design.xlsx"):
        assert expected in text and expected in printed
    assert saved.read_text(encoding="utf-8").startswith("=" * 72)
    assert not summary.passed


def test_a_clean_run_says_nothing_needs_attention():
    assert "Nothing needs attention." in RunSummary("sdt beams").add("Beams", 3).text()


def test_long_lists_are_cut_with_the_total():
    assert listed(["a", "b", "c"]) == "a, b, c"
    assert listed([str(n) for n in range(20)], 3) == "0, 1, 2, ... (20 in all)"


def test_the_column_summary_groups_the_failing_columns_by_check():
    report = pd.DataFrame({
        "UniqueName": ["C1", "C1", "C2", "C3"], "Combo": ["U1", "U2", "U1", "U1"],
        "Column_Design_Status": ["FAIL", "PASS", "FAIL", "PASS"],
        "Flexure_Check": ["PASS"] * 4, "Axial_Check": ["PASS"] * 4,
        "Slenderness_Check": ["PASS - not slender", "PASS - not slender",
                              "FAIL: SLENDERNESS - axis 3", "PASS - slender, moments magnified"],
        "Shear_Check": ["FAIL", "PASS", "PASS", "PASS"],
        "BCC_Status": ["PASS"] * 4, "Joint_Shear_Status": ["PASS"] * 4,
        "Transverse_Reinforcement_Check": ["PASS"] * 4, "SMRF_Dimension_Check": ["PASS"] * 4,
    })
    text = cw.column_summary(report, "sdt columns", check_top_level=False).text()
    assert "Columns designed  3" in text and "Failing           2" in text
    assert "Column shear: C1" in text and "Slenderness (ACI 6.2.6): C2" in text
    assert "Slender columns   2" in text
    assert "not checked at the topmost level" in text


def test_the_beam_summary_groups_the_failing_beams_by_reason():
    results = pd.DataFrame({
        "UniqueName": ["B1", "B1", "B2", "B2", "B3", "B3"],
        "Design_Status": ["OK", "OK", "FAILED: DEFLECTION (ACI 24.2.2)", "OK", "OK", "OK"],
        "Deflection_Check": ["PASS", "PASS", "FAIL", "FAIL", "PASS", "PASS"]})
    text = cw.beam_summary(results, "sdt beams", combos=25, earth_stories=["GF"]).text()
    assert "Beams designed  3" in text and "Failing         1" in text
    assert "FAILED: DEFLECTION (ACI 24.2.2): B2" in text
    assert "75 mm cover on  GF" in text and "Deflection      1 failing of 3" in text


def test_the_loop_summary_reports_the_net_size_changes_not_every_step():
    changes = [dl.Change("2-C1", "CR_500X500_C04_G60", "CR_600X600_C04_G60", "grow"),
               dl.Change("2-C1", "CR_600X600_C04_G60", "CR_700X700_C04_G60", "grow"),
               dl.Change("2GX-1", "G_300X500_C04_G60", "G_300X600_C04_G60", "grow"),
               dl.Change("2GX-1", "G_300X600_C04_G60", "G_300X500_C04_G60", "shrink")]
    summary = {"status": "converged", "iterations": 7, "changes": changes,
               "failing": ["GF-C1"], "beams": None, "columns": None,
               "drift": {"failed": [], "path": "drift.txt"}}
    settings = dl.LoopSettings(combos=["U1", "U2"], check_foundation_level=False,
                               beam_earth_cover_stories=("GF",))
    text = dl.loop_summary(summary, settings, "m.EDB", "m - DESIGN.EDB", "log.txt", "out").text()
    assert "converged after 7 iterations" in text
    assert "4 in all; 1 members end with another size" in text
    assert "CR_500X500_C04_G60 -> CR_700X700_C04_G60" in text and "2-C1" in text
    assert "G_300X600" not in text                      # back to its first size: not listed
    assert "Members still failing: GF-C1" in text
    assert "every check passes" in text and "foundation level" in text
    assert "75 mm beam cover on: GF" in text


# --------------------------------------------------------------------------
# P-DELTA MUST BE ON FOR THE SWAY EFFECT
# --------------------------------------------------------------------------
class _Connector:
    def __init__(self, table):
        self.table = table

    def get_data(self, name):
        assert name == "P-Delta Option Definition"
        if isinstance(self.table, Exception):
            raise self.table
        return self.table


@pytest.mark.parametrize("method, expected", [
    ("Iterative Based on Loads", "Iterative Based on Loads"), ("None", "None")])
def test_the_pdelta_option_is_read_from_the_model(method, expected):
    table = pd.DataFrame({"AutoMethod": [method, None], "LoadPattern": ["SELFWEIGHT", "SIDL"]})
    assert cw.pdelta_method(_Connector(table)) == expected


def test_columns_go_on_without_asking_when_pdelta_is_on_or_unknown():
    on = pd.DataFrame({"AutoMethod": ["Iterative Based on Loads"]})
    assert cw.confirm_pdelta(_Connector(on), "Column Design") is True
    assert cw.confirm_pdelta(_Connector(RuntimeError("no table")), "Column Design") is True


def test_columns_ask_before_going_on_with_pdelta_off(monkeypatch):
    off = pd.DataFrame({"AutoMethod": ["None"]})
    asked = []

    def choose(title, prompt, options, default_index=0):
        asked.append(prompt)
        return options[choice]

    monkeypatch.setattr("utilities._gui_helpers.select_option", choose)
    choice = 0
    assert cw.confirm_pdelta(_Connector(off), "Column Design") is False
    choice = 1
    assert cw.confirm_pdelta(_Connector(off), "Column Design") is True
    assert "P-delta is off" in asked[0]


def test_a_command_opens_its_summary_in_a_window_and_tests_do_not(monkeypatch, capsys):
    opened = []
    monkeypatch.setattr("utilities._gui_helpers.show_summary",
                        lambda text, title="": opened.append((title, text)))
    summary = RunSummary("sdt design").add("Status", "converged").fail("GF-C1 still fails")
    summary.show()                                    # the default: terminal only
    assert opened == []
    summary.show(popup=True)
    assert opened[0][0] == "Summary - sdt design"
    assert "Needs attention (1):" in opened[0][1] and "GF-C1 still fails" in opened[0][1]
    assert capsys.readouterr().out.count("SUMMARY - sdt design") == 2   # still printed


def test_the_terminal_keeps_the_detail_and_the_window_gets_the_summary(monkeypatch, capsys):
    """A command that printed its detailed results shows the summary in the window only."""
    opened = []
    monkeypatch.setattr("utilities._gui_helpers.show_summary",
                        lambda text, title="": opened.append(text))
    print("=== Iteration 1: beams ===")                    # the command's own detail
    RunSummary("sdt design").add("Status", "converged").show(popup=True, echo=False)
    out = capsys.readouterr().out
    assert "=== Iteration 1: beams ===" in out and "SUMMARY - sdt design" not in out
    assert "SUMMARY - sdt design" in opened[0]


def test_the_design_loop_prints_every_log_line_again(tmp_path, capsys):
    bench = dl.Workbench.__new__(dl.Workbench)
    bench.log_path = str(tmp_path / "log.txt")
    bench.log("=== Iteration 3: columns ===")
    assert "=== Iteration 3: columns ===" in capsys.readouterr().out
    assert "Iteration 3" in (tmp_path / "log.txt").read_text(encoding="utf-8")


def test_a_summary_window_that_cannot_open_does_not_stop_the_command(monkeypatch):
    def broken(text, title=""):
        raise RuntimeError("no display")

    monkeypatch.setattr("utilities._gui_helpers.show_summary", broken)
    assert "SUMMARY - sdt beams" in RunSummary("sdt beams").show(popup=True)


def test_a_summary_that_cannot_open_its_window_falls_back_to_the_terminal(monkeypatch, capsys):
    def broken(text, title=""):
        raise RuntimeError("no display")

    monkeypatch.setattr("utilities._gui_helpers.show_summary", broken)
    RunSummary("sdt check").show(popup=True, echo=False)
    assert "SUMMARY - sdt check" in capsys.readouterr().out


def test_the_loop_draws_the_column_schedule_with_the_chosen_tie_style(tmp_path, monkeypatch):
    """sdt design asks the interior tie style, like sdt columns, and uses it."""
    drawn = []
    monkeypatch.setattr(col, "export_column_cad_drawings",
                        lambda report, folder, dmain, cover, smrf, style, connectivity:
                        drawn.append(style) or [])
    monkeypatch.setattr(col, "write_column_results_xlsx", lambda *a, **k: "x")
    monkeypatch.setattr(col, "export_column_pdf", lambda *a, **k: None)
    working = tmp_path / "m - DESIGN.EDB"
    working.write_text("")
    bench = dl.Workbench.__new__(dl.Workbench)
    bench.progress = lambda text: None
    bench.settings = dl.LoopSettings(combos=["U1"], inner_tie_style="hoops")
    bench.tables = {"CONNECTIVITY": pd.DataFrame()}
    bench.column_groups = [("G", ["UniqueName"])]
    columns = pd.DataFrame({"UniqueName": ["C1"]})
    dl.save_final_design(bench, {"beams": None, "columns": columns}, str(working), str(tmp_path))
    assert drawn == ["hoops"]
    assert dl.LoopSettings(combos=[]).inner_tie_style == "crossties"   # the default
