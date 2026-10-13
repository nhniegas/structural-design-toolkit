"""Cell values of the result workbooks: rounded numbers, stated blanks, and
the marking of the members that fail.

A result file should not show a long unrounded number, and an empty cell
should not be mistaken for a missing calculation: a value that does not
apply is written as ``N/A`` with the reason. A member that fails is seen
while scrolling: its rows are light red, and its name and the results that
fail are red. A member that passes is left as it is.
"""

from __future__ import annotations

import math
import numbers

import pandas as pd

DEFAULT_DECIMALS = 2
NOT_APPLICABLE = "N/A"
# A result that starts with one of these is not a pass.
FAILING = ("FAIL", "ERROR", "BLOCKED", "NOT OK")
FAIL_ROW_FILL, FAIL_CELL_FILL, FAIL_FONT = "FDECEA", "FFC7CE", "9C0006"


def is_blank(value) -> bool:
    """None, NaN, NaT or an empty text."""
    if value is None:
        return True
    if isinstance(value, str):
        return not value.strip()
    try:
        return bool(value != value)  # NaN and NaT are not equal to themselves
    except (TypeError, ValueError):
        return value is pd.NA


def cell_value(value, decimals: int = DEFAULT_DECIMALS, blank: str = NOT_APPLICABLE):
    """The value to write in a result cell.

    Whole numbers stay whole, other numbers are rounded to ``decimals``, a
    blank becomes ``blank`` and text is written as it is.
    """
    if is_blank(value):
        return blank
    if isinstance(value, bool) or not isinstance(value, numbers.Number):
        return value
    if isinstance(value, numbers.Integral):
        return int(value)
    number = float(value)
    if math.isinf(number):
        return "Infinite" if number > 0 else "-Infinite"
    rounded = round(number, decimals)
    return int(rounded) if rounded == int(rounded) and abs(rounded) >= 1 else rounded


def is_failure(value) -> bool:
    """Whether a cell holds a result that fails (``FAIL``, ``FAILED: ...``, ``ERROR``)."""
    return isinstance(value, str) and value.strip().upper().startswith(FAILING)


def mark_failures(sheet, first_row: int, name_column: int) -> list[str]:
    """Mark the members that fail in a result sheet, and return their names.

    The rows from ``first_row`` down are those of the members, named in
    ``name_column`` (from 1); a member can have several rows. A member fails
    when a cell of any of its rows holds a failing result. Each of its rows is
    then filled light red, and its name and the failing cells are red and bold.
    """
    from copy import copy

    from openpyxl.styles import Font, PatternFill

    rows_of: dict[str, list[tuple]] = {}
    failing_cells: dict[str, list] = {}
    for row in sheet.iter_rows(min_row=first_row):
        name = str(row[name_column - 1].value)
        rows_of.setdefault(name, []).append(row)
        cells = [cell for cell in row if is_failure(cell.value)]
        if cells:
            failing_cells.setdefault(name, []).extend(cells)
    light = PatternFill("solid", fgColor=FAIL_ROW_FILL)
    strong = PatternFill("solid", fgColor=FAIL_CELL_FILL)
    red = Font(color=FAIL_FONT, bold=True)
    styles: dict = {}

    def paint(cell, kind: str) -> None:
        # Setting a fill on a cell is slow in openpyxl, and a large model has
        # millions of cells: the style of each kind of cell is worked out once.
        key = (tuple(cell._style) if cell._style is not None else None, kind)
        if key not in styles:
            cell.fill = light if kind == "row" else strong
            if kind == "cell":
                cell.font = red
            styles[key] = copy(cell._style)
        else:
            cell._style = copy(styles[key])

    for name, cells in failing_cells.items():
        for row in rows_of[name]:
            for cell in row:
                paint(cell, "row")
            cells.append(row[name_column - 1])
        for cell in cells:
            paint(cell, "cell")
    return list(failing_cells)
