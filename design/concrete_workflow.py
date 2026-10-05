"""Concrete beam and column design from the terminal (``python main.py beams`` / ``columns``).

``sdt beams``
    Attaches to the model open in ETABS, asks the inputs, reads the forces
    from the analysis results (with the service loads for deflection), designs
    every beam, and saves the results (.xlsx), the calculation report (.pdf)
    and the beam schedules (.dxf) in the folder you choose.

``sdt deflection``
    Checks only the deflection of the beams and girders: the bars of the
    last ``sdt beams`` of the model with the service moments read again from
    ETABS, and a results file of the deflections.

``sdt columns``
    Reads what ``sdt beams`` stored for the model, asks the column inputs,
    designs every column (the SMRF joint checks use the beam bars), and saves
    the results, the calculation report and the column schedule.

Between the two steps everything is kept in a binary store next to the model,
``<model> - design data.pkl``: the extracted tables, the inputs and the
results. It is internal; the files you save are made from it.
"""

from __future__ import annotations

import json
import os
import pickle
import re
import sys
from dataclasses import dataclass, field

import pandas as pd

if __package__ in (None, ""):
    sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from design import dcr_targets  # noqa: E402
from utilities.latex_help import missing_pdf_reason  # noqa: E402

STORE_SUFFIX = " - design data.pkl"
# Columns that hold names: kept as text (everything else that looks like a
# number becomes one, as the old workbook did).
TEXT_COLUMNS = {
    "UniqueName", "Label", "Story", "Combo", "SectProp", "Name", "GUID", "Element",
    "UniquePtI", "UniquePtJ", "UniquePt1", "UniquePt2", "UniquePt3", "UniquePt4",
    "BeamBay", "ColumnBay", "WallBay", "PointBay", "Beam", "Column", "DesignType",
    "Shape", "Material", "Tributary method", "Governed by",
}


def numeric_like(table: pd.DataFrame) -> pd.DataFrame:
    """Number-like text as numbers, except in the name columns."""
    table = table.copy()
    for column in table.columns:
        if column in TEXT_COLUMNS or table[column].dtype != object:
            continue
        values = table[column].replace({"None": None, "": None})
        converted = pd.to_numeric(values, errors="coerce")
        if converted.notna().sum() == values.notna().sum():
            table[column] = converted
        else:
            table[column] = values
    return table


@dataclass
class DesignStore:
    """What the beam and column steps share for one model."""

    model_path: str
    model_saved: float                      # model file time when extracted
    tables: dict[str, pd.DataFrame] = field(default_factory=dict)
    inputs: dict = field(default_factory=dict)
    beam_results: pd.DataFrame | None = None
    column_report: pd.DataFrame | None = None
    column_groups: list | None = None
    joint_results: pd.DataFrame | None = None

    @staticmethod
    def path_for(model_path: str) -> str:
        return os.path.splitext(os.path.normpath(model_path))[0] + STORE_SUFFIX

    @property
    def path(self) -> str:
        return self.path_for(self.model_path)

    def save(self) -> str:
        with open(self.path, "wb") as handle:
            pickle.dump(self, handle, protocol=pickle.HIGHEST_PROTOCOL)
        return self.path

    @classmethod
    def load(cls, model_path: str) -> "DesignStore | None":
        try:
            with open(cls.path_for(model_path), "rb") as handle:
                return pickle.load(handle)
        except (OSError, pickle.UnpicklingError, EOFError, AttributeError):
            return None

    def is_stale(self) -> bool:
        """True when the model file was saved after the extraction."""
        try:
            return os.path.getmtime(self.model_path) > self.model_saved + 1.0
        except OSError:
            return False


def extract(connector, combos: list[str], options, members: list[str] | None = None,
            progress=None, deflection_roles: dict[str, str] | None = None
            ) -> tuple[dict[str, pd.DataFrame], list[str]]:
    """Every table the beam and column design need, from the open model.

    ``deflection_roles`` maps each standard deflection combination to the
    model's own combination for that role (``model_inputs.ask_deflection_roles``).
    """
    from etabs_api.workflows.exporter import ETABSDataExporter

    exporter = ETABSDataExporter(connector)
    exporter.deflection_roles = deflection_roles
    if members is None:
        members = exporter.get_available_members()
    if progress:
        progress("Reading the analysis results and building the factored forces")
    notes = exporter.extract_all(combos, members, options)
    return {name: numeric_like(table) for name, table in exporter.tables.items()}, notes


# =============================================================================
# DIALOGS
# =============================================================================
BEAM_FIELDS = {
    "Beam main bar (mm)": ("dm", 28),
    "Beam stirrup (mm)": ("ds", 12),
    "Beam web (side) bar (mm)": ("dw", 12),
    "Web bar fy (MPa)": ("fyw", 414),
    "Beam cover (mm)": ("cc", 40),
}
COLUMN_FIELDS = {
    "Column main bar (mm)": ("dmain", 25),
    "Column tie (mm)": ("dties", 12),
    "Column cover (mm)": ("cover", 40),
}


def _settings_file() -> str:
    from utilities.user_settings import settings_path as user_settings_path

    return user_settings_path("concrete_design.json")


def _last() -> dict:
    try:
        with open(_settings_file(), encoding="utf-8") as handle:
            return json.load(handle)
    except (OSError, ValueError):
        return {}


def _remember(values: dict) -> None:
    saved = _last()
    saved.update(values)
    os.makedirs(os.path.dirname(_settings_file()), exist_ok=True)
    with open(_settings_file(), "w", encoding="utf-8") as handle:
        json.dump(saved, handle, indent=2)


def _ask_numbers(title: str, prompt: str, fields: dict, last: dict) -> dict | None:
    from utilities._gui_helpers import enter_values, show_warning

    defaults = {label: f"{float(last.get(key, value)):g}" for label, (key, value) in fields.items()}
    while True:
        typed = enter_values(title, prompt, list(fields), defaults)
        if typed is None:
            return None
        try:
            values = {key: float(typed[label]) for label, (key, _) in fields.items()}
            if min(values.values()) <= 0:
                raise ValueError
            return values
        except ValueError:
            show_warning("Every value must be a positive number.", title=title)
            defaults = typed


def attach_etabs():
    """The ETABS session that is running, as a connector; None with a message."""
    from etabs_api.core.connection import attach_running_etabs

    return attach_running_etabs("Design")


def pdelta_method(connector) -> str:
    """The P-delta option of the open model ("None" when it is off, "" when it
    cannot be read)."""
    try:
        table = connector.get_data("P-Delta Option Definition")
    except RuntimeError:
        return ""
    if table is None or "AutoMethod" not in getattr(table, "columns", ()):
        return ""
    methods = table["AutoMethod"].dropna().astype(str)
    return methods.iloc[0] if len(methods) else "None"


def confirm_pdelta(connector, title: str) -> bool:
    """False when the user stops because P-delta is off in the model.

    The column slenderness adds only the member effect; the sway effect must
    be in the analysis forces (ACI 6.7).
    """
    from utilities._gui_helpers import select_option

    method = pdelta_method(connector)
    if not method or "iterative" in method.lower():
        return True
    chosen = select_option(
        title, f"P-delta is off in this model (P-Delta option: {method}). The column "
        "slenderness adds the member effect only; the sway of the storeys must come from a "
        "P-delta analysis, or the second-order moments are understated. sdt setup defines it.",
        ["Stop (turn P-delta on, analyse and run sdt beams again)",
         "Continue without the sway effect"])
    return bool(chosen) and chosen.startswith("Continue")


def model_path_of(connector) -> str | None:
    path = os.path.splitext(os.path.normpath(str(connector.sap_model.GetModelFilename())))[0]
    path += ".EDB"
    return path if os.path.isfile(path) else None


def _gravity_options(combos: list[str]) -> list[str]:
    """Combinations offered for the beam seismic shear: no seismic or wind term."""
    plain = [c for c in combos if not re.search(r"\b(EQ|RSA)\d|\bW[XY]\b", c)]
    return plain or combos


# =============================================================================
# BEAMS
# =============================================================================
def run_beams() -> DesignStore | None:
    """Extract, design the beams and save the results, the calculations and the schedules."""
    from design.beam_designer_aci318 import (
        ask_deflection_limit,
        ask_earth_cover_stories,
        design_beams,
        export_beam_dxf,
        export_beam_pdf,
        write_beam_results_xlsx,
    )
    from etabs_api.workflows.analysis_forces import ask_force_options
    from utilities._gui_helpers import (
        LoadingWindow,
        select_option,
        select_output_directory,
        show_warning,
    )

    title = "Beam Design"
    connector = attach_etabs()
    if connector is None:
        return None
    model_path = model_path_of(connector)
    if model_path is None:
        show_warning("Save the ETABS model first: it has no file yet.", title=title)
        return None
    last = _last()

    ready = prepare_model(connector, model_path, title, "sdt beams", last)
    if ready is None:
        return None
    combos, seismic_key = ready.combos, ready.seismic_key
    options = ask_force_options()
    if options is None:
        return None
    smrf = select_option(title, "SMRF (seismic) design? Gravity beams (on other beams only) "
                         "are designed for gravity in any case.", ["Yes", "No"],
                         default_index=0 if last.get("smrf", True) else 1)
    if smrf is None:
        return None
    smrf = smrf == "Yes"
    gravity = None
    if smrf:
        choices = ready.gravity_choices[:24]  # what one choice dialog can show
        gravity = select_option(title, VE_GRAVITY_PROMPT, choices,
                                default_index=choices.index(ready.gravity_default)
                                if ready.gravity_default in choices else 0)
        if gravity is None:
            return None
    bars = _ask_numbers(title, "Beam bars and cover.", BEAM_FIELDS, last)
    if bars is None:
        return None
    earth_stories = ask_earth_cover_stories(beam_stories(connector), bars["cc"],
                                            last.get("beam_earth_cover_stories"))
    if earth_stories is None:
        return None
    divisor = ask_deflection_limit()
    if divisor is None:
        return None
    beam_types = (dcr_targets.GIRDER, dcr_targets.BEAM)
    targets = ask_dcr_targets(model_path, title, beam_types)
    if targets is None:
        return None
    carrier_depth = ask_carrier_depth(title, last)
    if carrier_depth is None:
        return None
    compatibility_torsion = ask_compatibility_torsion(title, last)
    if compatibility_torsion is None:
        return None
    office_bar_spacing = ask_bar_spacing(title, last)
    if office_bar_spacing is None:
        return None
    folder = select_output_directory("Folder for the beam results, calculations and schedules")
    if not folder:
        return None
    _remember({**({"seismic": seismic_key} if seismic_key else {}), "smrf": smrf,
               "gravity_combo": gravity, **bars, "beam_earth_cover_stories": earth_stories,
               "carrier_depth": carrier_depth,
               "compatibility_torsion": compatibility_torsion,
               "office_bar_spacing": office_bar_spacing})

    stem = os.path.splitext(os.path.basename(model_path))[0]
    with dcr_targets.use(targets), LoadingWindow("Beam design") as window:
        tables, notes = extract(connector, combos, options, progress=window.update,
                                deflection_roles=ready.deflection_roles)
        store = DesignStore(model_path, os.path.getmtime(model_path), tables, {
            "combos": combos, "seismic": seismic_key, "smrf": smrf, "gravity_combo": gravity,
            "beam_bars": bars, "long_limit": divisor,
            "beam_earth_cover_stories": earth_stories,
            "deflection_roles": ready.deflection_roles,
            "dcr_targets": targets.to_saved(), "carrier_depth": carrier_depth,
            "compatibility_torsion": compatibility_torsion,
            "office_bar_spacing": office_bar_spacing,
            "sources": {"model": list(ready.sources.model),
                        "answered": list(ready.sources.answered),
                        "assumed": list(ready.sources.assumed)},
        })
        results = design_beams(tables, smrf, gravity, bars, divisor, progress=window.update,
                               earth_cover_stories=earth_stories, carrier_depth=carrier_depth,
                               compatibility_torsion=compatibility_torsion,
                               office_bar_spacing=office_bar_spacing)
        store.beam_results = results
        store.save()
        window.update("Saving 1 of 3: the results workbook (.xlsx)")
        xlsx = write_beam_results_xlsx(results, os.path.join(folder, f"{stem} - Beam Design.xlsx"))
        window.update("Saving 2 of 3: the beam schedules (.dxf)")
        dxf = export_beam_dxf(results, folder)
        window.update("Saving 3 of 3: the calculation report (.pdf, LaTeX)")
        pdf = export_beam_pdf(results, os.path.join(folder, f"{stem} - Beam Calculations.pdf"),
                              smrf, gravity)
    summary = beam_summary(results, "sdt beams", model_path, len(combos), earth_stories)
    add_targets_to(summary, targets, beam_types)
    if carrier_depth:
        summary.add("Carrier depth", "a beam is at least as deep as the beams it carries")
    if compatibility_torsion:
        summary.add("Beam torsion", "at most phi Tcr (compatibility torsion, ACI 22.7.3.2)")
        summary.note(TORSION_NOTE)
    summary.add("Beam bar spacing", "office rule, 150 mm clear" if office_bar_spacing
                else "crack control only (ACI 24.3.2)")
    ready.sources.add_to(summary)
    for note in notes:
        summary.note(note)
    summary.file("Results", xlsx)
    summary.file("Schedules", ", ".join(os.path.basename(p) for p in dxf))
    summary.file("Calculations", pdf or missing_pdf_reason())
    summary.file("Design data", store.path)
    summary.show(os.path.join(folder, f"{stem} - Beam Design summary.txt"), popup=True)
    return store


@dataclass
class ModelReady:
    """What a design command settled about the model before its own questions."""

    combos: list[str]
    seismic_key: str | None          # EQ / RSA / Both for the ULS names; None when picked
    gravity_choices: list[str]
    gravity_default: str | None      # the choice closest to 1.2 D + f1 L, for Ve
    deflection_roles: dict[str, str]
    sources: object                  # model_inputs.Sources


VE_GRAVITY_PROMPT = (
    "Gravity combination for the beam seismic shear Ve. It is the factored gravity load on "
    "the span, 1.2 D + f1 L (ACI 318-14 18.6.5.1, NSCP 2015 203.3); the closest combination "
    "is selected. A combination without live load, such as 1.4 D, leaves the live load out.")
SEISMIC_CHOICES = {"Static (EQ)": "EQ", "Response spectrum (RSA)": "RSA", "Both": "Both"}
# The first choice changes nothing in the model. Tagging renames the members of
# the open model, unlocks it and drops its results.
UNNAMED_OPTIONS = {
    "Design them as they are, with their ETABS numbers": "numbers",
    "Tag them now, in this model (it renames the members and drops the results)": "tag",
    "Leave the unnamed members out": "skip",
}


def member_names(connector) -> list[str]:
    """Every beam and column of the open model, named or not."""
    names = []
    for table_name in ("Beam Object Connectivity", "Column Object Connectivity"):
        try:
            table = connector.get_data(table_name)
        except RuntimeError:
            continue
        if table is not None and "UniqueName" in getattr(table, "columns", ()):
            names += table["UniqueName"].astype(str).tolist()
    return list(dict.fromkeys(names))


def prepare_model(connector, model_path: str, title: str, command: str,
                  last: dict | None = None) -> ModelReady | None:
    """Settle what the design needs from the model before the command's own questions.

    On a model set up by ``sdt setup`` and tagged this only asks which seismic
    combinations to design for, as before. On any other model it first shows
    what was found and what is missing, then asks for it: the strength
    combinations, the deflection combinations, and what to do with members
    that still have their ETABS number. The answers are saved with the model.
    None when a dialog is closed or the user stops.
    """
    from etabs_api.workflows import model_inputs as mi
    from utilities._gui_helpers import select_option

    last = last or {}
    combinations = mi.read_combinations(connector)
    names = combinations.names
    standard = mi.standard_strength(names)
    saved = mi.load(model_path)
    roles_now = mi.deflection_roles(names, saved.get("deflection"))
    members = member_names(connector)
    unnamed = mi.unnamed_members(members)
    walls = mi.wall_count(connector)
    sources = mi.Sources()

    ready = mi.Readiness(command)
    if standard:
        ready.found.append(f"{len(standard)} strength combinations named ULS")
    else:
        ready.missing.append(
            "strength combinations: none is named ULS. You pick them from the model's "
            f"{len(combinations.linear)} linear combinations"
            + (f" ({len(combinations.envelopes)} envelopes cannot be designed for)"
               if combinations.envelopes else ""))
    absent = [mi.DEFLECTION_ROLES[role] for role, name in roles_now.items() if name is None]
    if absent:
        ready.missing.append("deflection combinations for: " + ", ".join(absent)
                             + ". You pick an existing one for each, or let the toolkit add it")
    else:
        ready.found.append("the four deflection combinations")
    if unnamed:
        ready.missing.append(
            f"names: {len(unnamed)} of {len(members)} beams and columns still have their "
            "ETABS number. You choose to tag them, design them as they are, or leave them out")
    else:
        ready.found.append(f"{len(members)} named beams and columns")
    if walls:
        ready.warnings.append(f"{walls} wall panels are in the model: walls are not designed")
    method = pdelta_method(connector)
    if method and "iterative" not in method.lower():
        ready.warnings.append(f"P-delta is off (option: {method}): the column slenderness "
                              "needs the sway effect from a P-delta analysis")
    if not ready.confirm(title):
        return None

    # ---- members without a name ----
    if unnamed:
        options = [text for text, key in UNNAMED_OPTIONS.items()
                   if key != "skip" or len(unnamed) < len(members)]
        chosen = select_option(
            title, f"{len(unnamed)} of {len(members)} beams and columns still have the number "
            "ETABS gave them. Tagging gives each its level, type and number (2GX-1, 3-C5), "
            "which the schedules use. Without tags the design still runs: beam lines are then "
            "found from the geometry.", options)
        if chosen is None:
            return None
        choice = UNNAMED_OPTIONS[chosen]
        if choice == "tag":
            from etabs_api.workflows.frame_tagger import auto_tag_frames

            if auto_tag_frames(in_place=True, connector=connector) is None:
                return None
            sources.answered.append("members tagged now")
        elif choice == "numbers":
            connector.include_numeric_members = True
            sources.answered.append(f"{len(unnamed)} members designed with their ETABS numbers")
        else:
            sources.answered.append(f"{len(unnamed)} unnamed members left out")

    # ---- strength combinations ----
    seismic_key = None
    if standard:
        default = ["EQ", "RSA", "Both"].index(last.get("seismic", "RSA")) \
            if last.get("seismic", "RSA") in ("EQ", "RSA", "Both") else 1
        seismic = select_option(title, "Seismic combinations to design for (gravity and wind "
                                "ULS combinations are always included):",
                                list(SEISMIC_CHOICES), default_index=default)
        if seismic is None:
            return None
        seismic_key = SEISMIC_CHOICES[seismic]
        combos = mi.standard_strength(names, seismic_key)
        sources.model.append(f"{len(combos)} strength combinations (ULS names)")
    else:
        picked = mi.ask_strength_combinations(connector, model_path, title, combinations)
        if picked is None:
            return None
        combos, origin = picked
        sources.answered.append(f"{len(combos)} strength combinations ({origin})")
    if not combos:
        return None

    # ---- deflection combinations ----
    roles = mi.ask_deflection_roles(connector, model_path, title, combinations)
    if roles is None:
        return None
    own = [role for role, name in roles.items() if name != role]
    added = [role for role, name in roles.items() if name == role and role not in names]
    if own:
        sources.answered.append(
            "deflection: " + ", ".join(f"{mi.DEFLECTION_ROLES[r]} = {roles[r]}" for r in own))
    if added:
        sources.assumed.append("deflection combinations added by the toolkit from the pattern "
                               "types: " + ", ".join(added))
    if not own and not added:
        sources.model.append("deflection combinations (DEF names)")
    gravity_choices = mi.gravity_options(combinations, combos)
    return ModelReady(combos, seismic_key, gravity_choices,
                      mi.ve_gravity_default(combinations, gravity_choices[:24]), roles, sources)


CODE_LIMITS = "Use the code limits (every ratio 1.00)"
SET_TARGETS = "Set target ratios"
CARRIER_NO = "No - the depths are as modelled"
CARRIER_YES = "Yes - a beam is at least as deep as the beams it carries"


def ask_dcr_targets(model_path: str, title: str, members: tuple[str, ...]
                    ) -> dcr_targets.Targets | None:
    """The target ratios of these member types; None when a dialog is closed.

    The code passes a check at a ratio of 1.00. A target below it is a margin
    the engineer chooses, by member type and check. The targets are saved
    with the model and offered again.
    """
    from etabs_api.workflows import model_inputs as mi
    from utilities._gui_helpers import enter_values, select_option, show_warning

    saved = dcr_targets.Targets.from_saved(mi.load(model_path).get("dcr"))
    mine = saved.lines(members)
    options = [CODE_LIMITS] + (["Use the targets saved for this model"] if mine else []) \
        + [SET_TARGETS]
    chosen = select_option(
        title, "Target ratios (demand / capacity). The code passes a check at 1.00; a lower "
        "target is a margin of your own, by member type and check. The members are designed "
        "to stay at or below it."
        + ("\n\nSaved for this model:\n  " + "\n  ".join(mine) if mine else ""),
        options, default_index=1 if mine else 0)
    if chosen is None:
        return None
    if chosen == CODE_LIMITS:
        return dcr_targets.Targets()
    if chosen != SET_TARGETS:
        return saved
    targets = dcr_targets.Targets(dict(saved.values))
    for member in members:
        checks = dcr_targets.CHECKS[member]
        labels = {words: check for check, words in checks.items()}
        defaults = {words: f"{targets.get(member, check):g}" for words, check in labels.items()}
        while True:
            typed = enter_values(
                title, f"{dcr_targets.TYPE_NAMES[member]}: target ratio of each check. Leave "
                "1 for the code limit; the strong column ratio is at least "
                f"{dcr_targets.code_limit(dcr_targets.STRONG_COLUMN):g}.",
                list(labels), defaults)
            if typed is None:
                return None
            try:
                for words, check in labels.items():
                    targets.set(member, check, float(typed[words]))
                break
            except ValueError as error:
                defaults = dict(typed)
                show_warning(f"{dcr_targets.TYPE_NAMES[member]}: {error} Type a number in "
                             "every box.", title=title)
    mi.save(model_path, dcr=targets.to_saved())
    return targets


def ask_carrier_depth(title: str, last: dict | None = None) -> bool | None:
    """Whether a beam must be at least as deep as the beams it carries."""
    from utilities._gui_helpers import select_option

    chosen = select_option(
        title, "Should a girder or beam be at least as deep as the beams that frame into it? "
        "With a shallower carrier, the bottom bars of the carried beam pass below the "
        "carrier's bottom bars and cannot rest on them. This is a detailing rule of your "
        "own, not a code clause: a carrier that is shallower fails.",
        [CARRIER_NO, CARRIER_YES],
        default_index=1 if (last or {}).get("carrier_depth") else 0)
    return None if chosen is None else chosen == CARRIER_YES


TORSION_ANALYSIS = "The analysis torsion (as ETABS gives it)"
TORSION_COMPATIBILITY = "At most phi Tcr: compatibility torsion (ACI 22.7.3.2)"
TORSION_NOTE = (
    "Beam torsion was limited to phi Tcr (ACI 22.7.3.2), which holds where the torsion can "
    "redistribute after cracking. ACI 22.7.3.3 then requires the adjoining members to be "
    "designed for the redistributed moments and shears: an analysis with the torsional "
    "stiffness of the beams reduced (a small J modifier) gives them. Cantilevers keep their "
    "analysis torsion.")


def ask_compatibility_torsion(title: str, last: dict | None = None) -> bool | None:
    """Whether the beam torsion is designed as compatibility torsion."""
    from utilities._gui_helpers import select_option

    chosen = select_option(
        title, "Which torsion are the beams designed for?\n\nThe analysis torsion is right "
        "for equilibrium torsion, and for a model whose beams have a reduced torsional "
        "stiffness (J modifier). In a model with the full torsional stiffness, beams that "
        "frame into each other pick up large compatibility torsion; ACI 22.7.3.2 lets that be "
        "taken as at most phi Tcr, because it redistributes when the beam cracks. Cantilevers "
        "keep the analysis torsion either way.",
        [TORSION_ANALYSIS, TORSION_COMPATIBILITY],
        default_index=1 if (last or {}).get("compatibility_torsion") else 0)
    return None if chosen is None else chosen == TORSION_COMPATIBILITY


SPACING_OFFICE = "Office rule: at most 150 mm clear between the bars of a face"
SPACING_CODE = "Code only: crack control spacing (ACI 24.3.2)"


def ask_bar_spacing(title: str, last: dict | None = None) -> bool | None:
    """Whether the beams follow the office rule of 150 mm clear between bars.
    True for the office rule, False for the code spacing, None when closed."""
    from utilities._gui_helpers import select_option

    chosen = select_option(
        title, "How many bars does a beam face need for spacing?\n\nThe office rule keeps "
        "the clear spacing at 150 mm or less, so a wide beam gets more bars than its strength "
        "needs. The code limits the spacing for crack control only (ACI 24.3.2, about 250 mm "
        "centre to centre). The extra bars of the office rule raise the probable moments of "
        "the beam, and with them the capacity shear of the beam and the joint shear and "
        "capacity shear of the columns.",
        [SPACING_OFFICE, SPACING_CODE],
        default_index=1 if (last or {}).get("office_bar_spacing") is False else 0)
    return None if chosen is None else chosen == SPACING_OFFICE


def add_targets_to(summary, targets: dcr_targets.Targets, members: tuple[str, ...]) -> None:
    """List the target ratios in a summary, with what a beam target does to the columns."""
    lines = targets.lines(members)
    for line in lines:
        summary.add("Target ratio", line)
    beams = {dcr_targets.GIRDER, dcr_targets.BEAM} & set(members)
    if any((member, dcr_targets.FLEXURE) in targets.values for member in beams):
        summary.note("A flexure target below 1 adds beam bars. More beam steel raises the "
                     "probable moments, so the column shear Ve and the joint shear demand "
                     "rise with it.")


def beam_stories(connector) -> list[str]:
    """The story of every beam of the open model, in the order ETABS lists them."""
    table = connector.get_data("Beam Object Connectivity")
    if table is None or "Story" not in getattr(table, "columns", ()):
        return []
    return list(dict.fromkeys(table["Story"].dropna().astype(str)))


def beam_summary(results: pd.DataFrame, command: str, model_path: str | None = None,
                 combos: int | None = None, earth_stories=()):
    """The closing summary of a beam design: counts and what fails, by reason."""
    from utilities.run_summary import RunSummary, listed

    summary = RunSummary(command, model_path)
    status = results.groupby(results["UniqueName"].astype(str))["Design_Status"].apply(
        lambda s: next((v for v in s.astype(str) if v != "OK"), "OK"))
    failed = status[status != "OK"]
    summary.add("Beams designed", len(status))
    if combos is not None:
        summary.add("Combinations", combos)
    summary.add("Passing", int((status == "OK").sum()))
    summary.add("Failing", len(failed))
    if earth_stories:
        summary.add("75 mm cover on", ", ".join(map(str, earth_stories)))
    if "Deflection_Check" in results.columns:
        checks = results.drop_duplicates("UniqueName")["Deflection_Check"].astype(str)
        summary.add("Deflection", f"{int(checks.str.startswith('FAIL').sum())} failing of "
                                  f"{len(checks)}")
    for reason, names in failed.groupby(failed).groups.items():
        summary.fail(f"{reason}: {listed(names)}")
    return summary


# =============================================================================
# DEFLECTION ONLY
# =============================================================================
DEFLECTION_REPORT_COLUMNS = ["Story", "UniqueName", "SectProp", "SupportStatus", "Width",
                             "Depth"]


def deflection_table(results: pd.DataFrame) -> pd.DataFrame:
    """One row per beam: the section and its deflection results."""
    from design.beam_deflection import DEFLECTION_COLUMNS

    top = results[results["Face"].astype(str).eq("TOP")] if "Face" in results.columns \
        else results.drop_duplicates("UniqueName")
    keep = [c for c in DEFLECTION_REPORT_COLUMNS if c in top.columns]
    keep += [c for c in DEFLECTION_COLUMNS if c in top.columns]
    table = top[keep].reset_index(drop=True)
    return table.rename(columns=DEFLECTION_COLUMNS)


def write_deflection_xlsx(table: pd.DataFrame, path: str) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    sheet = book.active
    sheet.title = "DEFLECTION"
    sheet.append(table.columns.tolist())
    for row in table.astype(object).where(pd.notna(table), None).itertuples(index=False):
        sheet.append(list(row))
    for index, name in enumerate(table.columns, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DEEBF7")
        sheet.column_dimensions[get_column_letter(index)].width = min(max(len(str(name)) + 2,
                                                                         10), 32)
    check = table.columns.get_loc("Deflection check") + 1 if "Deflection check" in table else 0
    for row in range(2, len(table) + 2):
        if check and str(sheet.cell(row=row, column=check).value) == "FAIL":
            sheet.cell(row=row, column=check).font = Font(color="C00000", bold=True)
    sheet.freeze_panes = "C2"
    book.save(path)
    return path


def run_deflection() -> pd.DataFrame | None:
    """Deflection checks only, with the bars of the last beam design of the model."""
    from design.beam_deflection import DEFLECTION_COLUMNS, add_deflection_columns
    from design.beam_designer_aci318 import ask_deflection_limit
    from etabs_api.workflows.exporter import ETABSDataExporter
    from utilities._gui_helpers import LoadingWindow, select_output_directory, show_warning

    title = "Deflection Check"
    connector = attach_etabs()
    if connector is None:
        return None
    model_path = model_path_of(connector)
    if model_path is None:
        show_warning("Save the ETABS model first: it has no file yet.", title=title)
        return None
    store = DesignStore.load(model_path)
    stem = os.path.splitext(os.path.basename(model_path))[0]
    if store is None or store.beam_results is None or store.beam_results.empty:
        show_warning(f"No beam design for {stem}: the deflection needs the beam bars. "
                     "Run sdt beams first.", title=title)
        return None
    from etabs_api.workflows import model_inputs as mi

    roles = mi.ask_deflection_roles(connector, model_path, title)
    if roles is None:
        return None
    if any(str(name).isnumeric() for name in store.beam_results["UniqueName"].astype(str)):
        connector.include_numeric_members = True  # the beams were designed with their numbers
    divisor = ask_deflection_limit()
    if divisor is None:
        return None
    folder = select_output_directory("Folder for the deflection results")
    if not folder:
        return None
    targets = dcr_targets.Targets.from_saved(store.inputs.get("dcr_targets"))
    with dcr_targets.use(targets), LoadingWindow("Deflection check") as window:
        window.update("Reading the service moments of the deflection combinations from ETABS")
        exporter = ETABSDataExporter(connector)
        exporter.deflection_roles = roles
        beams = list(store.beam_results["UniqueName"].astype(str).unique())
        service = exporter.display_service_loads(beams)
        exporter.display_connectivity_data()
        if service is None:
            show_warning("The model has no deflection combinations or no beam results.",
                         title=title)
            return None
        service = numeric_like(service)
        connectivity = numeric_like(exporter.tables["CONNECTIVITY"])
        results = store.beam_results.drop(columns=[c for c in DEFLECTION_COLUMNS
                                                   if c in store.beam_results.columns])
        checked = add_deflection_columns(results, service, divisor, connectivity,
                                         progress=window.update)
        table = deflection_table(checked)
        window.update("Saving the deflection workbook (.xlsx)")
        path = write_deflection_xlsx(table, os.path.join(folder, f"{stem} - Deflection.xlsx"))
    from utilities.run_summary import RunSummary, listed

    status = table["Deflection check"].astype(str)
    failing = table.loc[status.str.startswith("FAIL"), "UniqueName"].astype(str).tolist() \
        if "UniqueName" in table.columns else []
    summary = RunSummary("sdt deflection", model_path)
    summary.add("Beams checked", len(table))
    summary.add("Limits", f"L/360 live, L/{divisor} after partitions")
    summary.add("Failing", int(status.str.startswith("FAIL").sum()))
    own = [f"{mi.DEFLECTION_ROLES[role]} = {name}" for role, name in roles.items()
           if name != role]
    summary.add("Combinations", "; ".join(own) if own else "the DEF names of the model")
    add_targets_to(summary, targets, (dcr_targets.GIRDER, dcr_targets.BEAM))
    if failing:
        summary.fail("Deflection (ACI 24.2.2): " + listed(failing))
    summary.file("Results", path)
    summary.show(popup=True)
    return table


# =============================================================================
# COLUMNS
# =============================================================================
def run_columns() -> DesignStore | None:
    """Design the columns from the stored beam step and save the outputs."""
    from design.column_designer_aci318 import (
        ask_capacity_check_levels,
        ask_column_design_options,
        ask_inner_tie_style,
        design_columns,
        export_column_cad_drawings,
        export_column_pdf,
        write_column_results_xlsx,
    )
    from utilities._gui_helpers import (
        LoadingWindow,
        select_option,
        select_output_directory,
        show_warning,
    )

    title = "Column Design"
    connector = attach_etabs()
    if connector is None:
        return None
    model_path = model_path_of(connector)
    if model_path is None:
        show_warning("Save the ETABS model first: it has no file yet.", title=title)
        return None
    store = DesignStore.load(model_path)
    stem = os.path.splitext(os.path.basename(model_path))[0]
    if store is None or store.beam_results is None or store.beam_results.empty:
        show_warning(f"No beam design for {stem}. Run sdt beams first.", title=title)
        return None
    if not confirm_pdelta(connector, title):
        return None
    if store.is_stale():
        go_on = select_option(title, f"The model was saved after the beam design of {stem}. "
                              "The stored forces may be out of date.",
                              ["Stop (run sdt beams again)", "Continue with the stored data"])
        if go_on is None or go_on.startswith("Stop"):
            return None
    last = _last()
    bars = _ask_numbers(title, "Column bars and cover.", COLUMN_FIELDS, last)
    if bars is None:
        return None
    answers = ask_column_design_options(normal_cover=bars["cover"])
    if answers is None:
        return None
    smrf = bool(store.inputs.get("smrf", True))
    levels = ask_capacity_check_levels(smrf, last)
    if levels is None:
        return None
    tie_style = ask_inner_tie_style()
    if tie_style is None:
        return None
    targets = ask_dcr_targets(model_path, title, (dcr_targets.COLUMN,))
    if targets is None:
        return None
    folder = select_output_directory("Folder for the column results, calculations and schedule")
    if not folder:
        return None
    _remember({**bars, "check_top_level": levels[0], "check_foundation_level": levels[1]})

    with dcr_targets.use(targets), LoadingWindow("Column design") as window:
        report, groups, joints = design_columns(
            store.tables, store.beam_results, smrf, bars["dmain"], bars["dties"], bars["cover"],
            progress=window.update, continuous_vertical_bars=answers[0],
            bottom_story_cover=answers[1], check_top_level=levels[0],
            check_foundation_level=levels[1])
        store.column_report, store.column_groups, store.joint_results = report, groups, joints
        store.inputs.update({"column_bars": bars, "continuous_bars": answers[0],
                             "bottom_cover": answers[1], "check_top_level": levels[0],
                             "check_foundation_level": levels[1],
                             "column_dcr_targets": targets.to_saved()})
        store.save()
        window.update("Saving 1 of 3: the results workbook (.xlsx)")
        xlsx = write_column_results_xlsx(report, groups,
                                         os.path.join(folder, f"{stem} - Column Design.xlsx"))
        window.update("Saving 2 of 3: the column schedule (.dxf)")
        dxf = export_column_cad_drawings(report, folder, bars["dmain"], bars["cover"], smrf,
                                         tie_style, store.tables["CONNECTIVITY"])
        window.update("Saving 3 of 3: the calculation report (.pdf, LaTeX)")
        pdf = export_column_pdf(report, os.path.join(folder, f"{stem} - Column Calculations.pdf"),
                                smrf, bars["dmain"], bars["dties"], bars["cover"])
    summary = column_summary(report, "sdt columns", model_path, levels[0], levels[1])
    add_targets_to(summary, targets, (dcr_targets.COLUMN,))
    from etabs_api.workflows import model_inputs as mi

    # the forces are those of the beam step: so are the inputs they came from
    mi.Sources(**(store.inputs.get("sources") or {})).add_to(summary)
    walls = mi.wall_count(connector)
    if walls:
        summary.note(f"{walls} wall panels are in the model: walls are not designed.")
    summary.file("Results", xlsx)
    summary.file("Schedule", ", ".join(os.path.basename(p) for p in dxf))
    summary.file("Calculations", pdf or missing_pdf_reason())
    summary.show(os.path.join(folder, f"{stem} - Column Design summary.txt"), popup=True)
    return store


# What a failing column fails in, from its report rows: (label, column, text that fails).
_COLUMN_FAILURES = (
    ("Flexure / axial", "Flexure_Check", "FAIL"),
    ("Flexure / axial", "Axial_Check", "FAIL"),
    ("Slenderness (ACI 6.2.6)", "Slenderness_Check", "FAIL"),
    ("Column shear", "Shear_Check", "FAIL"),
    ("Strong column - weak beam (BCC)", "BCC_Status", "FAIL"),
    ("Joint shear", "Joint_Shear_Status", "FAIL"),
    ("Transverse detailing", "Transverse_Reinforcement_Check", "FAIL"),
    ("SMRF dimensions", "SMRF_Dimension_Check", "FAIL"),
)


def column_summary(report: pd.DataFrame, command: str, model_path: str | None = None,
                   check_top_level: bool = True, check_foundation_level: bool = True):
    """The closing summary of a column design: counts and what fails, by check."""
    from utilities.run_summary import RunSummary, listed

    summary = RunSummary(command, model_path)
    names = report["UniqueName"].astype(str)
    failing = report.loc[report["Column_Design_Status"].astype(str) == "FAIL"]
    failed_names = sorted(set(failing["UniqueName"].astype(str)))
    summary.add("Columns designed", names.nunique())
    summary.add("Combinations", report["Combo"].astype(str).nunique())
    summary.add("Passing", names.nunique() - len(failed_names))
    summary.add("Failing", len(failed_names))
    if "Slenderness_Check" in report.columns:
        slender = report.loc[report["Slenderness_Check"].astype(str).str.contains(
            "slender, moments magnified|FAIL", regex=True), "UniqueName"].astype(str).nunique()
        summary.add("Slender columns", f"{slender} (member effects, ACI 6.6.4.5; sway effects "
                                       "from the ETABS P-delta analysis)")
    if not check_top_level:
        summary.note("BCC and joint shear were not checked at the topmost level (your choice).")
    if not check_foundation_level:
        summary.note("BCC, joint shear and the probable-moment shear Ve were not checked at "
                     "the foundation level (your choice): those columns use the analysis shear.")
    seen: dict[str, set[str]] = {}
    for label, column, text in _COLUMN_FAILURES:
        if column not in failing.columns:
            continue
        hit = failing.loc[failing[column].astype(str).str.upper().str.startswith(text),
                          "UniqueName"].astype(str)
        if len(hit):
            seen.setdefault(label, set()).update(hit)
    for label, members in seen.items():
        summary.fail(f"{label}: {listed(sorted(members))}")
    other = [name for name in failed_names if not any(name in m for m in seen.values())]
    if other:
        summary.fail(f"Other (see the Design Status Reason column): {listed(other)}")
    return summary


if __name__ == "__main__":
    run_beams()
