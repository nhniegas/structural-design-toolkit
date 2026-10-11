"""The column report: one row per column, end and combination, and the results workbook."""

import math
import numpy as np
import pandas as pd
from design.concrete.columns.slenderness import SLENDERNESS_REPORT_FIELDS
from utilities._calc_report import is_blank, number


BCC_EXEMPT_TEXT = ("N/A - ACI 18.7.3.1: the column stops at this joint and "
                   "Pu < 0.1 Ag fc'")


# Capacity-design checks the user chose not to run at a level (sdt columns asks).
TOP_LEVEL_SKIPPED = "Not checked - topmost level (user choice)"


FOUNDATION_LEVEL_SKIPPED = "Not checked - foundation level (user choice)"


CAPACITY_SHEAR_SKIPPED = "Not used - foundation level (user choice): analysis shear"


TOP_CAPACITY_SHEAR_SKIPPED = "Not used - topmost level (user choice): analysis shear"


def _skip_joint_checks(joints: pd.DataFrame, reasons: dict[str, str]) -> pd.DataFrame:
    """Mark the strong column - weak beam and joint shear checks of the joints in
    ``reasons`` (joint -> text) as not checked. Their values are kept out of the
    report and they never ask for more column bars or a larger section."""
    if joints.empty or not reasons or "Joint_Point" not in joints.columns:
        return joints
    joints = joints.copy()
    for column in ("Strong_Column_Check", "Joint_Shear_Check", "Joint_Check_Reason"):
        if column in joints.columns:
            joints[column] = joints[column].astype(object)
    skipped = joints["Joint_Point"].astype(str).map(reasons)
    rows = skipped.notna()
    if not rows.any():
        return joints
    joints.loc[rows, "BCC_Exempt"] = True
    joints.loc[rows, "Column_Beam_Ratio"] = np.nan
    joints.loc[rows, "Joint_Shear_Utilization"] = np.nan
    for column in ("Strong_Column_Check", "Joint_Shear_Check", "Joint_Check_Reason"):
        joints.loc[rows, column] = skipped[rows]
    return joints


# Per-end report fields of the joint checks. X is the column's width direction
# (local 3) and Y its depth direction (local 2): "X" values are for the beams
# framing along X.
BCC_REPORT_FIELDS = [
    "Column_Reinforcement_At_Joint",
    "Beam_Reinforcement_X",
    "Sum_Column_Mn_X_kNm",
    "Sum_Beam_Mn_X_kNm",
    "BCC_Ratio_X",
    "Beam_Reinforcement_Y",
    "Sum_Column_Mn_Y_kNm",
    "Sum_Beam_Mn_Y_kNm",
    "BCC_Ratio_Y",
    "BCC_Status",
]


JOINT_SHEAR_REPORT_FIELDS = [
    "Joint_Shear_Demand_X_kN",
    "Joint_Shear_Capacity_X_kN",
    "Joint_Shear_Utilization_X",
    "Joint_Shear_Demand_Y_kN",
    "Joint_Shear_Capacity_Y_kN",
    "Joint_Shear_Utilization_Y",
    "Joint_Shear_Status",
]


SLENDERNESS_NOT_RUN = "N/A - slenderness not run"


JOINT_OTHER_COMBINATION = ("N/A - the joint is checked in the combinations that govern it "
                           "(see the other rows of this end)")


def _build_consolidated_column_report(
    column_results: pd.DataFrame,
    load_checks: pd.DataFrame,
    shear_checks: pd.DataFrame,
    joint_results: pd.DataFrame,
    column_labels: dict[str, str] | None = None,
    level_elevations: dict[str, float] | None = None,
    joint_na_reason: str = "N/A",
) -> tuple[pd.DataFrame, list[tuple[str, list[str]]]]:
    """Combine column, force, shear, and joint checks into one I/J report.

    Joint checks are reported for each column axis. The values shown for an axis
    belong to one governing case: the lowest strong-column ratio, and the highest
    joint shear utilization. ``joint_na_reason`` is written where an end has no
    joint check at all.
    """
    def summarize_status(values: list[str]) -> str:
        """Apply conservative PASS/FAIL/BLOCKED precedence to paired joint checks."""
        normalized = list(dict.fromkeys(value.strip() for value in values if value.strip()))
        blocked = [value for value in normalized if value.startswith("BLOCKED")]
        if blocked:
            return " / ".join(blocked)
        failed = [value for value in normalized if value.startswith("FAIL")]
        if failed:   # with its reason, when the check gave one
            reasons = [value for value in failed if value != "FAIL"]
            return " / ".join(reasons) if reasons else "FAIL"
        if normalized and all(value.startswith("PASS") for value in normalized):
            return "PASS"
        return " / ".join(normalized) if normalized else "N/A"

    groups: list[tuple[str, list[str]]] = [
        (
            "COLUMN IDENTIFICATION",
            ["Column_Label", "UniqueName", "Story", "Level_Elevation_m", "Combo"],
        ),
        (
            "SECTION / REINFORCEMENT",
            [
                "Section",
                "Shape",
                "Width_mm",
                "Depth_mm",
                "Diameter_mm",
                "Cover_mm",
                "f\'c_MPa",
                "fy_MPa",
                "fyt_MPa",
                "Longitudinal_Bars",
                "Bars_X_Edge",
                "Bars_Y_Edge",
                "Bundle_Layout",
                "Vertical_Bar_Continuity",
                "Bar_Layout_Data",
                "Reinforcement_Ratio",
                "Reinforcement_Ratio_Limit",
            ],
        ),
    ]
    for end in ("I", "J"):
        groups.append(
            (
                f"FLEXURE / AXIAL - {end}",
                [
                    f"Pu_kN_{end}",
                    f"Mu2_kNm_{end}",
                    f"Mu3_kNm_{end}",
                    f"phi_Mn_kNm_{end}",
                    f"Flexure_Utilization_{end}",
                    f"Axial_Check_{end}",
                    f"Flexure_Check_{end}",
                ],
            )
        )
    for end in ("I", "J"):
        groups.append(
            (f"SLENDERNESS - {end}", [f"{name}_{end}" for name in SLENDERNESS_REPORT_FIELDS])
        )
    for end in ("I", "J"):
        groups.append(
            (
                f"SHEAR - {end}",
                [
                    f"Analysis_Vu_kN_{end}",
                    f"Probable_Ve_kN_{end}",
                    f"Design_Vu_kN_{end}",
                    f"Vc_kN_{end}",
                    f"Concrete_Shear_Neglected_{end}",
                    f"phi_Vn_kN_{end}",
                    f"Shear_Utilization_{end}",
                    f"Shear_Check_{end}",
                ],
            )
        )
    for end in ("I", "J"):
        groups.append(
            (
                f"BEAM-COLUMN CAPACITY - {end}",
                [f"{name}_{end}" for name in BCC_REPORT_FIELDS],
            )
        )
    for end in ("I", "J"):
        groups.append(
            (
                f"JOINT SHEAR - {end}",
                [f"{name}_{end}" for name in JOINT_SHEAR_REPORT_FIELDS],
            )
        )
    groups.append(
        (
            "TRANSVERSE DETAILING",
            [
                "Transverse_Type",
                "Transverse_Provision_Summary",
                "Tie_Bar_mm",
                "Transverse_Spacing_mm",
                "Transverse_Legs_X",
                "Transverse_Legs_Y",
                "Confinement_Criteria_Governing",
                "High_Axial_or_High_fc_Check",
                "Confinement_18_7_5_a",
                "Confinement_18_7_5_b",
                "Confinement_18_7_5_c",
                "Confinement_18_7_5_d",
                "Confinement_18_7_5_e",
                "Confinement_18_7_5_f",
                "Kf",
                "Kn",
                "Required_Ash_s_Ratio_X",
                "Provided_Ash_s_Ratio_X",
                "Required_Ash_s_Ratio_Y",
                "Provided_Ash_s_Ratio_Y",
                "Confinement_Check",
                "Transverse_Spacing_Check",
                "Alternating_Support_Check_X",
                "Alternating_Support_Check_Y",
                "Alternating_Support_Added_Legs",
                "Tie_Diameter_Check",
                "SMRF_Dimension_Check",
                "Transverse_Reinforcement_Check",
                "Column_Design_Status",
            ],
        )
    )

    force_map = {
        (str(row["UniqueName"]), str(row["Combo"]), str(row["End"])): row
        for row in load_checks.to_dict("records")
    }
    combos_of: dict[str, set[str]] = {}
    for force_member, combo, _end in force_map:
        combos_of.setdefault(force_member, set()).add(combo)
    # Shear rows of one column, combination and end, reduced to their worst values once.
    shear_groups: dict[tuple[str, str, str], dict] = {}
    if not shear_checks.empty:
        table = shear_checks.assign(
            _member=shear_checks["UniqueName"].astype(str),
            _combo=shear_checks["Combo"].astype(str),
            _end=shear_checks["End"].astype(str),
            _neglected=shear_checks["Concrete_Shear_Strength_Neglected"].eq(True),
            _passes=shear_checks["Shear_Check"].eq("PASS"),
        )
        reduced = table.groupby(["_member", "_combo", "_end"], dropna=False).agg(
            Analysis_Shear_kN=("Analysis_Shear_kN", "max"),
            Capacity_Based_Ve_kN=("Capacity_Based_Ve_kN", "max"),
            Design_Shear_kN=("Design_Shear_kN", "max"),
            Vc_kN=("Vc_kN", "min"),
            neglected=("_neglected", "all"),
            phi_Vn_kN=("phi_Vn_kN", "min"),
            Shear_Utilization=("Shear_Utilization", "max"),
            passes=("_passes", "all"),
            **({"note": ("Capacity_Shear_Note", "max")}
               if "Capacity_Shear_Note" in table.columns else {}),
        )
        shear_groups = {key: values for key, values in zip(
            reduced.index, reduced.to_dict("records"))}
    # Joint rows per column, combination, end and column axis. Entries are
    # "member:end:axis"; an empty axis means the column orientation is unknown.
    joint_groups: dict[tuple[str, str, str, str], list[pd.Series]] = {}
    if not joint_results.empty:
        for joint_row in joint_results.to_dict("records"):
            combo = str(joint_row.get("Load_Combo", ""))
            for entry in str(joint_row.get("Column_End_Members", "")).split(","):
                parts = [part.strip() for part in entry.rsplit(":", 2)]
                if len(parts) == 3 and parts[1] in ("I", "J"):
                    member, end, axis = parts
                elif len(parts) >= 2:
                    member, end, axis = ":".join(parts[:-1]), parts[-1], ""
                else:
                    continue
                joint_groups.setdefault((member, combo, end, axis), []).append(joint_row)
    # Ends that have a joint check in some combination: where another combination
    # has none, that is because only the governing ones are evaluated.
    joint_ends = {(member, end) for member, _combo, end, _axis in joint_groups}

    def number(item, key: str) -> float:
        """Read a numeric joint value; NaN when it is missing or text."""
        try:
            return float(item.get(key))
        except (TypeError, ValueError):
            return np.nan

    def unique_text(items: list[pd.Series], key: str) -> str:
        """Join the distinct text values of one joint field."""
        return " | ".join(
            dict.fromkeys(str(item.get(key)) for item in items if pd.notna(item.get(key)))
        )

    report_rows: list[dict] = []
    for column in column_results.to_dict("records"):
        member = str(column["UniqueName"])
        combos = sorted(combos_of.get(member, ()))
        if not combos:
            combos = [""]
        for combo in combos:
            report_row = {
                "UniqueName": member,
                "Column_Label": (column_labels or {}).get(member, member),
                "Story": column.get("Story"),
                "Level_Elevation_m": (level_elevations or {}).get(member, np.nan),
                "Combo": combo,
                "Section": column.get("SectProp"),
                "Shape": column.get("Shape"),
                "Width_mm": column.get("Width_mm"),
                "Depth_mm": column.get("Depth_mm"),
                "Diameter_mm": column.get("Diameter_mm"),
                "Cover_mm": column.get("Cover_mm"),
                "f'c_MPa": column.get("f'c_MPa"),
                "fy_MPa": column.get("fy_MPa"),
                "fyt_MPa": column.get("fyt_MPa"),
                "Longitudinal_Bars": column.get("Longitudinal_Bars"),
                "Bars_X_Edge": column.get("Bars_X_Edge"),
                "Bars_Y_Edge": column.get("Bars_Y_Edge"),
                "Bundle_Layout": column.get("Longitudinal_Bar_Layout"),
                "Vertical_Bar_Continuity": column.get("Vertical_Bar_Continuity"),
                "Bar_Layout_Data": column.get("Longitudinal_Bar_Coordinates"),
                "Reinforcement_Ratio": column.get("Reinforcement_Ratio"),
                "Reinforcement_Ratio_Limit": column.get(
                    "Reinforcement_Ratio_Limit"
                ),
                "Transverse_Type": column.get("Transverse_Type"),
                "Transverse_Provision_Summary": column.get(
                    "Transverse_Provision_Summary"
                ),
                "Tie_Bar_mm": column.get("Tie_Bar_mm"),
                "Transverse_Spacing_mm": column.get(
                    "Transverse_Spacing_Provided_mm"
                ),
                "Transverse_Legs_X": column.get("Transverse_Legs_X"),
                "Transverse_Legs_Y": column.get("Transverse_Legs_Y"),
                "Confinement_Criteria_Governing": column.get(
                    "Confinement_Criteria_Governing"
                ),
                "High_Axial_or_High_fc_Check": column.get(
                    "High_Axial_or_High_fc_Check"
                ),
                "Confinement_18_7_5_a": column.get(
                    "Confinement_18_7_5_a"
                ),
                "Confinement_18_7_5_b": column.get(
                    "Confinement_18_7_5_b"
                ),
                "Confinement_18_7_5_c": column.get(
                    "Confinement_18_7_5_c"
                ),
                "Confinement_18_7_5_d": column.get(
                    "Confinement_18_7_5_d"
                ),
                "Confinement_18_7_5_e": column.get(
                    "Confinement_18_7_5_e"
                ),
                "Confinement_18_7_5_f": column.get(
                    "Confinement_18_7_5_f"
                ),
                "Kf": column.get("Kf"),
                "Kn": column.get("Kn"),
                "Required_Ash_s_Ratio_X": column.get(
                    "Required_Ash_s_Ratio_X"
                ),
                "Provided_Ash_s_Ratio_X": column.get(
                    "Provided_Ash_s_Ratio_X"
                ),
                "Required_Ash_s_Ratio_Y": column.get(
                    "Required_Ash_s_Ratio_Y"
                ),
                "Provided_Ash_s_Ratio_Y": column.get(
                    "Provided_Ash_s_Ratio_Y"
                ),
                "Confinement_Check": column.get("Confinement_Check"),
                "Transverse_Spacing_Check": column.get(
                    "Transverse_Spacing_Check"
                ),
                "Alternating_Support_Check_X": column.get(
                    "Alternating_Support_Check_X"
                ),
                "Alternating_Support_Check_Y": column.get(
                    "Alternating_Support_Check_Y"
                ),
                "Alternating_Support_Added_Legs": column.get(
                    "Alternating_Support_Added_Legs"
                ),
                "Tie_Diameter_Check": column.get("Tie_Diameter_Check"),
                "SMRF_Dimension_Check": column.get("SMRF_Dimension_Check"),
                "Transverse_Reinforcement_Check": column.get(
                    "Transverse_Reinforcement_Check"
                ),
                "Column_Design_Status": column.get("Design_Status"),
            }
            for end in ("I", "J"):
                force = force_map.get((member, combo, end))
                for key, source in (
                    (f"Pu_kN_{end}", "Pu_kN"),
                    (f"Mu2_kNm_{end}", "Mu2_kNm"),
                    (f"Mu3_kNm_{end}", "Mu3_kNm"),
                    (f"phi_Mn_kNm_{end}", "phi_Mn_kNm"),
                    (f"Flexure_Utilization_{end}", "Flexure_Utilization"),
                    (f"Axial_Check_{end}", "Axial_Check"),
                    (f"Flexure_Check_{end}", "Flexure_Check"),
                ):
                    report_row[key] = force.get(source) if force is not None else "N/A"
                for name in SLENDERNESS_REPORT_FIELDS:
                    report_row[f"{name}_{end}"] = (
                        force.get(name, SLENDERNESS_NOT_RUN) if force is not None else "N/A")

                shear_rows = shear_groups.get((member, combo, end))
                if shear_rows is None:
                    for prefix in (
                        "Analysis_Vu_kN",
                        "Probable_Ve_kN",
                        "Design_Vu_kN",
                        "Vc_kN",
                        "Concrete_Shear_Neglected",
                        "phi_Vn_kN",
                        "Shear_Utilization",
                        "Shear_Check",
                    ):
                        report_row[f"{prefix}_{end}"] = "N/A"
                else:
                    report_row[f"Analysis_Vu_kN_{end}"] = shear_rows["Analysis_Shear_kN"]
                    report_row[f"Probable_Ve_kN_{end}"] = (
                        shear_rows.get("note") or shear_rows["Capacity_Based_Ve_kN"])
                    report_row[f"Design_Vu_kN_{end}"] = shear_rows["Design_Shear_kN"]
                    report_row[f"Vc_kN_{end}"] = shear_rows["Vc_kN"]
                    report_row[f"Concrete_Shear_Neglected_{end}"] = (
                        "YES" if shear_rows["neglected"] else "NO"
                    )
                    report_row[f"phi_Vn_kN_{end}"] = shear_rows["phi_Vn_kN"]
                    report_row[f"Shear_Utilization_{end}"] = shear_rows["Shear_Utilization"]
                    report_row[f"Shear_Check_{end}"] = (
                        "PASS" if shear_rows["passes"] else "FAIL"
                    )

                end_rows: list[pd.Series] = []
                for axis in ("X", "Y"):
                    axis_rows = joint_groups.get(
                        (member, combo, end, axis), []
                    ) + joint_groups.get((member, combo, end, ""), [])
                    end_rows.extend(axis_rows)
                    missing = (
                        unique_text(axis_rows, "Joint_Check_Reason")
                        if axis_rows
                        else f"N/A - no beam frames in along {axis}"
                    )
                    # Strong column - weak beam: the case with the lowest ratio.
                    rated = [
                        item for item in axis_rows
                        if pd.notna(number(item, "Column_Beam_Ratio"))
                    ]
                    if rated:
                        governing = min(
                            rated, key=lambda item: number(item, "Column_Beam_Ratio")
                        )
                        report_row[f"Beam_Reinforcement_{axis}_{end}"] = governing.get(
                            "Framing_Beam_Reinforcement"
                        )
                        report_row[f"Sum_Column_Mn_{axis}_kNm_{end}"] = number(
                            governing, "Sum_Column_Mn_kNm"
                        )
                        report_row[f"Sum_Beam_Mn_{axis}_kNm_{end}"] = number(
                            governing, "Sum_Beam_Mn_kNm"
                        )
                        report_row[f"BCC_Ratio_{axis}_{end}"] = number(
                            governing, "Column_Beam_Ratio"
                        )
                    else:
                        for prefix in (
                            f"Beam_Reinforcement_{axis}",
                            f"Sum_Column_Mn_{axis}_kNm",
                            f"Sum_Beam_Mn_{axis}_kNm",
                            f"BCC_Ratio_{axis}",
                        ):
                            report_row[f"{prefix}_{end}"] = missing
                    # Joint shear: the case with the highest utilization.
                    sheared = [
                        item for item in axis_rows
                        if pd.notna(number(item, "Joint_Shear_Utilization"))
                    ]
                    if sheared:
                        governing = max(
                            sheared,
                            key=lambda item: number(item, "Joint_Shear_Utilization"),
                        )
                        report_row[f"Joint_Shear_Demand_{axis}_kN_{end}"] = number(
                            governing, "Joint_Shear_Demand_kN"
                        )
                        report_row[f"Joint_Shear_Capacity_{axis}_kN_{end}"] = number(
                            governing, "phi_Vn_kN"
                        )
                        report_row[f"Joint_Shear_Utilization_{axis}_{end}"] = number(
                            governing, "Joint_Shear_Utilization"
                        )
                    else:
                        for prefix in (
                            f"Joint_Shear_Demand_{axis}_kN",
                            f"Joint_Shear_Capacity_{axis}_kN",
                            f"Joint_Shear_Utilization_{axis}",
                        ):
                            report_row[f"{prefix}_{end}"] = missing

                if end_rows:
                    report_row[f"Column_Reinforcement_At_Joint_{end}"] = unique_text(
                        end_rows, "Column_Reinforcement_At_Joint"
                    )
                    report_row[f"BCC_Status_{end}"] = summarize_status(
                        [str(item.get("Strong_Column_Check", "N/A")) for item in end_rows]
                    )
                    report_row[f"Joint_Shear_Status_{end}"] = summarize_status(
                        [str(item.get("Joint_Shear_Check", "N/A")) for item in end_rows]
                    )
                else:
                    reason = (JOINT_OTHER_COMBINATION if (member, end) in joint_ends
                              else joint_na_reason)
                    for name in BCC_REPORT_FIELDS + JOINT_SHEAR_REPORT_FIELDS:
                        report_row[f"{name}_{end}"] = reason
            report_rows.append(report_row)

    report = pd.DataFrame(
        report_rows,
        columns=[column for _, column_names in groups for column in column_names],
    )
    report["_story_sort"] = report["Story"].astype(str)
    report = (
        report.sort_values(
            ["Column_Label", "Level_Elevation_m", "_story_sort", "UniqueName", "Combo"],
            ascending=[True, False, True, True, True],
            kind="stable",
        )
        .drop(columns="_story_sort")
        .reset_index(drop=True)
    )
    return _expand_column_report_hierarchy(report)


def _expand_column_report_hierarchy(
    wide_report: pd.DataFrame,
) -> tuple[pd.DataFrame, list[tuple[str, list[str]]]]:
    """Expand I/J columns into hierarchical label, member, end, and combo rows."""
    shared_columns = [
        "Section",
        "Shape",
        "Width_mm",
        "Depth_mm",
        "Diameter_mm",
        "Cover_mm",
        "f\'c_MPa",
        "fy_MPa",
        "fyt_MPa",
        "Longitudinal_Bars",
        "Bars_X_Edge",
        "Bars_Y_Edge",
        "Bundle_Layout",
        "Vertical_Bar_Continuity",
        "Bar_Layout_Data",
        "Reinforcement_Ratio",
        "Reinforcement_Ratio_Limit",
    ]
    end_groups = {
        "FLEXURE / AXIAL": [
            "Pu_kN",
            "Mu2_kNm",
            "Mu3_kNm",
            "phi_Mn_kNm",
            "Flexure_Utilization",
            "Axial_Check",
            "Flexure_Check",
        ],
        "SLENDERNESS": list(SLENDERNESS_REPORT_FIELDS),
        "COLUMN SHEAR": [
            "Analysis_Vu_kN",
            "Probable_Ve_kN",
            "Design_Vu_kN",
            "Vc_kN",
            "Concrete_Shear_Neglected",
            "phi_Vn_kN",
            "Shear_Utilization",
            "Shear_Check",
        ],
        "BEAM-COLUMN CAPACITY": list(BCC_REPORT_FIELDS),
        "JOINT SHEAR": list(JOINT_SHEAR_REPORT_FIELDS),
    }
    detailing_columns = [
        "Transverse_Type",
        "Transverse_Provision_Summary",
        "Tie_Bar_mm",
        "Transverse_Spacing_mm",
        "Transverse_Legs_X",
        "Transverse_Legs_Y",
        "Confinement_Criteria_Governing",
        "High_Axial_or_High_fc_Check",
        "Confinement_18_7_5_a",
        "Confinement_18_7_5_b",
        "Confinement_18_7_5_c",
        "Confinement_18_7_5_d",
        "Confinement_18_7_5_e",
        "Confinement_18_7_5_f",
        "Kf",
        "Kn",
        "Required_Ash_s_Ratio_X",
        "Provided_Ash_s_Ratio_X",
        "Required_Ash_s_Ratio_Y",
        "Provided_Ash_s_Ratio_Y",
        "Confinement_Check",
        "Transverse_Spacing_Check",
        "Alternating_Support_Check_X",
        "Alternating_Support_Check_Y",
        "Alternating_Support_Added_Legs",
        "Tie_Diameter_Check",
        "SMRF_Dimension_Check",
        "Transverse_Reinforcement_Check",
    ]
    identity_columns = [
        "Column_Label",
        "UniqueName",
        "Story",
        "End",
        "Combo",
    ]
    report_groups: list[tuple[str, list[str]]] = [
        ("COLUMN LABEL / LEVEL", identity_columns),
        ("SECTION / LONGITUDINAL REINFORCEMENT", shared_columns),
    ]
    for group_name, names in end_groups.items():
        report_groups.append((group_name, names))
    report_groups.extend(
        [
            ("TRANSVERSE REINFORCEMENT DETAILING", detailing_columns),
            ("DESIGN STATUS", ["Column_Design_Status", "Design_Status_Reason"]),
        ]
    )

    long_rows: list[dict] = []
    for wide_row in wide_report.to_dict("records"):
        for end in ("I", "J"):
            result = {
                key: wide_row.get(key)
                for key in ("Column_Label", "UniqueName", "Story", "Combo")
            }
            result["End"] = end
            result.update(
                {key: wide_row.get(key) for key in shared_columns + detailing_columns}
            )
            for names in end_groups.values():
                for name in names:
                    result[name] = wide_row.get(f"{name}_{end}")
            statuses = [
                str(result.get(name, ""))
                for name in (
                    "Flexure_Check",
                    "Axial_Check",
                    "Slenderness_Check",
                    "Shear_Check",
                    "BCC_Status",
                    "Joint_Shear_Status",
                    "Confinement_Check",
                    "Transverse_Spacing_Check",
                    "Alternating_Support_Check_X",
                    "Alternating_Support_Check_Y",
                    "Tie_Diameter_Check",
                    "SMRF_Dimension_Check",
                    "Transverse_Reinforcement_Check",
                )
            ]
            statuses.append(
                str(wide_row.get("Column_Design_Status", "")).strip()
            )
            statuses = [status.strip() for status in statuses if status.strip()]
            normalized_statuses = [status.upper() for status in statuses]
            blocked_reasons = list(
                dict.fromkeys(
                    status
                    for status, normalized in zip(statuses, normalized_statuses)
                    if normalized.startswith("BLOCKED")
                )
            )
            if any(status.startswith("FAIL") for status in normalized_statuses):
                result["Column_Design_Status"] = "FAIL"
                result["Design_Status_Reason"] = "; ".join(
                    status
                    for status, normalized in zip(statuses, normalized_statuses)
                    if normalized.startswith("FAIL")
                )
            elif any(status.startswith("ERROR") for status in normalized_statuses):
                result["Column_Design_Status"] = "ERROR"
                result["Design_Status_Reason"] = "; ".join(
                    status
                    for status, normalized in zip(statuses, normalized_statuses)
                    if normalized.startswith("ERROR")
                )
            elif any(status.startswith("BLOCKED") for status in normalized_statuses):
                result["Column_Design_Status"] = "BLOCKED"
                result["Design_Status_Reason"] = "; ".join(blocked_reasons)
            else:
                result["Column_Design_Status"] = "PASS"
                result["Design_Status_Reason"] = "All applicable checks passed."
            long_rows.append(result)

    long_report = pd.DataFrame(
        long_rows,
        columns=[
            column
            for _, column_names in report_groups
            for column in column_names
        ],
    )
    long_report["_end_order"] = long_report["End"].map({"J": 0, "I": 1})
    elevations = pd.to_numeric(wide_report["Level_Elevation_m"], errors="coerce")
    long_report["_elevation"] = np.repeat(elevations.fillna(-math.inf).to_numpy(float), 2)
    long_report = (
        long_report.assign(
            _label_sort=long_report["Column_Label"].astype(str),
            _story_sort=long_report["Story"].astype(str),
            _member_sort=long_report["UniqueName"].astype(str),
        )
        .sort_values(
            [
                "_label_sort",
                "_elevation",
                "_story_sort",
                "_member_sort",
                "_end_order",
                "Combo",
            ],
            ascending=[True, False, True, True, True, True],
            kind="stable",
        )
        .drop(columns=["_end_order", "_elevation", "_label_sort", "_story_sort", "_member_sort"])
        .reset_index(drop=True)
    )
    return long_report, report_groups


# Worksheet header of every report field. The sheet is read back through the
# same names (DXF schedule, calculation report).
COLUMN_REPORT_LABELS = {
    "Column_Label": "Column Label",
    "UniqueName": "Unique Name",
    "Story": "Story",
    "End": "End",
    "Combo": "Load Combination",
    "Section": "Section",
    "Shape": "Shape",
    "Width_mm": "Width (mm)",
    "Depth_mm": "Depth (mm)",
    "Diameter_mm": "Diameter (mm)",
    "Cover_mm": "Concrete Cover (mm)",
    "f'c_MPa": "f′c (MPa)",
    "fy_MPa": "fᵧ (MPa)",
    "fyt_MPa": "fᵧₜ (MPa)",
    "Longitudinal_Bars": "Longitudinal Bars",
    "Bars_X_Edge": "Bars on X Edge",
    "Bars_Y_Edge": "Bars on Y Edge",
    "Bundle_Layout": "Bundle Layout",
    "Vertical_Bar_Continuity": "Vertical Bar Continuity",
    "Bar_Layout_Data": "Bar Layout Data (x, y, n)",
    "Beam_Reinforcement_X": "Beam Bars at Joint, X",
    "Sum_Column_Mn_X_kNm": "ΣMₙ,Column X (kN·m)",
    "Sum_Beam_Mn_X_kNm": "ΣMₙ,Beam X (kN·m)",
    "BCC_Ratio_X": "ΣMₙ,Column / ΣMₙ,Beam, X",
    "Beam_Reinforcement_Y": "Beam Bars at Joint, Y",
    "Sum_Column_Mn_Y_kNm": "ΣMₙ,Column Y (kN·m)",
    "Sum_Beam_Mn_Y_kNm": "ΣMₙ,Beam Y (kN·m)",
    "BCC_Ratio_Y": "ΣMₙ,Column / ΣMₙ,Beam, Y",
    "Joint_Shear_Demand_X_kN": "Joint Shear Demand X (kN)",
    "Joint_Shear_Capacity_X_kN": "ϕ Joint Shear Capacity X (kN)",
    "Joint_Shear_Utilization_X": "Joint Shear Utilization X",
    "Joint_Shear_Demand_Y_kN": "Joint Shear Demand Y (kN)",
    "Joint_Shear_Capacity_Y_kN": "ϕ Joint Shear Capacity Y (kN)",
    "Joint_Shear_Utilization_Y": "Joint Shear Utilization Y",
    "Alternating_Support_Check_X": "Bar Support Check, X Edge",
    "Alternating_Support_Check_Y": "Bar Support Check, Y Edge",
    "Reinforcement_Ratio": "ρ Longitudinal",
    "Reinforcement_Ratio_Limit": "ρ Limit",
    "Pu_kN": "Pᵤ (kN)",
    "Mu2_kNm": "Mᵤ₂ design (kN·m)",
    "Mu3_kNm": "Mᵤ₃ design (kN·m)",
    "phi_Mn_kNm": "ϕMₙ (kN·m)",
    "Flexure_Utilization": "Flexure Utilization",
    "Axial_Check": "Axial Check",
    "Flexure_Check": "Flexure Check",
    "Slender_lu3_mm": "lᵤ, about 3 (mm)",
    "Slender_k3": "k, about 3",
    "Slender_ratio3": "k lᵤ / r, about 3",
    "Slender_limit3": "Limit 34 + 12 M₁/M₂ ≤ 40, about 3",
    "Slender_Cm3": "Cₘ, about 3",
    "Slender_Pc3_kN": "P꜀, about 3 (kN)",
    "Slender_delta3": "δₙₛ, about 3",
    "Slender_Mmin3_kNm": "M₂,min, about 3 (kN·m)",
    "Mu3_analysis_kNm": "Mᵤ₃ from analysis (kN·m)",
    "Slender_lu2_mm": "lᵤ, about 2 (mm)",
    "Slender_k2": "k, about 2",
    "Slender_ratio2": "k lᵤ / r, about 2",
    "Slender_limit2": "Limit 34 + 12 M₁/M₂ ≤ 40, about 2",
    "Slender_Cm2": "Cₘ, about 2",
    "Slender_Pc2_kN": "P꜀, about 2 (kN)",
    "Slender_delta2": "δₙₛ, about 2",
    "Slender_Mmin2_kNm": "M₂,min, about 2 (kN·m)",
    "Mu2_analysis_kNm": "Mᵤ₂ from analysis (kN·m)",
    "Slender_beta_dns": "βdns",
    "Slender_beta_basis": "βdns basis",
    "Slender_base": "Base of the column (ψ at the support)",
    "Slenderness_Check": "Slenderness Check",
    "Analysis_Vu_kN": "Analysis Vᵤ (kN)",
    "Probable_Ve_kN": "Probable Vₑ (kN)",
    "Design_Vu_kN": "Design Vᵤ (kN)",
    "Vc_kN": "V꜀ (kN)",
    "Concrete_Shear_Neglected": "Concrete Shear Neglected",
    "phi_Vn_kN": "ϕVₙ (kN)",
    "Shear_Utilization": "Shear Utilization",
    "Shear_Check": "Shear Check",
    "Column_Reinforcement_At_Joint": "Column Bars Contributing at Joint",
    "Beam_Reinforcement": "Beam Bars at Joint",
    "BCC_Ratio": "ΣMₙ,Column / ΣMₙ,Beam",
    "BCC_Status": "BCC Check",
    "Joint_Shear_Demand_kN": "Joint Shear Demand (kN)",
    "Joint_Shear_Capacity_kN": "ϕ Joint Shear Capacity (kN)",
    "Joint_Shear_Utilization": "Joint Shear Utilization",
    "Joint_Shear_Status": "Joint Shear Check",
    "Transverse_Type": "Hoop / Spiral Type",
    "Transverse_Provision_Summary": "Transverse Reinforcement Provision",
    "Tie_Bar_mm": "Tie Bar Diameter (mm)",
    "Transverse_Spacing_mm": "Tie / Spiral Spacing (mm)",
    "Transverse_Legs_X": "Tie Legs along X Edge",
    "Transverse_Legs_Y": "Tie Legs along Y Edge",
    "Confinement_Criteria_Governing": "Governing ACI 18.7.5 Expression",
    "High_Axial_or_High_fc_Check": "High Pᵤ / f′c Condition",
    "Confinement_18_7_5_a": "ρₛ,req from 18.7.5(a)",
    "Confinement_18_7_5_b": "ρₛ,req from 18.7.5(b)",
    "Confinement_18_7_5_c": "ρₛ,req from 18.7.5(c)",
    "Confinement_18_7_5_d": "ρₛ,req from 18.7.5(d)",
    "Confinement_18_7_5_e": "ρₛ,req from 18.7.5(e)",
    "Confinement_18_7_5_f": "ρₛ,req from 18.7.5(f)",
    "Kf": "kᶠ",
    "Kn": "kₙ",
    "Required_Ash_s_Ratio_X": "Required Aₛₕ/(s b꜀), X",
    "Provided_Ash_s_Ratio_X": "Provided Aₛₕ/(s b꜀), X",
    "Required_Ash_s_Ratio_Y": "Required Aₛₕ/(s b꜀), Y",
    "Provided_Ash_s_Ratio_Y": "Provided Aₛₕ/(s b꜀), Y",
    "Confinement_Check": "Confinement Ratio Check",
    "Transverse_Spacing_Check": "Tie Spacing Check",
    "Alternating_Support_Check": "Longitudinal Bar Support Check",
    "Alternating_Support_Added_Legs": "Added Hoop Legs",
    "Tie_Diameter_Check": "Tie Diameter Check",
    "SMRF_Dimension_Check": "SMRF Column Dimension Check",
    "Transverse_Reinforcement_Check": "Transverse Reinforcement Check",
    "Column_Design_Status": "Overall Design Status",
    "Design_Status_Reason": "Design Status Reason",
}


_COLUMN_GROUP_FILLS = {
    "COLUMN LABEL / LEVEL": "BDD7EE",
    "SECTION / LONGITUDINAL REINFORCEMENT": "E2EFDA",
    "FLEXURE / AXIAL": "E2EFDA",
    "SLENDERNESS": "EDEDED",
    "COLUMN SHEAR": "FFF2CC",
    "BEAM-COLUMN CAPACITY": "FCE4D6",
    "JOINT SHEAR": "FCE4D6",
    "TRANSVERSE REINFORCEMENT DETAILING": "DEEBF7",
    "DESIGN STATUS": "D9D2E9",
}


# Dimensionless ratios far below 1: two decimals would round them to 0.01 or 0.00.
_COLUMN_FOUR_DECIMALS = {
    "Reinforcement_Ratio", "Reinforcement_Ratio_Limit", "Confinement_18_7_5_a",
    "Confinement_18_7_5_b", "Confinement_18_7_5_c", "Confinement_18_7_5_d",
    "Confinement_18_7_5_e", "Confinement_18_7_5_f", "Required_Ash_s_Ratio_X",
    "Provided_Ash_s_Ratio_X", "Required_Ash_s_Ratio_Y", "Provided_Ash_s_Ratio_Y",
}


_LOW_AXIAL = "N/A - Pu <= 0.3 Ag f'c and f'c <= 70 MPa"


def column_blank_reason(name: str, row: dict) -> str:
    """Why a cell of the column report has no value."""
    circular = str(row.get("Shape", "")).strip().casefold() == "circular"
    if str(row.get("Confinement_Check", "")).startswith("N/A - non-SMRF") and (
            name.startswith(("Confinement_18_7_5", "Required_Ash", "Provided_Ash"))
            or name in ("Kf", "Kn")):
        return "N/A - seismic design is off"
    if name in ("Width_mm", "Depth_mm", "Bars_X_Edge", "Bars_Y_Edge"):
        return "N/A - circular column" if circular else "N/A"
    if name == "Diameter_mm":
        return "N/A - rectangular column"
    if name in ("Confinement_18_7_5_a", "Confinement_18_7_5_b", "Kn"):
        return "N/A - spiral column: 18.7.5(d) to (f) apply" if circular else "N/A"
    if name == "Confinement_18_7_5_c":
        return "N/A - spiral column: 18.7.5(d) to (f) apply" if circular else _LOW_AXIAL
    if name in ("Confinement_18_7_5_d", "Confinement_18_7_5_e"):
        return "N/A" if circular else "N/A - rectangular hoops: 18.7.5(a) to (c) apply"
    if name == "Confinement_18_7_5_f":
        return _LOW_AXIAL if circular else "N/A - rectangular hoops: 18.7.5(a) to (c) apply"
    if name == "Kf":
        return _LOW_AXIAL
    return "N/A"


def column_report_display(report: pd.DataFrame) -> pd.DataFrame:
    """The report with its readable headers and ends (Bottom (I), Top (J))."""
    clean = report.rename(columns={
        name: COLUMN_REPORT_LABELS.get(name, name.replace("_", " ")) for name in report.columns
    })
    if "End" in clean.columns:
        clean["End"] = clean["End"].map(
            lambda end: {"I": "Bottom (I)", "J": "Top (J)"}.get(end, end))
    return clean


def write_column_results_xlsx(report: pd.DataFrame, groups: list[tuple[str, list[str]]],
                              path: str) -> str:
    """The column design report as a formatted Excel file (no Excel needed)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    ordered = [name for _, names in groups for name in names if name in report.columns]
    clean = column_report_display(report[ordered])
    book = Workbook()
    sheet = book.active
    sheet.title = "COLUMN DESIGN"
    sheet["A1"] = "COLUMN DESIGN - CONSOLIDATED CHECKS"
    sheet["A1"].font = Font(size=14, bold=True)
    thick, thin = Side(style="medium"), Side(style="thin")
    column = 1
    for label, names in groups:
        present = [n for n in names if n in report.columns]
        if not present:
            continue
        last = column + len(present) - 1
        sheet.merge_cells(start_row=2, start_column=column, end_row=2, end_column=last)
        cell = sheet.cell(row=2, column=column, value=label)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        for index in range(column, last + 1):
            sheet.cell(row=2, column=index).fill = PatternFill(
                "solid", fgColor=_COLUMN_GROUP_FILLS.get(label, "BDD7EE"))
            sheet.cell(row=2, column=index).border = Border(top=thick, bottom=thick)
        column = last + 1
    for index, name in enumerate(clean.columns, start=1):
        cell = sheet.cell(row=3, column=index, value=name)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="F2F2F2")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=thin)
        sheet.column_dimensions[get_column_letter(index)].width = min(
            max(12, len(str(name)) * 0.9), 30)
    from utilities._xlsx_values import cell_value, is_blank

    # Rounded numbers, and a stated reason wherever a value does not apply.
    values = []
    for record, shown in zip(report[ordered].to_dict("records"),
                             clean.astype(object).values.tolist()):
        values.append([
            cell_value(value, 4 if name in _COLUMN_FOUR_DECIMALS else 2,
                       column_blank_reason(name, record) if is_blank(value) else "")
            for name, value in zip(ordered, shown)
        ])
    names = clean["Unique Name"].tolist() if "Unique Name" in clean.columns else []
    grey = PatternFill("solid", fgColor="F2F2F2")
    shade, previous = False, None
    for offset, row in enumerate(values):
        excel_row = 4 + offset
        if names and names[offset] != previous:
            shade, previous = not shade, names[offset]
            if offset:
                for index in range(1, len(row) + 1):
                    cell = sheet.cell(row=excel_row - 1, column=index)
                    cell.border = Border(bottom=thick)
        for index, value in enumerate(row, start=1):
            cell = sheet.cell(row=excel_row, column=index, value=value)
            if shade:
                cell.fill = grey
            text = str(value).upper() if value is not None else ""
            if text.startswith("FAIL"):
                cell.font = Font(color="C00000", bold=True)
    sheet.freeze_panes = "E4"
    book.save(path)
    return path
