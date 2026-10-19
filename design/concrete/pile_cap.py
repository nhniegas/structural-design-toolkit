"""Pile cap design to NSCP 2015 (``sdt pilecap``).

A reinforced concrete pile cap under one column, standalone like ``sdt
composite``: the inputs are typed in dialogs, the loads are typed or read from
an Excel or CSV file, and the result is printed and/or written as a PDF
calculation report.

NSCP 2015 governs (Section 413.4.2 for pile caps, with Sections 407, 409,
421, 422, 423, 424, 425 and 418.13). Where NSCP leaves a point open, the
program follows the documents it names in the report:

* CRSI, Design Guide for Pile Caps, 1st ed., 2015 (written to ACI 318-14, the
  basis of NSCP 2015 Chapter 4): rigid cap and pinned piles (3.3, 6.2), pile
  placement tolerance (5.2), minimum steel (4.3), uniform short bars (4.2), the
  special shear checks near the column (5.3 and 5.4, limit states 4 to 6 and
  P1 to P4), development to the pile centre for deep caps (5.3).
* FHWA GEC 12, Design and Construction of Driven Pile Foundations, 2016,
  quoting AASHTO LRFD 10.7.1.2: pile spacing, edge distance, embedment
  (7.10.1) and pile head fixity (7.3.7.1).

Units inside: N, mm, MPa. Inputs and outputs: kN, kN-m, mm, MPa, kPa.
Signs: P is compression positive. Mx increases the reaction of the piles on
+Y, My that of the piles on +X. Vx and Vy act at the top of the cap along +X
and +Y.
"""

from __future__ import annotations

import csv
import math
import os
import re
import sys
from dataclasses import dataclass, field

if __package__ in (None, ""):
    sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))

from utilities.design_cli import Field, run_design  # noqa: E402

# --------------------------------------------------------------------------- #
# Code values (each with its source)
# --------------------------------------------------------------------------- #
PHI_FLEXURE = 0.90          # NSCP Table 421.2.1(a), tension controlled (421.2.2)
PHI_SHEAR = 0.75            # NSCP Table 421.2.1(b)
PHI_STM = 0.75              # NSCP Table 421.2.1(g)
PHI_BEARING = 0.65          # NSCP Table 421.2.1(d)
MIN_DOWEL_RATIO = 0.005     # NSCP 416.3.4.1, of the column gross area
LAMBDA = 1.0                # normal-weight concrete, NSCP 419.2.4
SQRT_FC_MAX = 8.3           # NSCP 422.5.3.1 and 422.6.3.1, MPa
MIN_EFFECTIVE_DEPTH = 300   # NSCP 413.4.2.1, mm
COVER_EARTH = 75            # NSCP Table 420.6.1.3.1, cast against earth, mm
BETA_S_PILE_CAP = 0.60      # NSCP 413.4.2.4 with Table 423.4.3(c)
BETA_N = {"CCC": 1.0, "CCT": 0.80, "CTT": 0.60}   # NSCP Table 423.9.2
MIN_STRUT_ANGLE = 25.0      # NSCP 423.2.7, degrees
LF_FIRM, LF_SOFT = 1500.0, 3000.0   # NSCP 306.6, depth to fixity below ground, mm
MAX_DEPTH = 4000.0          # the least-depth search stops here, mm
DEPTH_STEP = 50.0           # the least-depth search step, mm
SPACING_STEP = 25.0         # bar spacings are rounded down to this, mm
AGGREGATE = 25.0            # nominal maximum aggregate size for bar clear spacing, mm
SIDE_COVER = 75.0           # cover from the cap side to the bar ends, mm
LATERAL_WORDS = re.compile(r"(^|[^A-Z])(E|EQ|EL|EM|EX|EY|RSA|RS|W|WX|WY|WL|WIND|EARTHQUAKE|"
                           r"SEISMIC)([^A-Z]|$)")
SHORT_TON_KN = 8.896        # CRSI tabulates pile loads in (short) tons


# --------------------------------------------------------------------------- #
# Inputs
# --------------------------------------------------------------------------- #
LOADS_EXAMPLE = "D+L, S, 1600, 0, 0, 0, 0; 1.2D+1.6L, F, 2200, 0, 0, 0, 0"

GROUP_PILES = [
    Field("column_shape", "Column shape", "Rectangular", "choice", ("Rectangular", "Circular")),
    Field("cx", "Column size along X, or diameter (mm)", 600.0),
    Field("cy", "Column size along Y (mm)", 600.0),
    Field("col_bar", "Column main bar diameter (mm), for its development into the cap; "
          "blank: not checked", 25.0, optional=True),
    Field("layout", "Pile layout", "Standard", "choice", ("Standard", "Coordinates")),
    Field("n_piles", "Number of piles (2 to 10, or 12)", 4.0),
    Field("spacing", "Pile spacing, centre to centre (mm)", 1200.0),
    Field("edge", "Edge distance, pile centre to cap edge (mm)", 500.0),
    Field("coords", "Pile coordinates x,y (mm, from the column centre), one pile after "
          "another with ; between, e.g. -600,-600; 600,-600; 0,600", "", "text", optional=True),
    Field("cap_lx", "Cap length along X (mm)", None, optional=True),
    Field("cap_ly", "Cap length along Y (mm)", None, optional=True),
    Field("dp", "Pile diameter or side (mm)", 400.0),
    Field("pile_shape", "Pile shape", "Round", "choice", ("Round", "Square")),
    Field("embed", "Pile embedment into the cap (mm)", 100.0),
    Field("tolerance", "Pile placement tolerance (mm)", 75.0),
]
GROUP_CAPACITY = [
    Field("capacity_method", "Pile capacity method", "Allowable", "choice",
          ("Allowable", "Allowable + 1/3 for wind and earthquake", "Factored resistance")),
    Field("q_comp", "Pile capacity in compression (kN)", 1000.0),
    Field("q_tens", "Pile capacity in tension (kN)", 300.0),
    Field("q_lat", "Pile capacity, lateral (kN)", 100.0),
    Field("head", "Pile heads", "Pinned", "choice", ("Pinned", "Fixed")),
    Field("ground", "Ground at the pile heads", "Firm", "choice", ("Firm", "Soft")),
    Field("lf", "Depth to fixity Lf (m); blank: 1.5 firm, 3.0 soft (NSCP 306.6)", None,
          optional=True),
]
GROUP_CAP = [
    Field("h", "Cap thickness h (mm)", 900.0),
    Field("find_depth", "Find the least thickness that passes", "Yes", "choice", ("Yes", "No")),
    Field("fc", "Concrete strength f'c (MPa)", 28.0),
    Field("fy", "Bar yield strength fy (MPa)", 414.0),
    Field("bottom_bar", "Bottom bar diameter (mm)", 25.0),
    Field("bottom_spacing", "Bottom bar spacing (mm); blank: designed", None, optional=True),
    Field("top_bar", "Top bar diameter (mm)", 16.0),
    Field("top_spacing", "Top bar spacing (mm); blank: designed", None, optional=True),
    Field("clear_above_pile", "Clear cover from the pile tops to the bottom bars (mm)", 75.0),
    Field("top_cover", "Cover to the top bars (mm)", 75.0),
    Field("gamma_c", "Concrete unit weight (kN/m3)", 24.0),
    Field("soil_h", "Soil depth over the cap (m)", 0.5),
    Field("gamma_s", "Soil unit weight (kN/m3)", 18.0),
    Field("surcharge", "Surcharge on the ground (kPa)", 0.0),
    Field("flexure_section", "Critical section for moment", "Column face",
          "choice", ("Column face", "c/4 from the column centre")),
    Field("deep_method", "Shear of a deep cap (piles near the column)", "Strut-and-tie",
          "choice", ("Strut-and-tie", "CRSI special checks")),
]
GROUP_LOADS = [
    Field("load_source", "Column loads from", "Typed", "choice", ("Typed", "File")),
    Field("loads", "Loads, one combination per line: name, S (service) or F (factored), "
          "P, Mx, My, Vx, Vy (kN, kN-m), ending with ;", LOADS_EXAMPLE, "text", optional=True),
    Field("loads_file", "Loads file (.xlsx or .csv): columns Name, Type, P, Mx, My, Vx, Vy",
          "", "text", optional=True),
]
# an input shown in the window only when a choice has one of these values
SHOWN_WHEN = {
    "cy": ("column_shape", ("Rectangular",)),
    "n_piles": ("layout", ("Standard",)), "spacing": ("layout", ("Standard",)),
    "edge": ("layout", ("Standard",)),
    "coords": ("layout", ("Coordinates",)), "cap_lx": ("layout", ("Coordinates",)),
    "cap_ly": ("layout", ("Coordinates",)),
    "ground": ("head", ("Fixed",)), "lf": ("head", ("Fixed",)),
    "loads": ("load_source", ("Typed",)), "loads_file": ("load_source", ("File",)),
}
INPUT_GROUPS = [
    ("Pile Cap: column and piles", GROUP_PILES),
    ("Pile Cap: pile capacity", GROUP_CAPACITY),
    ("Pile Cap: cap and reinforcement", GROUP_CAP),
    ("Pile Cap: column loads", GROUP_LOADS),
]
INPUTS = [f for _, group in INPUT_GROUPS for f in group]


# --------------------------------------------------------------------------- #
# Data
# --------------------------------------------------------------------------- #
@dataclass
class Combo:
    """One load combination at the column base (kN, kN-m)."""

    name: str
    factored: bool
    P: float
    Mx: float = 0.0
    My: float = 0.0
    Vx: float = 0.0
    Vy: float = 0.0
    lateral: bool | None = None      # has wind or earthquake (for the 1/3 increase)
    df: float | None = None           # factor on the cap and soil weight
    lf: float | None = None           # factor on the surcharge

    def has_lateral(self) -> bool:
        if self.lateral is not None:
            return self.lateral
        return bool(LATERAL_WORDS.search(self.name.upper()))


@dataclass
class Case:
    """A combination with the factors used on the cap's own weight."""

    combo: Combo
    df: float
    lf: float

    @property
    def label(self) -> str:
        return f"{self.combo.name} (cap weight x{self.df:g})"


@dataclass
class Check:
    """One check: demand, capacity, ratio, and where it comes from."""

    group: str
    name: str
    demand: float | None
    capacity: float | None
    unit: str
    clause: str
    case: str = ""
    note: str = ""
    governs: bool = True       # False: shown for information only
    ratio_override: float | None = None

    @property
    def ratio(self) -> float | None:
        if self.ratio_override is not None:
            return self.ratio_override
        if self.demand is None or self.capacity is None:
            return None
        if self.capacity <= 0:
            return math.inf if self.demand > 0 else 0.0
        return self.demand / self.capacity

    @property
    def passed(self) -> bool:
        return self.ratio is None or self.ratio <= 1.0 + 1e-9


@dataclass
class Bars:
    direction: str             # "X" or "Y": the bars run along this axis
    face: str                  # "Bottom" or "Top"
    diameter: float
    spacing: float
    count: int
    As_req: float              # mm2 over the full width
    As_prov: float
    As_flexure: float
    As_min: float
    As_tie: float
    d: float
    width: float
    anchorage: str             # "Straight", "Hooked" or "NOT DEVELOPED"
    ld: float
    ldh: float
    available: float
    typed: bool
    Mu: float                  # kN-m over the full width
    note: str = ""


@dataclass
class PileCap:
    """Inputs (in mm, MPa, kN) and results of one pile cap."""

    values: dict
    piles: list[tuple[float, float]]
    Lx: float
    Ly: float
    h: float
    combos: list[Combo]
    checks: list[Check] = field(default_factory=list)
    bars: list[Bars] = field(default_factory=list)
    reactions: dict = field(default_factory=dict)      # case label -> list of R (kN)
    lateral_per_pile: dict = field(default_factory=dict)
    geometry: dict = field(default_factory=dict)
    notes: list[str] = field(default_factory=list)
    deep: bool = False
    searched: list[tuple[float, bool, str]] = field(default_factory=list)
    h_typed: float = 0.0

    # ---- convenience
    @property
    def fc(self) -> float:
        return self.values["fc"]

    @property
    def fy(self) -> float:
        return self.values["fy"]

    @property
    def dp(self) -> float:
        return self.values["dp"]

    @property
    def n(self) -> int:
        return len(self.piles)

    def failing(self) -> list[Check]:
        return [c for c in self.checks if c.governs and not c.passed]

    @property
    def passed(self) -> bool:
        return not self.failing()


# --------------------------------------------------------------------------- #
# Pile layouts
# --------------------------------------------------------------------------- #
def standard_layout(n: int, s: float) -> list[tuple[float, float]]:
    """Pile centres of a standard group of ``n`` piles at spacing ``s`` (mm),
    centred on the column. Similar to the CRSI arrangements (Fig. 4.4)."""
    h3 = s * math.sqrt(3) / 2
    grid = lambda nx, ny: [((i - (nx - 1) / 2) * s, (j - (ny - 1) / 2) * s)  # noqa: E731
                           for j in range(ny) for i in range(nx)]
    if n == 1:
        return [(0.0, 0.0)]
    if n == 2:
        return grid(2, 1)
    if n == 3:
        return [(0.0, 2 * h3 / 3), (-s / 2, -h3 / 3), (s / 2, -h3 / 3)]
    if n == 4:
        return grid(2, 2)
    if n == 5:
        a = s / math.sqrt(2)
        return [(-a, -a), (a, -a), (0.0, 0.0), (-a, a), (a, a)]
    if n == 6:
        return grid(3, 2)
    if n == 7:
        return [(0.0, 0.0)] + [(s * math.cos(math.radians(60 * k)),
                                s * math.sin(math.radians(60 * k))) for k in range(6)]
    if n == 8:
        return [(-s, -h3), (0.0, -h3), (s, -h3), (-s / 2, 0.0), (s / 2, 0.0),
                (-s, h3), (0.0, h3), (s, h3)]
    if n == 9:
        return grid(3, 3)
    if n == 10:
        return [(-s, -h3), (0.0, -h3), (s, -h3), (-1.5 * s, 0.0), (-s / 2, 0.0), (s / 2, 0.0),
                (1.5 * s, 0.0), (-s, h3), (0.0, h3), (s, h3)]
    if n == 12:
        return grid(4, 3)
    raise ValueError("A standard layout has 2 to 10 or 12 piles; type the coordinates "
                     "for any other group.")


def parse_coordinates(text: str) -> list[tuple[float, float]]:
    piles = []
    for part in [p for p in re.split(r"[;\n]", str(text)) if p.strip()]:
        numbers = [float(v) for v in re.split(r"[,\s]+", part.strip()) if v]
        if len(numbers) != 2:
            raise ValueError(f"A pile is given as x,y; {part.strip()!r} is not.")
        piles.append((numbers[0], numbers[1]))
    if len(piles) < 2:
        raise ValueError("Type at least two piles, as x,y; x,y.")
    return piles


# --------------------------------------------------------------------------- #
# Loads
# --------------------------------------------------------------------------- #
def _kind(text: str) -> bool:
    word = str(text).strip().upper()
    if word in ("F", "FACTORED", "U", "ULS", "LRFD", "STRENGTH"):
        return True
    if word in ("S", "SERVICE", "SLS", "ASD", "ALLOWABLE"):
        return False
    raise ValueError(f"A combination is S (service) or F (factored); {text!r} is neither.")


def _yes(text) -> bool | None:
    word = str(text).strip().upper()
    if word in ("", "NONE", "NAN"):
        return None
    return word in ("Y", "YES", "TRUE", "1")


def parse_loads(text: str) -> list[Combo]:
    """Combinations typed as ``name, S or F, P, Mx, My, Vx, Vy; ...``."""
    combos = []
    for part in [p for p in re.split(r"[;\n]", str(text)) if p.strip()]:
        cells = [c.strip() for c in part.split(",")]
        if len(cells) < 3:
            raise ValueError(f"A combination needs at least a name, S or F, and P: {part!r}")
        numbers = [float(c) if c else 0.0 for c in cells[2:7]]
        numbers += [0.0] * (5 - len(numbers))
        combos.append(Combo(cells[0], _kind(cells[1]), *numbers))
    if not combos:
        raise ValueError("No load combination was given.")
    return combos


HEADERS = {"name": ("name", "combination", "combo", "load combination"),
           "type": ("type", "kind", "s/f", "service/factored"),
           "P": ("p", "p (kn)", "axial", "fz"), "Mx": ("mx", "mx (kn-m)", "mx (kn·m)"),
           "My": ("my", "my (kn-m)", "my (kn·m)"), "Vx": ("vx", "vx (kn)", "fx"),
           "Vy": ("vy", "vy (kn)", "fy"), "lateral": ("we", "wind/eq", "lateral", "w/e"),
           "df": ("df", "dead factor"), "lf": ("lf", "surcharge factor")}


def _rows_to_combos(rows: list[list]) -> list[Combo]:
    rows = [r for r in rows if any(str(c).strip() for c in r if c is not None)]
    if not rows:
        raise ValueError("The loads file is empty.")
    head = [str(c).strip().lower() for c in rows[0]]
    at = {}
    for key, names in HEADERS.items():
        for i, cell in enumerate(head):
            if cell in names:
                at[key] = i
                break
    for key in ("name", "type", "P"):
        if key not in at:
            raise ValueError("The loads file needs the columns Name, Type (S or F) and P; "
                             "optional Mx, My, Vx, Vy, WE (Yes/No), DF, LF.")

    def cell(row, key, default=None):
        i = at.get(key)
        if i is None or i >= len(row) or row[i] is None or str(row[i]).strip() == "":
            return default
        return row[i]

    combos = []
    for row in rows[1:]:
        if cell(row, "name") is None:
            continue
        num = lambda key: float(cell(row, key, 0.0))  # noqa: E731
        df, lf = cell(row, "df"), cell(row, "lf")
        combos.append(Combo(str(cell(row, "name")), _kind(cell(row, "type")), num("P"),
                            num("Mx"), num("My"), num("Vx"), num("Vy"),
                            _yes(cell(row, "lateral", "")),
                            None if df is None else float(df), None if lf is None else float(lf)))
    if not combos:
        raise ValueError("The loads file has no combination under its heading row.")
    return combos


def read_loads_file(path: str) -> list[Combo]:
    """Combinations from an Excel (.xlsx) or CSV file: one row each, with the
    columns Name, Type (S or F), P, Mx, My, Vx, Vy and the optional WE, DF, LF."""
    if not path or not os.path.isfile(path):
        raise ValueError(f"The loads file was not found: {path!r}")
    if path.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        sheet = load_workbook(path, read_only=True, data_only=True).active
        return _rows_to_combos([list(r) for r in sheet.iter_rows(values_only=True)])
    with open(path, newline="", encoding="utf-8-sig") as handle:
        return _rows_to_combos(list(csv.reader(handle)))


def weight_cases(combo: Combo) -> list[Case]:
    """The factors on the cap's own weight to check a combination with.

    The column loads come factored; the cap, the soil over it and the
    surcharge are added here. Unless the combination gives its own factors,
    it is checked with the largest and the smallest weight: 1.2 and 0.9 on
    the dead weight (1.6 and 0 on the surcharge) for a factored combination
    (NSCP 203.3.1), 1.0 and 0.6 (1.0 and 0) for a service one (NSCP 203.4.1).
    """
    if combo.df is not None:
        return [Case(combo, combo.df, combo.lf if combo.lf is not None else combo.df)]
    if combo.factored:
        return [Case(combo, 1.2, 1.6), Case(combo, 0.9, 0.0)]
    return [Case(combo, 1.0, 1.0), Case(combo, 0.6, 0.0)]


# --------------------------------------------------------------------------- #
# Geometry
# --------------------------------------------------------------------------- #
def _shapes():
    from shapely.geometry import Point, box

    return Point, box


def column_shape(v: dict):
    Point, box = _shapes()
    if v["column_shape"] == "Circular":
        return Point(0, 0).buffer(v["cx"] / 2, quad_segs=32)
    return box(-v["cx"] / 2, -v["cy"] / 2, v["cx"] / 2, v["cy"] / 2)


def equivalent_column(v: dict) -> tuple[float, float]:
    """Column sides for the critical sections: a circular column is taken as a
    square of the same area (NSCP 413.2.7.3)."""
    if v["column_shape"] == "Circular":
        side = math.sqrt(math.pi) / 2 * v["cx"]
        return side, side
    return v["cx"], v["cy"]


def pile_shape(v: dict, x: float, y: float):
    Point, box = _shapes()
    r = v["dp"] / 2
    if v["pile_shape"] == "Round":
        return Point(x, y).buffer(r, quad_segs=32)
    return box(x - r, y - r, x + r, y + r)


def critical_perimeter(shape, offset: float, cap):
    """The two-way critical section at ``offset`` from a shape (NSCP 422.6.4.1),
    kept inside the cap: (its length in mm, the region it encloses, the number
    of cap edges it reaches)."""
    join = 2 if shape.geom_type == "Polygon" and len(shape.exterior.coords) <= 6 else 1
    region = shape.buffer(offset, join_style=join, quad_segs=32)
    inside = region.intersection(cap)
    edge_length = inside.boundary.intersection(cap.boundary.buffer(1e-6)).length
    b0 = max(inside.boundary.length - edge_length, 0.0)
    minx, miny, maxx, maxy = cap.bounds
    rx0, ry0, rx1, ry1 = region.bounds
    edges = sum((rx0 < minx - 1e-6, rx1 > maxx + 1e-6, ry0 < miny - 1e-6, ry1 > maxy + 1e-6))
    return b0, inside, edges


def share_outside(distance_outside: float, dp: float) -> float:
    """The part of a pile reaction that loads a section (NSCP 413.4.2.5): all of
    it with the centre dp/2 or more outside, none dp/2 or more inside, and a
    straight line between."""
    if distance_outside >= dp / 2:
        return 1.0
    if distance_outside <= -dp / 2:
        return 0.0
    return 0.5 + distance_outside / dp


# --------------------------------------------------------------------------- #
# Strengths and detailing (NSCP)
# --------------------------------------------------------------------------- #
def sqrt_fc(fc: float) -> float:
    return min(math.sqrt(fc), SQRT_FC_MAX)


def shrinkage_ratio(fy: float) -> float:
    """NSCP Table 424.4.3.2."""
    if fy < 420:
        return 0.0020
    return max(0.0018 * 420 / fy, 0.0014)


def flexural_steel(Mu_Nmm: float, b: float, d: float, fc: float, fy: float) -> tuple[float, str]:
    """Bars for a moment, tension controlled with phi = 0.90 (NSCP 421.2.2):
    (As in mm2, "" or the reason it cannot be done)."""
    if Mu_Nmm <= 0:
        return 0.0, ""
    Rn = Mu_Nmm / (PHI_FLEXURE * b * d * d)
    inside = 1 - 2 * Rn / (0.85 * fc)
    if inside <= 0:
        return math.inf, "the section is too small for the moment"
    rho = 0.85 * fc / fy * (1 - math.sqrt(inside))
    As = rho * b * d
    beta1 = min(max(0.85 - 0.05 * (fc - 28) / 7, 0.65), 0.85)
    c = As * fy / (0.85 * fc * b) / beta1
    if c / d > 0.375:      # net tensile strain under 0.005: not tension controlled
        return math.inf, "not tension controlled (c/d > 0.375): deepen the cap"
    return As, ""


def minimum_steel(As: float, b: float, d: float, h: float, fc: float, fy: float) -> float:
    """Minimum flexural steel of a pile cap, as the CRSI guide (4.3) reads
    NSCP 409.6.1.2 and 409.6.1.3 with the shrinkage ratio of Table 424.4.3.2:
    the beam minimum, or 4/3 of the steel needed, but not less than the
    shrinkage steel."""
    beam_min = max(0.25 * math.sqrt(fc) / fy, 1.4 / fy) * b * d
    shrink = shrinkage_ratio(fy) * b * h
    return max(As, min(beam_min, max(4 / 3 * As, shrink)))


def development(db: float, fc: float, fy: float, top: bool,
                 cb: float | None = None) -> tuple[float, float]:
    """(straight ld, hooked ldh) in mm (NSCP 425.4.2 and 425.4.3.1).

    Straight, the shorter of the two equations NSCP permits, at least 300 mm
    (425.4.2.1), with psi_t = 1.3 for top bars:
    * Table 425.4.2.2: (fy psi_t psi_e / (2.1 lambda sqrt f'c)) db for 20 mm bars
      and smaller, 1.7 for larger;
    * Eq. 425.4.2.3a, when ``cb`` (the smaller of the cover to the bar centre and
      half the bar spacing) is given: (fy / (1.1 lambda sqrt f'c))
      (psi_t psi_e psi_s / ((cb + Ktr)/db)) db, with Ktr = 0, the confinement
      term at most 2.5, and psi_s = 0.8 for 20 mm bars and smaller.
    Hooked: 0.24 fy psi_e psi_c psi_r / (lambda sqrt f'c) db, at least 8 db and
    150 mm, with psi_c = 0.7 for the side cover of a cap (Table 425.4.3.2)."""
    root = min(math.sqrt(fc), SQRT_FC_MAX)          # NSCP 425.4.1.4
    psi_t = 1.3 if top else 1.0
    factor = 2.1 if db <= 20 else 1.7
    ld = fy * psi_t / (factor * LAMBDA * root) * db
    if cb:
        psi_s = 0.8 if db <= 20 else 1.0
        confinement = min(cb / db, 2.5)
        ld = min(ld, fy / (1.1 * LAMBDA * root) * psi_t * psi_s / confinement * db)
    ld = max(ld, 300.0)
    ldh = max(0.24 * fy * 0.7 / (LAMBDA * root) * db, 8 * db, 150.0)
    return ld, ldh


def two_way_vc(fc: float, beta: float, alpha_s: float, d: float, b0: float) -> float:
    """vc in MPa (NSCP Table 422.6.5.2), the least of (a), (b) and (c)."""
    r = LAMBDA * sqrt_fc(fc)
    return min(0.33 * r, 0.17 * (1 + 2 / beta) * r, 0.083 * (2 + alpha_s * d / b0) * r)


def alpha_s(edges: int) -> float:
    """NSCP 422.6.5.3: 40 interior, 30 edge, 20 corner."""
    return {0: 40.0, 1: 30.0}.get(edges, 20.0)


# --------------------------------------------------------------------------- #
# Calculation
# --------------------------------------------------------------------------- #
def _num(values: dict, key: str, default=None):
    value = values.get(key)
    return default if value is None or value == "" else float(value)


def prepare(values: dict) -> tuple[list[tuple[float, float]], float, float]:
    """Pile centres and cap size (mm) from the inputs."""
    if values.get("layout", "Standard") == "Standard":
        n = int(round(_num(values, "n_piles", 4)))
        s = _num(values, "spacing")
        piles = standard_layout(n, s)
        e = _num(values, "edge")
        xs, ys = [p[0] for p in piles], [p[1] for p in piles]
        # the cap is centred on the column: each half reaches the farthest pile plus e
        Lx = 2 * max(max(xs) + e, -(min(xs) - e))
        Ly = 2 * max(max(ys) + e, -(min(ys) - e))
        return piles, Lx, Ly
    piles = parse_coordinates(values.get("coords", ""))
    Lx, Ly = _num(values, "cap_lx"), _num(values, "cap_ly")
    if not Lx or not Ly:
        raise ValueError("With typed coordinates, give the cap length along X and along Y.")
    for x, y in piles:
        if abs(x) >= Lx / 2 or abs(y) >= Ly / 2:
            raise ValueError(f"The pile at ({x:g}, {y:g}) is outside the cap.")
    return piles, Lx, Ly


def calculate(values: dict) -> PileCap:
    """Design (or check) the pile cap from a dict of the inputs (keys of ``INPUTS``)."""
    v = dict(values)
    combos = v.get("combos") or _load_combos(v)
    piles, Lx, Ly = prepare(v)
    h0 = _num(v, "h")
    if not h0 or h0 <= 0:
        raise ValueError("The cap thickness must be positive.")
    if v.get("find_depth", "Yes") == "Yes":
        searched = []
        start = _least_start(v)
        h = start
        best = None
        while h <= MAX_DEPTH:
            cap = _design(v, combos, piles, Lx, Ly, h)
            searched.append((h, cap.passed, ", ".join(sorted({c.group for c in cap.failing()}))))
            if cap.passed:
                best = cap
                break
            if _depth_cannot_help(cap):
                best = cap
                break
            h += DEPTH_STEP
        cap = best or _design(v, combos, piles, Lx, Ly, h0)
        cap.searched = searched
    else:
        cap = _design(v, combos, piles, Lx, Ly, h0)
    cap.h_typed = h0
    return cap


def _least_start(v: dict) -> float:
    """The thinnest cap the search starts from: d of at least 300 mm (NSCP 413.4.2.1)."""
    h = (MIN_EFFECTIVE_DEPTH + v["embed"] + v["clear_above_pile"] + 1.5 * v["bottom_bar"])
    return math.ceil(h / DEPTH_STEP) * DEPTH_STEP


def _depth_cannot_help(cap: PileCap) -> bool:
    """True when only checks that a thicker cap does not change fail (pile capacity,
    a moment the pile group cannot resist)."""
    groups = {c.group for c in cap.failing()}
    return bool(groups) and groups <= {"Pile capacity", "Pile group", "Bearing"}


def _load_combos(v: dict) -> list[Combo]:
    if v.get("load_source") == "File":
        return read_loads_file(str(v.get("loads_file") or ""))
    return parse_loads(v.get("loads") or "")


def _design(v: dict, combos: list[Combo], piles, Lx: float, Ly: float, h: float) -> PileCap:
    Point, box = _shapes()
    cap = PileCap(values=v, piles=piles, Lx=Lx, Ly=Ly, h=h, combos=combos)
    db, dbt = v["bottom_bar"], v["top_bar"]
    tol = v.get("tolerance") or 0.0
    cap_poly = box(-Lx / 2, -Ly / 2, Lx / 2, Ly / 2)
    col = column_shape(v)
    cx_eq, cy_eq = equivalent_column(v)

    # ---- depths: the bars along the longer side are the lower layer (CRSI 5.2)
    bottom_of_bars = v["embed"] + v["clear_above_pile"]
    long_axis = "X" if Lx >= Ly else "Y"
    d_lower = h - bottom_of_bars - db / 2
    d = {long_axis: d_lower, ("Y" if long_axis == "X" else "X"): d_lower - db}
    d_avg = (d["X"] + d["Y"]) / 2
    d_top = h - v["top_cover"] - dbt / 2
    cap.geometry = dict(d=d, d_avg=d_avg, d_top=d_top, long_axis=long_axis, cap=cap_poly,
                        column=col, cx_eq=cx_eq, cy_eq=cy_eq)

    # ---- weights (kN), at the cap centre
    area = Lx * Ly / 1e6
    col_area = col.area / 1e6
    cap.geometry["weights"] = dict(cap=v["gamma_c"] * area * h / 1000,
                                   soil=v["gamma_s"] * v["soil_h"] * (area - col_area),
                                   surcharge=v["surcharge"] * (area - col_area))

    # ---- pile reactions (rigid cap, pinned piles: CRSI 3.3 and 6.2)
    xs = [p[0] for p in piles]
    ys = [p[1] for p in piles]
    n = len(piles)
    xc, yc = sum(xs) / n, sum(ys) / n
    u = [x - xc for x in xs]
    w_ = [y - yc for y in ys]
    Suu = sum(a * a for a in u)
    Svv = sum(b * b for b in w_)
    Suv = sum(a * b for a, b in zip(u, w_))
    det = Suu * Svv - Suv * Suv
    collinear = abs(det) < 1e-6 * max(Suu * Svv, 1.0)
    h_lever = (h - v["embed"]) / 1000          # m: from the top of the cap to the pile heads
    cases = [case for combo in combos for case in weight_cases(combo)]
    unresisted = []
    for case in cases:
        c = case.combo
        W = (case.df * (cap.geometry["weights"]["cap"] + cap.geometry["weights"]["soil"])
             + case.lf * cap.geometry["weights"]["surcharge"])
        P = c.P + W
        My = c.My + c.Vx * h_lever - P * xc / 1000      # kN-m about the group centroid
        Mx = c.Mx + c.Vy * h_lever - P * yc / 1000
        if collinear:
            # piles in one line: only the moment about the line's normal is resisted
            if Suu >= Svv:
                b_ = My / Suu * 1000 if Suu else 0.0
                c_ = 0.0
                if abs(Mx) > 1e-6:
                    unresisted.append((case.label, "Mx", Mx))
            else:
                c_ = Mx / Svv * 1000 if Svv else 0.0
                b_ = 0.0
                if abs(My) > 1e-6:
                    unresisted.append((case.label, "My", My))
        else:
            b_ = (My * Svv - Mx * Suv) / det * 1000
            c_ = (Mx * Suu - My * Suv) / det * 1000
        cap.reactions[case.label] = [P / n + b_ * a + c_ * b for a, b in zip(u, w_)]
        cap.lateral_per_pile[case.label] = math.hypot(c.Vx, c.Vy) / n
    case_of = {case.label: case for case in cases}

    # ---- pile group and capacity
    _pile_capacity_checks(cap, cases, unresisted)

    # ---- moments and bars
    _flexure(cap, cases, case_of, tol)

    # ---- shear
    _one_way_shear(cap, cases, tol)
    _two_way_column(cap, cases, tol)
    _pile_punching(cap, cases)

    # ---- deep caps: CRSI special checks and strut-and-tie (NSCP 423)
    _deep_cap(cap, cases, tol)
    _column_bearing(cap, cases)

    # ---- bars: flexure, minimum steel and the strut-and-tie ties
    _finalize_bars(cap)

    # ---- detailing and recommendations
    _detailing(cap)
    return cap


def _pile_capacity_checks(cap: PileCap, cases: list[Case], unresisted) -> None:
    v = cap.values
    method = v.get("capacity_method", "Allowable")
    factored_method = method == "Factored resistance"
    used = [c for c in cases if c.combo.factored == factored_method]
    if not used:
        cap.notes.append(f"No {'factored' if factored_method else 'service'} combination was "
                         f"given, so the pile capacity ({method}) was not checked.")
    qc, qt, ql = v["q_comp"], v["q_tens"], v["q_lat"]
    clause = ("NSCP 413.4.1.1, 203.3" if factored_method else "NSCP 413.4.1.1, 203.4.1")
    if method.startswith("Allowable +"):
        clause = "NSCP 413.4.1.1, 203.4.2"
    worst = {"comp": None, "tens": None, "lat": None}
    for case in used:
        increase = 4 / 3 if method.startswith("Allowable +") and case.combo.has_lateral() else 1.0
        R = cap.reactions[case.label]
        for key, demand, capacity in (("comp", max(R), qc * increase),
                                      ("tens", max(-min(R), 0.0), qt * increase),
                                      ("lat", cap.lateral_per_pile[case.label], ql * increase)):
            ratio = demand / capacity if capacity > 0 else (math.inf if demand > 0 else 0)
            if worst[key] is None or ratio > worst[key][0]:
                worst[key] = (ratio, demand, capacity, case.label, increase)
    names = {"comp": "Compression in the most loaded pile", "tens": "Tension in a pile",
             "lat": "Lateral load per pile (shared equally)"}
    for key, item in worst.items():
        if item is None:
            continue
        _, demand, capacity, label, increase = item
        note = "capacity increased by 1/3 (wind or earthquake)" if increase > 1 else ""
        if key == "lat":
            note = (note + "; " if note else "") + "CRSI 3.3: equal shares, passive soil " \
                "neglected"
        cap.checks.append(Check("Pile capacity", names[key], demand, capacity, "kN", clause,
                                label, note))
    for label, axis, M in unresisted[:1]:
        cap.checks.append(Check(
            "Pile group", f"{axis} across a single line of piles", abs(M), 0.0, "kN-m",
            "NSCP 413.2.6.1", label,
            "the piles are in one line and cannot resist this moment; provide tie beams "
            "or a wider group"))


def _flexure(cap: PileCap, cases, case_of, tol: float) -> None:
    v = cap.values
    fc, fy = v["fc"], v["fy"]
    g = cap.geometry
    factored = [c for c in cases if c.combo.factored]
    if not factored:
        cap.notes.append("No factored combination was given: the cap was not designed.")
        return
    cx_eq, cy_eq = g["cx_eq"], g["cy_eq"]
    at_face = v.get("flexure_section", "Column face") == "Column face"
    section = {"X": cx_eq / 2 if at_face else cx_eq / 4, "Y": cy_eq / 2 if at_face else cy_eq / 4}
    clause_sec = "NSCP 413.2.7.1, 413.2.7.3" if at_face else "CRSI 5.2 (c/4)"
    w_area = {}
    for case in factored:
        w_area[case.label] = (case.df * (v["gamma_c"] * cap.h / 1000 + v["gamma_s"] * v["soil_h"])
                              + case.lf * v["surcharge"]) / 1e3       # kPa to N/mm2
    head_moment = _pile_head_moments(cap, factored)
    sag = {"X": (0.0, ""), "Y": (0.0, "")}
    hog = {"X": (0.0, ""), "Y": (0.0, "")}
    for axis in ("X", "Y"):
        L = cap.Lx if axis == "X" else cap.Ly
        B = cap.Ly if axis == "X" else cap.Lx
        s = section[axis]
        for case in factored:
            R = cap.reactions[case.label]
            for side in (1, -1):
                M = 0.0
                for (x, y), r in zip(cap.piles, R):
                    pos = (x if axis == "X" else y) * side
                    if pos > s:
                        M += r * 1000 * (pos - s + tol)               # N-mm
                overhang = L / 2 - s
                M -= w_area[case.label] * B * overhang * overhang / 2
                beyond = sum(1 for (x, y) in cap.piles if (x if axis == "X" else y) * side > s)
                extra = head_moment.get(case.label, {}).get(axis, 0.0) * beyond * 1e6
                for key, store, value in (("sag", sag, M + extra), ("hog", hog, -M + extra)):
                    if value / 1e6 > store[axis][0]:
                        store[axis] = (value / 1e6, case.label)
    g["Mu"] = dict(sag=sag, hog=hog, section=section)
    g["clause_section"] = clause_sec
    needs = []
    for axis in ("X", "Y"):
        B = cap.Ly if axis == "X" else cap.Lx
        dd = g["d"][axis]
        Mu, label = sag[axis]
        As_f, why = flexural_steel(Mu * 1e6, B, dd, fc, fy)
        if why:
            cap.checks.append(Check("Flexure", f"Bottom bars along {axis}", Mu, None, "kN-m",
                                    "NSCP 421.2.2", label, why, ratio_override=math.inf))
        As_min = minimum_steel(As_f if math.isfinite(As_f) else 0.0, B, dd, cap.h, fc, fy)
        if cap.n == 2 and axis != ("X" if cap.Lx >= cap.Ly else "Y"):
            As_min = max(As_f if math.isfinite(As_f) else 0.0,
                         shrinkage_ratio(fy) * B * cap.h)       # CRSI 4.3: 2-pile caps
        As_req = As_min
        short = (axis == "X") != (cap.Lx >= cap.Ly)
        beta = max(cap.Lx, cap.Ly) / min(cap.Lx, cap.Ly)
        note = ""
        if short and beta > 1.0 and cap.n > 2:
            As_req *= 2 * beta / (beta + 1)
            note = (f"short bars x 2 beta/(beta+1) = {2 * beta / (beta + 1):.3f} for a "
                    "uniform spacing (NSCP 413.3.3.3, CRSI 4.2)")
        needs.append(dict(axis=axis, face="Bottom", db=v["bottom_bar"],
                          typed=v.get("bottom_spacing"), As_req=As_req, As_f=As_f, As_min=As_min,
                          d=dd, B=B, Mu=Mu, section=section[axis], note=note, label=label))
    # top bars wherever a section has top tension (NSCP 414.1.5, 418.13.2.4)
    for axis in ("X", "Y"):
        Mu, label = hog[axis]
        if Mu <= 0.5:
            continue
        B = cap.Ly if axis == "X" else cap.Lx
        dd = g["d_top"] - (v["top_bar"] if axis != g["long_axis"] else 0.0)
        As_f, why = flexural_steel(Mu * 1e6, B, dd, fc, fy)
        if why:
            cap.checks.append(Check("Flexure", f"Top bars along {axis}", Mu, None, "kN-m",
                                    "NSCP 421.2.2", label, why, ratio_override=math.inf))
        As_min = minimum_steel(As_f if math.isfinite(As_f) else 0.0, B, dd, cap.h, fc, fy)
        needs.append(dict(axis=axis, face="Top", db=v["top_bar"], typed=v.get("top_spacing"),
                          As_req=As_min, As_f=As_f, As_min=As_min, d=dd, B=B, Mu=Mu,
                          section=section[axis],
                          note="top tension under a pile in tension or the cap's weight",
                          label=label))
    g["needs"] = needs


def _finalize_bars(cap: PileCap) -> None:
    """Choose the bars of each face for the larger of the flexural steel and the
    tie of the strut-and-tie model, then check their spacing and development."""
    g = cap.geometry
    ties = g.get("ties", {})
    cap.bars = []
    for need in g.get("needs", []):
        As_tie = ties.get(need["axis"], (0.0, 0.0, ""))[0] if need["face"] == "Bottom" else 0.0
        As_req = max(need["As_req"], As_tie)
        bars = _choose_bars(cap, need["axis"], need["face"], need["db"], need["typed"], As_req,
                            need["As_f"], need["As_min"], need["d"], need["B"], need["Mu"],
                            need["section"], need["note"])
        bars.As_tie = As_tie
        cap.bars.append(bars)
        cap.checks.append(Check(
            "Flexure", f"{bars.face} bars along {bars.direction}: {bars.count} - "
            f"{bars.diameter:g} mm @ {bars.spacing:g}", bars.As_req, bars.As_prov, "mm2",
            f"NSCP 409.6.1.2, 409.6.1.3, Table 424.4.3.2; {g['clause_section']}", need["label"],
            f"Mu = {bars.Mu:.1f} kN-m over {bars.width:.0f} mm, d = {bars.d:.0f} mm"
            + (f"; includes the strut-and-tie tie, {As_tie:,.0f} mm2" if As_tie > need["As_req"]
               else "")))
        if bars.anchorage == "Not needed":
            continue
        cap.checks.append(Check(
            "Detailing", f"Development of the {bars.face.lower()} bars along {bars.direction}",
            bars.ldh if bars.anchorage != "Straight" else bars.ld, bars.available, "mm",
            "NSCP 413.2.8, 425.4.2.2, 425.4.2.3, 425.4.3.1", "",
            f"{bars.anchorage.lower()} ends; ld = {bars.ld:.0f} mm, ldh = {bars.ldh:.0f} mm"))
    for axis, (As_tie, T, label) in ties.items():
        bars = next((b for b in cap.bars if b.face == "Bottom" and b.direction == axis), None)
        if bars:
            cap.checks.append(Check("Strut-and-tie", f"Tie along {axis} (bottom bars)", As_tie,
                                    bars.As_prov, "mm2", "NSCP 423.3.1, 423.7.2", label,
                                    f"tie force {T:.0f} kN; As = T / (phi fy)"))


def _pile_head_moments(cap: PileCap, cases) -> dict:
    """Fixed heads: each pile takes M = H Lf / 2 (kN-m) in each direction."""
    v = cap.values
    if v.get("head", "Pinned") != "Fixed":
        return {}
    lf = _num(v, "lf")
    lf_mm = lf * 1000 if lf else (LF_FIRM if v.get("ground", "Firm") == "Firm" else LF_SOFT)
    cap.geometry["Lf"] = lf_mm
    out = {}
    for case in cases:
        c = case.combo
        out[case.label] = {"X": abs(c.Vx) / cap.n * lf_mm / 2000,
                           "Y": abs(c.Vy) / cap.n * lf_mm / 2000}
    return out


def _choose_bars(cap: PileCap, axis, face, db, typed_spacing, As_req, As_f, As_min, d, B, Mu,
                 section, note) -> Bars:
    """Bars of a face: the count and spacing (rounded down to 25 mm), within
    the maximum spacing (NSCP 407.7.2.3) and the minimum clear spacing
    (NSCP 425.2.1), and how their ends are anchored."""
    Ab = math.pi * db * db / 4
    width = B - 2 * SIDE_COVER - db
    s_max = min(3 * cap.h, 450.0)
    s_min = db + max(25.0, db, 4 / 3 * AGGREGATE)
    if typed_spacing:
        s = float(typed_spacing)
        count = int(math.floor(width / s + 1e-9)) + 1
    else:
        count = max(int(math.ceil(As_req / Ab - 1e-9)), 2) if math.isfinite(As_req) else 2
        count = max(count, int(math.ceil(width / s_max)) + 1)
        s = math.floor(width / (count - 1) / SPACING_STEP) * SPACING_STEP
        while s > s_max:
            count += 1
            s = math.floor(width / (count - 1) / SPACING_STEP) * SPACING_STEP
        if s < s_min:     # this bar size cannot give the steel: as close as allowed
            s = math.ceil(s_min / SPACING_STEP) * SPACING_STEP
        count = int(math.floor(width / s + 1e-9)) + 1
    As_prov = count * Ab
    if s < s_min:
        cap.checks.append(Check("Detailing", f"Clear spacing of the {face.lower()} bars along "
                                f"{axis}", s_min, s, "mm", "NSCP 425.2.1", "",
                                "bars too close: use a larger bar"))
    if typed_spacing and s > s_max:
        cap.checks.append(Check("Detailing", f"Spacing of the {face.lower()} bars along {axis}",
                                s, s_max, "mm", "NSCP 407.7.2.3"))
    v = cap.values
    cover_to_centre = (v["top_cover"] if face == "Top"
                       else v["embed"] + v["clear_above_pile"]) + db / 2
    cb = min(cover_to_centre, SIDE_COVER + db / 2, s / 2)
    ld, ldh = development(db, cap.fc, cap.fy, top=(face == "Top"), cb=cb)
    L = cap.Lx if axis == "X" else cap.Ly
    # developed from the column face (NSCP 413.2.8.3), also when the moment is taken at c/4
    half_column = (cap.geometry["cx_eq"] if axis == "X" else cap.geometry["cy_eq"]) / 2
    available = L / 2 - SIDE_COVER - max(section, half_column)
    if Mu <= 0.5 and not (As_f and math.isfinite(As_f) and As_f > 0):
        anchorage = "Not needed"          # minimum steel only: no stress to develop
    elif ld <= available:
        anchorage = "Straight"
    elif ldh <= available:
        anchorage = "Hooked"
    else:
        anchorage = "NOT DEVELOPED"
    return Bars(axis, face, db, s, count, As_req, As_prov, As_f if math.isfinite(As_f) else 0.0,
                As_min, 0.0, d, B, anchorage, ld, ldh, available, bool(typed_spacing), Mu, note)


def _one_way_shear(cap: PileCap, cases, tol: float) -> None:
    v = cap.values
    g = cap.geometry
    factored = [c for c in cases if c.combo.factored]
    for axis in ("X", "Y"):
        dd = g["d"][axis]
        half = (g["cx_eq"] if axis == "X" else g["cy_eq"]) / 2
        L = cap.Lx if axis == "X" else cap.Ly
        B = cap.Ly if axis == "X" else cap.Lx
        s = half + dd                       # NSCP 413.2.7.2 with 407.4.3 / 408.4.3
        if s >= L / 2:
            cap.checks.append(Check("One-way shear", f"Section at d from the column, along {axis}",
                                    None, None, "kN", "NSCP 413.4.2.3(a)", "",
                                    "the section falls outside the cap: no one-way shear"))
            continue
        worst = (0.0, "")
        for case in factored:
            w = (case.df * (v["gamma_c"] * cap.h / 1000 + v["gamma_s"] * v["soil_h"])
                 + case.lf * v["surcharge"]) / 1e3          # kPa to N/mm2
            for side in (1, -1):
                V = 0.0
                for (x, y), r in zip(cap.piles, cap.reactions[case.label]):
                    pos = (x if axis == "X" else y) * side + tol
                    V += r * 1000 * share_outside(pos - s, cap.dp)
                V -= w * B * (L / 2 - s)
                if abs(V) / 1000 > worst[0]:
                    worst = (abs(V) / 1000, case.label)
        Vc = 0.17 * LAMBDA * sqrt_fc(v["fc"]) * B * dd / 1000
        cap.checks.append(Check("One-way shear",
                                f"At d = {dd:.0f} mm from the column, along {axis}",
                                worst[0], PHI_SHEAR * Vc, "kN",
                                "NSCP 413.4.2.3(a), 422.5.5.1, 413.4.2.5", worst[1],
                                f"phi Vc = 0.75 x 0.17 sqrt(f'c) b d, b = {B:.0f} mm"))


def _two_way_column(cap: PileCap, cases, tol: float) -> None:
    v = cap.values
    g = cap.geometry
    d = g["d_avg"]
    b0, region, edges = critical_perimeter(g["column"], d / 2, g["cap"])
    g["column_perimeter"] = region
    if b0 <= 0:
        cap.checks.append(Check("Two-way shear", "Around the column at d/2", None, None, "kN",
                                "NSCP 413.4.2.3(b)", "", "the section falls outside the cap"))
        return
    from shapely.geometry import Point

    beta = 1.0 if v["column_shape"] == "Circular" else max(v["cx"], v["cy"]) / min(v["cx"], v["cy"])
    vc = two_way_vc(v["fc"], beta, alpha_s(edges), d, b0)
    worst = (0.0, "")
    outside_area = (cap.Lx * cap.Ly - region.area)
    for case in [c for c in cases if c.combo.factored]:
        w = (case.df * (v["gamma_c"] * cap.h / 1000 + v["gamma_s"] * v["soil_h"])
             + case.lf * v["surcharge"]) / 1e3          # kPa to N/mm2
        V = 0.0
        for (x, y), r in zip(cap.piles, cap.reactions[case.label]):
            dist = region.exterior.distance(Point(x, y))
            inside = region.contains(Point(x, y))
            V += r * 1000 * share_outside(-dist + tol if inside else dist + tol, cap.dp)
        V -= w * outside_area
        if V / 1000 > worst[0]:
            worst = (V / 1000, case.label)
    cap.checks.append(Check("Two-way shear", f"Around the column at d/2, b0 = {b0:.0f} mm",
                            worst[0], PHI_SHEAR * vc * b0 * d / 1000, "kN",
                            "NSCP 413.4.2.3(b), 422.6.5.2, 413.4.2.5", worst[1],
                            f"vc = {vc:.3f} MPa (alpha_s = {alpha_s(edges):g}, beta = {beta:.2f}), "
                            f"d = {d:.0f} mm"))


def _pile_punching(cap: PileCap, cases) -> None:
    """Two-way shear around single piles, piles whose sections overlap, and corner
    piles (CRSI limit states P1 to P3), and one-way shear across a corner pile (P4)."""
    v = cap.values
    g = cap.geometry
    d = g["d_avg"]
    from shapely.geometry import Point
    from shapely.ops import unary_union

    factored = [c for c in cases if c.combo.factored]
    if not factored:
        return
    Rmax = [max(cap.reactions[c.label][i] for c in factored) for i in range(cap.n)]
    case_max = [max(factored, key=lambda c: cap.reactions[c.label][i]).label for i in range(cap.n)]
    # the candidates (CRSI 5.4): each pile alone (P1, or P3 at a corner), and each
    # pair of piles whose sections at d/2 overlap, around both together (P2)
    candidates = [[i] for i in range(cap.n)]
    reach = cap.dp + d
    for i in range(cap.n):
        for j in range(i + 1, cap.n):
            if math.dist(cap.piles[i], cap.piles[j]) < reach:
                candidates.append([i, j])
    worst = {}
    for members in candidates:
        if all(g["column"].contains(Point(cap.piles[i])) for i in members):
            continue          # a pile under the column carries its load straight down
        shape = unary_union([pile_shape(v, *cap.piles[i]) for i in members]).convex_hull
        b0, region, edges = critical_perimeter(shape, d / 2, g["cap"])
        if len(members) == 1 and edges >= 2:
            # a corner pile: the section may run straight to the two edges (CRSI P3)
            x, y = cap.piles[members[0]]
            ex, ey = cap.Lx / 2 - abs(x), cap.Ly / 2 - abs(y)
            b0 = min(b0, math.pi * reach / 4 + ex + ey)
        if b0 <= 0:
            continue
        Vu = sum(max(Rmax[i], 0.0) for i in members)
        vc = two_way_vc(v["fc"], 1.0, alpha_s(edges), d, b0)
        cap_ = PHI_SHEAR * vc * b0 * d / 1000
        kind = "pair" if len(members) == 2 else "single"
        name = (f"piles {members[0] + 1} and {members[1] + 1}" if kind == "pair"
                else f"pile {members[0] + 1}")
        item = (Vu / cap_ if cap_ else math.inf, name, Vu, cap_, b0, vc, edges,
                case_max[members[0]])
        if kind not in worst or item[0] > worst[kind][0]:
            worst[kind] = item
    for kind, limit_states in (("single", "P1, P3"), ("pair", "P2")):
        if kind not in worst:
            continue
        _, name, Vu, cap_, b0, vc, edges, label = worst[kind]
        where = {0: "interior", 1: "edge"}.get(edges, "corner")
        cap.checks.append(Check("Pile punching", f"Two-way shear around {name} ({where}), "
                                f"b0 = {b0:.0f} mm", Vu, cap_, "kN",
                                f"NSCP 422.6.5.2; CRSI 5.4 ({limit_states})", label,
                                f"vc = {vc:.3f} MPa, d = {d:.0f} mm, no reduction for the cap's "
                                "weight (CRSI)"))
    # P4: one-way shear across a corner pile (CRSI 5.4)
    worst4 = None
    for i, (x, y) in enumerate(cap.piles):
        ex, ey = cap.Lx / 2 - abs(x), cap.Ly / 2 - abs(y)
        E = min(ex, ey)
        if ex > cap.dp / 2 + d or ey > cap.dp / 2 + d:
            continue
        b = 2 * (E * math.sqrt(2) + cap.dp / 2 + min(d, 330.0))
        Vc = 0.17 * LAMBDA * sqrt_fc(v["fc"]) * b * d / 1000
        item = (max(Rmax[i], 0) / (PHI_SHEAR * Vc), i, b, Vc)
        if worst4 is None or item[0] > worst4[0]:
            worst4 = item
    if worst4:
        _, i, b, Vc = worst4
        cap.checks.append(Check("Pile punching", f"One-way shear across corner pile {i + 1}, "
                                f"b = {b:.0f} mm", max(Rmax[i], 0), PHI_SHEAR * Vc, "kN",
                                "NSCP 422.5.5.1; CRSI 5.4 (P4)", case_max[i]))


def _distance_to_face(cap: PileCap) -> dict:
    """w: the distance in plan from the column face to the nearest pile centre,
    along X and along Y, with the placement tolerance (CRSI 5.3)."""
    g = cap.geometry
    tol = cap.values.get("tolerance") or 0.0
    out = {}
    for axis, half in (("X", g["cx_eq"] / 2), ("Y", g["cy_eq"] / 2)):
        ws = [abs(p[0 if axis == "X" else 1]) - half + tol for p in cap.piles
              if abs(p[0 if axis == "X" else 1]) > half]
        out[axis] = min(ws) if ws else None
    return out


def _deep_cap(cap: PileCap, cases, tol: float) -> None:
    """Piles close to the column (CRSI 5.3, the condition of ACI R13.4.2.5):
    the CRSI shear checks at the column face, for information, and a
    strut-and-tie model (NSCP 413.2.6.3, Section 423), which then governs."""
    v = cap.values
    g = cap.geometry
    w = _distance_to_face(cap)
    d = g["d_avg"]
    near_one_way = {a: (w[a] is not None and w[a] < g["d"][a]) for a in ("X", "Y")}
    near_two_way = any(w[a] is not None and w[a] < d / 2 for a in ("X", "Y"))
    cap.deep = any(near_one_way.values()) or near_two_way
    g["w"] = w
    if not cap.deep:
        return
    factored = [c for c in cases if c.combo.factored]
    if not factored:
        return
    stm = v.get("deep_method", "Strut-and-tie") == "Strut-and-tie"
    g["deep_method"] = "Strut-and-tie" if stm else "CRSI special checks"
    info = "for information: the strut-and-tie model governs" if stm else \
        "for information: the CRSI special checks govern"
    for c in cap.checks:
        if c.group in ("One-way shear", "Two-way shear"):
            c.governs = False
            c.note = (c.note + "; " if c.note else "") + info
    root = sqrt_fc(v["fc"])
    c_side = min(g["cx_eq"], g["cy_eq"])
    # CRSI limit state 4: two-way shear at the column face
    if near_two_way:
        ws = [w[a] for a in ("X", "Y") if w[a] is not None]
        bs = 2 * (g["cx_eq"] + g["cy_eq"])
        caps_ = []
        for wa in ws:
            vc = min(d / max(wa, 1.0) * (1 + d / c_side) * 0.166 * root, 2.66 * root)
            caps_.append(PHI_SHEAR * vc * bs * d / 1000)
        phiVc = sum(caps_) / len(caps_)
        Vu, label = _shear_outside_column(cap, factored, tol)
        cap.checks.append(Check("Deep cap (CRSI)", "Two-way shear at the column face (LS 4)", Vu,
                                phiVc, "kN", "CRSI 5.3, 5.4; ACI R13.4.2.5", label,
                                info if stm else "", governs=not stm))
    for axis in ("X", "Y"):
        if not near_one_way[axis]:
            continue
        dd = g["d"][axis]
        B = cap.Ly if axis == "X" else cap.Lx
        half = (g["cx_eq"] if axis == "X" else g["cy_eq"]) / 2
        Vu, Mu, label = _face_shear_moment(cap, factored, axis, half, tol)
        r = min(max(Mu / (Vu * dd / 1000) if Vu > 0 else 1.0, 0.0), 1.0)
        vc = 0.83 * root if r <= 0 else min(
            dd / max(w[axis], 1.0) * (3.5 - 2.5 * r) * (0.158 + 0.0083 / r) * root, 0.83 * root)
        cap.checks.append(Check("Deep cap (CRSI)", f"One-way shear at the column face along {axis} "
                                f"(LS {5 if axis == 'X' else 6})", Vu,
                                PHI_SHEAR * vc * B * dd / 1000,
                                "kN", "CRSI 5.3, 5.4; ACI R13.4.2.5", label,
                                info if stm else "", governs=not stm))
    first = len(cap.checks)
    _strut_and_tie(cap, factored)
    if not stm:
        for c in cap.checks[first:]:
            c.governs = False
            c.note = (c.note + "; " if c.note else "") + info
        g["ties"] = {}


def _column_bearing(cap: PileCap, cases) -> None:
    """Bearing of the column on the cap (NSCP 422.8.3) and the bars across the
    interface (NSCP 416.3). Where a strut-and-tie model governs, its node under
    the column is checked instead (NSCP 422.8.1.2)."""
    if cap.deep and cap.geometry.get("deep_method") == "Strut-and-tie":
        return
    factored = [c for c in cases if c.combo.factored]
    if not factored:
        return
    v = cap.values
    g = cap.geometry
    col = g["column"]
    A1 = col.area
    minx, miny, maxx, maxy = col.bounds
    # the frustum sloped 1 vertical to 2 horizontal, within the cap (NSCP 422.8.3.2)
    e = max(min(2 * cap.h, cap.Lx / 2 - maxx, minx + cap.Lx / 2,
                cap.Ly / 2 - maxy, miny + cap.Ly / 2), 0.0)
    A2 = col.buffer(e, join_style=2).area if v["column_shape"] != "Circular" else col.buffer(e).area
    factor = min(math.sqrt(A2 / A1), 2.0)
    phiBn = PHI_BEARING * factor * 0.85 * v["fc"] * A1 / 1000
    weights = g["weights"]
    Pu, label = 0.0, ""
    for case in factored:
        W = case.df * (weights["cap"] + weights["soil"]) + case.lf * weights["surcharge"]
        P = sum(cap.reactions[case.label]) - W
        if P > Pu:
            Pu, label = P, case.label
    excess = max(Pu - phiBn, 0.0)
    As_min = MIN_DOWEL_RATIO * A1
    As_req = max(As_min, excess * 1000 / (PHI_BEARING * v["fy"]))
    g["bearing"] = {"A1": A1, "A2": A2, "e": e, "factor": factor, "phiBn": phiBn, "Pu": Pu,
                    "As_req": As_req}
    note = (f"phi 0.65 x {factor:.2f} x 0.85 f'c A1; A1 = {A1:,.0f} mm2, A2 = {A2:,.0f} mm2; "
            "the column's own bearing on its concrete is part of the column design (NSCP 416.3.3.4)")
    if excess > 0:
        note += (f"; the excess {excess:,.0f} kN goes to the column bars or dowels "
                 "(NSCP 416.3.1.2), see the bars across the interface")
    cap.checks.append(Check("Bearing", "Column on the cap", Pu, phiBn, "kN",
                            "NSCP 422.8.3, Table 422.8.3.2, Table 421.2.1(d)", label, note,
                            governs=excess == 0))
    cap.checks.append(Check("Detailing", "Column bars or dowels across the interface, area needed",
                            As_req, None, "mm2", "NSCP 416.3.1.2, 416.3.4.1, 416.3.5.1", label,
                            f"at least 0.005 Ag = {As_min:,.0f} mm2"
                            + (f"; the bearing excess over phi fy, phi = 0.65: "
                               f"{excess * 1000 / (PHI_BEARING * v['fy']):,.0f} mm2"
                               if excess > 0 else ""), governs=False))
    if excess > 0:
        cap.notes.append(f"The column load {Pu:,.0f} kN is over the bearing strength of the cap, "
                         f"{phiBn:,.0f} kN (NSCP 422.8.3): the column bars or dowels across the "
                         f"interface must carry the excess, at least {As_req:,.0f} mm2 of bars "
                         "developed both ways (NSCP 416.3.1.2, 416.3.5.1).")


def _shear_outside_column(cap, factored, tol):
    v = cap.values
    g = cap.geometry
    from shapely.geometry import Point

    worst = (0.0, "")
    for case in factored:
        w = (case.df * (v["gamma_c"] * cap.h / 1000 + v["gamma_s"] * v["soil_h"])
             + case.lf * v["surcharge"]) / 1e3          # kPa to N/mm2
        V = sum(r for (x, y), r in zip(cap.piles, cap.reactions[case.label])
                if not g["column"].contains(Point(x, y))) * 1000
        V -= w * (cap.Lx * cap.Ly - g["column"].area)
        if V / 1000 > worst[0]:
            worst = (V / 1000, case.label)
    return worst


def _face_shear_moment(cap, factored, axis, half, tol):
    v = cap.values
    worst = (0.0, 0.0, "")
    L = cap.Lx if axis == "X" else cap.Ly
    B = cap.Ly if axis == "X" else cap.Lx
    for case in factored:
        w = (case.df * (v["gamma_c"] * cap.h / 1000 + v["gamma_s"] * v["soil_h"])
             + case.lf * v["surcharge"]) / 1e3          # kPa to N/mm2
        for side in (1, -1):
            V = M = 0.0
            for (x, y), r in zip(cap.piles, cap.reactions[case.label]):
                pos = (x if axis == "X" else y) * side
                if pos > half:
                    V += r * 1000
                    M += r * 1000 * (pos - half + tol)
            over = L / 2 - half
            V -= w * B * over
            M -= w * B * over * over / 2
            if V / 1000 > worst[0]:
                worst = (V / 1000, M / 1e6, case.label)
    return worst


def _strut_and_tie(cap: PileCap, factored) -> None:
    """A three-dimensional strut-and-tie model (NSCP 413.2.6.3 and Section 423;
    ACI R13.2.6.3, Adebar et al. 1990): a strut from a node under the column to
    each pile, and ties along X and Y at the level of the bottom bars.

    * The node under a pile is the pile top; its tie zone is twice the depth of
      the bar centroid above the pile top. A node anchoring ties in two
      directions is CTT (beta_n = 0.60), in one direction CCT (0.80), NSCP
      Table 423.9.2.
    * The node under the column is CCC (beta_n = 1.0), at the quarter point of
      the column toward the pile; its depth a holds the horizontal strut that
      balances the tie: a = T / (phi 0.85 f'c b), with b the column side
      across the tie. The lever arm is z = d - a/2, found by iteration.
    * Struts use beta_s = 0.60 (NSCP 413.4.2.4, Table 423.4.3(c)). The end of a
      strut has the area A sin(theta) + b u cos(theta): the bearing area seen
      along the strut plus the tie or compression zone seen across it.
    """
    v = cap.values
    g = cap.geometry
    fc, fy = v["fc"], v["fy"]
    d = g["d_avg"]
    Ap = math.pi * cap.dp ** 2 / 4 if v["pile_shape"] == "Round" else cap.dp ** 2
    u_pile = 2 * (v["clear_above_pile"] + v["bottom_bar"])
    A_col = g["column"].area
    worst: dict = {}

    def keep(key, ratio, *item):
        if key not in worst or ratio > worst[key][0]:
            worst[key] = (ratio, *item)

    weights = g["weights"]
    for case in factored:
        R = cap.reactions[case.label]
        # the cap, soil and surcharge reach the piles without passing the column
        W = (case.df * (weights["cap"] + weights["soil"]) + case.lf * weights["surcharge"])
        Pu = sum(R) - W
        Rpos = sum(r for r in R if r > 0)
        column_bearing = PHI_STM * 0.85 * BETA_N["CCC"] * fc * A_col / 1000
        keep("column bearing", Pu / column_bearing, Pu, column_bearing, case.label)
        z = 0.9 * d
        a = 0.0
        struts = []
        T_x = T_y = 0.0
        for _ in range(8):          # the lever arm and the top node depth together
            Tx = {1: 0.0, -1: 0.0}
            Ty = {1: 0.0, -1: 0.0}
            struts = []
            for (x, y), r in zip(cap.piles, R):
                if r <= 0:
                    continue
                nx = math.copysign(g["cx_eq"] / 4, x) if abs(x) > g["cx_eq"] / 4 else x
                ny = math.copysign(g["cy_eq"] / 4, y) if abs(y) > g["cy_eq"] / 4 else y
                dx, dy = x - nx, y - ny
                length = math.hypot(dx, dy)
                if length < 1e-6:
                    continue
                theta = math.atan2(z, length)
                H = r / math.tan(theta)
                Tx[1 if x > 0 else -1] += H * abs(dx) / length
                Ty[1 if y > 0 else -1] += H * abs(dy) / length
                struts.append((x, y, r, theta, dx, dy))
            T_x, T_y = max(Tx.values()), max(Ty.values())
            a_x = T_x * 1000 / (PHI_STM * 0.85 * BETA_N["CCC"] * fc * g["cy_eq"])
            a_y = T_y * 1000 / (PHI_STM * 0.85 * BETA_N["CCC"] * fc * g["cx_eq"])
            a = max(a_x, a_y)
            z_new = max(d - a / 2, 0.2 * d)
            if abs(z_new - z) < 1.0:
                z = z_new
                break
            z = z_new
        for x, y, r, theta, dx, dy in struts:
            C = r / math.sin(theta)
            two_ties = abs(dx) > 1.0 and abs(dy) > 1.0
            kind = "CTT" if two_ties else "CCT"
            keep("angle", MIN_STRUT_ANGLE / math.degrees(theta), math.degrees(theta), case.label, z)
            bearing = PHI_STM * 0.85 * BETA_N[kind] * fc * Ap / 1000
            keep("pile bearing", r / bearing, r, bearing, case.label, kind)
            A_end = Ap * math.sin(theta) + cap.dp * u_pile * math.cos(theta)
            F = PHI_STM * 0.85 * min(BETA_S_PILE_CAP, BETA_N[kind]) * LAMBDA * fc * A_end / 1000
            keep("pile strut", C / F, C, F, case.label)
            share = A_col * r / Rpos if Rpos > 0 else A_col
            A_top = share * math.sin(theta) + math.sqrt(share) * a * math.cos(theta)
            F_top = PHI_STM * 0.85 * BETA_S_PILE_CAP * LAMBDA * fc * A_top / 1000
            C_col = max(r - W / cap.n, 0.0) / math.sin(theta)
            limit = PHI_STM * 0.85 * BETA_S_PILE_CAP * LAMBDA * fc * share / 1000
            keep("column strut", C_col / F_top, C_col, F_top, case.label,
                 max(r - W / cap.n, 0.0) / limit)
        keep("tieX", T_x, T_x, case.label, z)
        keep("tieY", T_y, T_y, case.label, z)
        if a > d:
            keep("top node", a / d, a, d, case.label)

    if "angle" in worst:
        _, theta, label, z = worst["angle"]
        cap.checks.append(Check("Strut-and-tie", "Least angle between a strut and the ties",
                                MIN_STRUT_ANGLE, theta, "deg", "NSCP 423.2.7", label,
                                f"lever arm z = {z:.0f} mm",
                                ratio_override=MIN_STRUT_ANGLE / theta))
    if "column bearing" in worst:
        _, Pu, F, label = worst["column bearing"]
        cap.checks.append(Check("Bearing", "Column on the cap (CCC node)", Pu, F, "kN",
                                "NSCP 423.9.1, Table 423.9.2(a)", label,
                                "phi 0.85 beta_n f'c A_column, beta_n = 1.0"))
    if "pile bearing" in worst:
        _, r, F, label, kind = worst["pile bearing"]
        cap.checks.append(Check("Bearing", f"Cap on the most loaded pile ({kind} node)", r, F, "kN",
                                f"NSCP 423.9.1, Table 423.9.2({'c' if kind == 'CTT' else 'b'})",
                                label,
                                f"phi 0.85 beta_n f'c A_pile, beta_n = {BETA_N[kind]:.2f}; a "
                                "thicker cap does not change it: a larger pile or a stronger "
                                "concrete does"))
    if "pile strut" in worst:
        _, C, F, label = worst["pile strut"]
        cap.checks.append(Check("Strut-and-tie", "Strut at the pile node", C, F, "kN",
                                "NSCP 423.3.1, 423.4.3, 413.4.2.4", label,
                                f"beta_s = 0.60; end area Ap sin(theta) + dp u cos(theta), "
                                f"u = {u_pile:.0f} mm"))
    if "column strut" in worst:
        _, C, F, label, at_vertical = worst["column strut"]
        hopeless = at_vertical > 1.0          # fails even with a vertical strut
        cap.checks.append(Check(
            "Bearing" if hopeless else "Strut-and-tie", "Strut at the column node", C, F, "kN",
            "NSCP 423.3.1, 423.4.3, 413.4.2.4", label,
            "beta_s = 0.60; the column load and area shared by the piles in proportion to their "
            "reactions (the cap's weight taken off)"
            + ("; no thickness helps: the column stress is over 0.75 x 0.85 x 0.60 f'c of the "
               "cap. Use a stronger concrete or a larger column, or the CRSI special checks"
               if hopeless else "")))
    if "top node" in worst:
        _, a, d_, label = worst["top node"]
        cap.checks.append(Check("Strut-and-tie", "Depth of the node under the column", a, d_,
                                "mm", "NSCP 423.9.1", label, "the compression zone needs more "
                                "depth than the cap has"))
    g["ties"] = {}
    for axis, key in (("X", "tieX"), ("Y", "tieY")):
        if key in worst:
            _, T, label, z = worst[key]
            g["ties"][axis] = (T * 1000 / (PHI_STM * fy), T, label)
    # ties are developed past the pile centre (NSCP 423.8.3; CRSI 5.3, tied arch)
    tol = v.get("tolerance") or 0.0
    _, ldh = development(v["bottom_bar"], fc, fy, top=False)
    for axis in ("X", "Y"):
        L = cap.Lx if axis == "X" else cap.Ly
        outer = max(abs(p[0 if axis == "X" else 1]) for p in cap.piles)
        available = L / 2 - SIDE_COVER - (outer + tol)
        cap.checks.append(Check("Strut-and-tie", f"Tie anchorage past the outer pile along {axis}",
                                ldh, available, "mm", "NSCP 423.8.3; CRSI 5.3 (tied arch)", "",
                                "hooked bars developed from the pile centre"
                                + ("" if ldh <= available else
                                   ": widen the edge distance or use headed bars")))


def _detailing(cap: PileCap) -> None:
    v = cap.values
    g = cap.geometry
    dp = cap.dp
    d_min = min(g["d"].values())
    cap.checks.append(Check("Detailing", "Effective depth of the bottom bars", MIN_EFFECTIVE_DEPTH,
                            d_min, "mm", "NSCP 413.4.2.1", "", "d of at least 300 mm"))
    cap.checks[-1].ratio_override = MIN_EFFECTIVE_DEPTH / d_min if d_min > 0 else math.inf
    bottom_cover = v["embed"] + v["clear_above_pile"]
    cap.checks.append(Check("Detailing", "Cover to the bottom bars (cast against earth)",
                            COVER_EARTH, bottom_cover, "mm", "NSCP Table 420.6.1.3.1", "",
                            "embedment plus clear cover above the piles"))
    cap.checks[-1].ratio_override = COVER_EARTH / bottom_cover if bottom_cover else math.inf
    # column bars into the cap (NSCP 418.13.2.1 and 418.13.2.2)
    db_col = _num(v, "col_bar")
    if db_col:
        _, ldh = development(db_col, cap.fc, cap.fy, top=False)
        ldh_plain = max(0.24 * cap.fy / (LAMBDA * math.sqrt(cap.fc)) * db_col, 8 * db_col, 150.0)
        available = cap.h - bottom_cover - 2 * v["bottom_bar"]
        cap.checks.append(Check("Detailing", f"Column bars ({db_col:g} mm) developed into the cap",
                                ldh_plain, available, "mm",
                                "NSCP 418.13.2.1, 418.13.2.2, 425.4.3.1",
                                "", "90-degree hooks resting on the bottom bars, turned toward "
                                "the column centre"))
    # recommendations: not code limits of NSCP, shown without failing the design
    spacing = min((math.dist(a, b) for i, a in enumerate(cap.piles) for b in cap.piles[i + 1:]),
                  default=None)
    if spacing:
        need = max(2.5 * dp, 760.0)
        cap.checks.append(Check("Recommendation", "Pile spacing, centre to centre", need, spacing,
                                "mm", "AASHTO LRFD 10.7.1.2 (FHWA GEC 12, 7.10.1)", "",
                                f"CRSI 4.3 recommends {max(3 * dp, dp + 600, 900):.0f} mm "
                                "(3 dp, dp + 600 and 900)", governs=False))
        cap.checks[-1].ratio_override = need / spacing
    edge_side = min(min(cap.Lx / 2 - abs(x), cap.Ly / 2 - abs(y)) for x, y in cap.piles) - dp / 2
    cap.checks.append(Check("Recommendation", "Cap edge to the side of a pile", 230.0, edge_side,
                            "mm", "AASHTO LRFD 10.7.1.2 (FHWA GEC 12, 7.10.1)", "",
                            "", governs=False))
    cap.checks[-1].ratio_override = 230.0 / edge_side if edge_side > 0 else math.inf
    tons = v["q_comp"] / SHORT_TON_KN
    crsi_e = 381 if tons <= 60 else 533 if tons <= 120 else 686 if tons < 200 else \
        762 if tons <= 280 else 914
    edge_centre = edge_side + dp / 2
    cap.checks.append(Check("Recommendation", "Edge distance to the pile centre", crsi_e,
                            edge_centre, "mm", "CRSI 4.3 (edge splitting)", "",
                            f"for a pile load of {tons:.0f} tons", governs=False))
    cap.checks[-1].ratio_override = crsi_e / edge_centre if edge_centre > 0 else math.inf
    cap.checks.append(Check("Recommendation", "Pile embedment into the cap", 300.0, v["embed"],
                            "mm", "AASHTO LRFD 10.7.1.2 (FHWA GEC 12, 7.10.1)", "",
                            "150 mm where the pile is anchored with bars; CRSI 4.3: 100 mm "
                            "for concrete piles, 150 mm for steel piles", governs=False))
    cap.checks[-1].ratio_override = 300.0 / v["embed"] if v["embed"] else math.inf
    if v.get("head") == "Fixed":
        need = 2 * dp
        cap.checks.append(Check("Recommendation", "Embedment of a fixed pile head", need,
                                v["embed"], "mm", "FHWA GEC 12, 7.3.7.1", "",
                                "two to three pile diameters, or a specially designed connection",
                                governs=False))
        cap.checks[-1].ratio_override = need / v["embed"] if v["embed"] else math.inf
        factored = [c for c in cap.reactions if any(
            c.startswith(cb.name + " (") and cb.factored for cb in cap.combos)]
        uplift = max([max(-min(cap.reactions[c]), 0.0) for c in factored] or [0.0])
        lf = g.get("Lf", LF_FIRM)
        H = max([cap.lateral_per_pile[c] for c in factored] or [0.0])
        cap.notes.append(
            f"Fixed pile heads: each pile head takes M = H Lf / 2 = {H:.1f} x {lf / 1000:.2f} / 2 "
            f"= {H * lf / 2000:.1f} kN-m (Lf from NSCP 306.6). The anchorage of the piles into "
            f"the cap must develop the uplift ({uplift:.0f} kN) and that moment, and at least "
            f"25 % of the pile's tension strength (NSCP 308.4.1).")
    if v.get("head") == "Pinned":
        cap.notes.append("Pinned pile heads (CRSI 3.3, 6.1): the piles take axial load and shear "
                         "only. Uplift must still be anchored into the cap (NSCP 308.4.1).")


# --------------------------------------------------------------------------- #
# Results
# --------------------------------------------------------------------------- #
def status(check: Check) -> str:
    """PASS or FAIL for a check that governs; OK or NOT MET for a recommendation;
    INFO OK or INFO EXCEEDED for a check shown for information only."""
    if check.ratio is None:
        return "N/A"
    if check.governs:
        return "PASS" if check.passed else "FAIL"
    if check.group == "Recommendation":
        return "OK" if check.passed else "NOT MET"
    return "INFO OK" if check.passed else "INFO EXCEEDED"


def _fmt(value, digits=1) -> str:
    if value is None:
        return "-"
    if isinstance(value, float) and not math.isfinite(value):
        return "inf"
    return f"{value:,.{digits}f}"


def result_rows(cap: PileCap) -> list[list[str]]:
    rows = []
    for c in cap.checks:
        rows.append([c.group, c.name, _fmt(c.demand), _fmt(c.capacity), c.unit,
                     _fmt(c.ratio, 3) if c.ratio is not None else "-", status(c), c.clause])
    return rows


def bar_rows(cap: PileCap) -> list[list[str]]:
    return [[b.face, b.direction, f"{b.count} - {b.diameter:g} mm @ {b.spacing:g}",
             _fmt(b.Mu), _fmt(b.d, 0), _fmt(b.As_flexure, 0), _fmt(b.As_min, 0), _fmt(b.As_tie, 0),
             _fmt(b.As_req, 0), _fmt(b.As_prov, 0), b.anchorage] for b in cap.bars]


def reaction_rows(cap: PileCap) -> list[list[str]]:
    rows = []
    for label, R in cap.reactions.items():
        rows.append([label] + [_fmt(r) for r in R])
    return rows


def headline(cap: PileCap) -> list[str]:
    v = cap.values
    lines = [f"Cap {cap.Lx:.0f} x {cap.Ly:.0f} x {cap.h:.0f} mm on {cap.n} piles of "
             f"{cap.dp:g} mm ({v['pile_shape'].lower()}), f'c {cap.fc:g} MPa, fy {cap.fy:g} MPa"]
    if cap.searched and cap.passed:
        lines.append(f"Least thickness that passes: {cap.h:.0f} mm (typed {cap.h_typed:.0f} mm)")
    elif cap.searched and _depth_cannot_help(cap):
        groups = ", ".join(sorted({c.group.lower() for c in cap.failing()}))
        lines.append(f"No thickness passes: what fails ({groups}) does not change with the "
                     f"thickness; shown at {cap.h:.0f} mm, where everything else passes")
    elif cap.searched:
        lines.append(f"No thickness up to {MAX_DEPTH:.0f} mm passes; shown at {cap.h:.0f} mm")
    if cap.deep and cap.geometry.get("deep_method") == "CRSI special checks":
        lines.append("Deep cap: the CRSI special checks at the column face govern the shear "
                     "(CRSI 5.3, ACI R13.4.2.5)")
    elif cap.deep:
        lines.append("Deep cap: the strut-and-tie model governs the shear (NSCP 413.2.6.3, 423)")
    else:
        lines.append("Ordinary cap: sectional shear (NSCP 413.4.2.3)")
    fails = cap.failing()
    lines.append("RESULT: " + ("PASS" if cap.passed else
                               "FAIL - " + "; ".join(sorted({c.group for c in fails}) or
                                                     ["bar development"])))
    return lines


def summary_text(cap: PileCap) -> str:
    out = ["", "=" * 78, "PILE CAP DESIGN (NSCP 2015)", "=" * 78]
    out += headline(cap)
    out += ["", "Bars:"]
    for b in cap.bars:
        out.append(f"  {b.face:<6} along {b.direction}: {b.count} - {b.diameter:g} mm @ "
                   f"{b.spacing:g} mm, As {b.As_prov:,.0f} mm2 (needs {b.As_req:,.0f}), "
                   f"{b.anchorage.lower()}")
    out += ["", f"{'Check':<62} {'Demand':>10} {'Capacity':>10} {'Ratio':>7}  Status"]
    for c in cap.checks:
        name = f"{c.group}: {c.name}"
        out.append(f"{name[:62]:<62} {_fmt(c.demand):>10} {_fmt(c.capacity):>10} "
                   f"{(_fmt(c.ratio, 3) if c.ratio is not None else '-'):>7}  {status(c)}")
    if cap.notes:
        out += ["", "Notes:"] + [f"  - {n}" for n in cap.notes]
    out.append("=" * 78)
    return "\n".join(out)


def terminal_report(cap: PileCap):
    from utilities.terminal import Report

    out = Report("PILE CAP DESIGN (NSCP 2015)")
    out.lines(headline(cap))
    out.heading("Bars")
    out.table(["Face", "Along", "Bars", "Mu (kN-m)", "d (mm)", "As flexure", "As min", "As tie",
               "As required", "As provided", "Ends"], bar_rows(cap))
    out.heading("Checks")
    out.table(["Group", "Check", "Demand", "Capacity", "Unit", "Ratio", "Status", "Clause"],
              result_rows(cap))
    out.heading("Pile reactions (kN, compression positive)")
    out.table(["Combination"] + [f"Pile {i + 1}" for i in range(cap.n)], reaction_rows(cap))
    out.heading("Equations of the checks")
    out.equations([
        "Pile reactions:  R = P/n + My x / Σx² + Mx y / Σy²  (rigid cap, pinned piles; CRSI 6.2)",
        "Flexure:         φMn = 0.90 As fy (d − a/2) ≥ Mu at the column face  (NSCP 413.2.7.1)",
        "One-way shear:   φVc = 0.75 × 0.17 λ √f′c b d  (NSCP 422.5.5.1)",
        "Two-way shear:   vc = least of 0.33, 0.17(1 + 2/β), 0.083(2 + αs d/b₀) × λ√f′c  "
        "(NSCP 422.6.5.2)",
        "Strut-and-tie:   φFns = 0.75 × 0.85 βs f′c Acs, βs = 0.60;  φFnt = 0.75 As fy  "
        "(NSCP 423)",
    ])
    if cap.notes:
        out.heading("Notes")
        out.lines(cap.notes)
    return out


def run(window: bool | None = None):
    """``sdt pilecap``: the window with the form, the live plan and the results.

    With an answers file (or ``window=False``) the inputs are asked in the four
    dialogs of ``utilities/design_cli.py`` instead, which the file answers."""
    from utilities import answers

    if window is None:
        window = not answers.replaying()
    if window:
        from design.concrete.pile_cap_window import open_window

        return open_window(sys.modules[__name__])
    return run_design(sys.modules[__name__], "Pile Cap (NSCP 2015)", "pile_cap",
                      "Pile_Cap_Report")


def complete(values: dict) -> dict:
    """After the dialogs: choose the loads file when it was not typed."""
    if values.get("load_source") == "File" and not values.get("loads_file"):
        from utilities._gui_helpers import select_open_path

        path = select_open_path("Select the loads file (Excel or CSV)", ".xlsx *.csv",
                                "Loads (Excel or CSV)")
        if not path:
            raise ValueError("No loads file was chosen.")
        values["loads_file"] = path
    return values


# --------------------------------------------------------------------------- #
# The window (utilities/design_window.py)
# --------------------------------------------------------------------------- #
PREVIEW_TITLE = "PLAN"
PREVIEW_KEYS = ["column_shape", "cx", "cy", "layout", "n_piles", "spacing", "edge", "coords",
                "cap_lx", "cap_ly", "dp", "pile_shape"]
MULTILINE = {"loads"}
FILE_INPUTS = {"loads_file": ("Loads (Excel or CSV)", "*.xlsx *.csv", ("load_source", "File"))}
RESULT_ROWS = 14


def preview(figure, values: dict, cap: "PileCap | None") -> str:
    """The plan for the window: the cap, the column, the piles numbered; after a
    run, each pile's largest factored reaction and uplift, its colour by service
    load over capacity, and the critical sections."""
    from matplotlib.patches import Rectangle

    from utilities.design_window import PLOT, style_axes

    piles, Lx, Ly = prepare(values)
    ax = figure.add_axes([0.08, 0.07, 0.9, 0.88])
    style_axes(ax)
    ax.add_patch(Rectangle((-Lx / 2, -Ly / 2), Lx, Ly, facecolor=PLOT["field"],
                           edgecolor=PLOT["text"], lw=1.5, zorder=1))
    x, y = column_shape(values).exterior.xy
    ax.fill(x, y, facecolor=PLOT["amber"], edgecolor=PLOT["on_amber"], lw=1.0, zorder=4)
    reactions = low = ratio = None
    if cap is not None:
        factored = [label for label in cap.reactions
                    if any(label.startswith(c.name + " (") and c.factored for c in cap.combos)]
        source = factored or list(cap.reactions)
        reactions = [max(cap.reactions[label][i] for label in source) for i in range(cap.n)]
        low = [min(cap.reactions[label][i] for label in source) for i in range(cap.n)]
        service = [label for label in cap.reactions if label not in factored]
        if service and values["q_comp"] > 0:
            ratio = [max(cap.reactions[label][i] for label in service) / values["q_comp"]
                     for i in range(cap.n)]
        region = cap.geometry.get("column_perimeter")
        if region is not None and not region.is_empty and region.geom_type == "Polygon":
            rx, ry = region.exterior.xy
            ax.plot(rx, ry, color=PLOT["red"], lw=0.9, ls="--", zorder=3)
        for axis in ("X", "Y"):
            half = (cap.geometry["cx_eq"] if axis == "X" else cap.geometry["cy_eq"]) / 2
            s = half + cap.geometry["d"][axis]
            for side in (1, -1):
                if axis == "X" and s < Lx / 2:
                    ax.plot([side * s] * 2, [-Ly / 2, Ly / 2], color=PLOT["green"], lw=0.8,
                            ls="-.", zorder=2)
                if axis == "Y" and s < Ly / 2:
                    ax.plot([-Lx / 2, Lx / 2], [side * s] * 2, color=PLOT["green"], lw=0.8,
                            ls="-.", zorder=2)
    for i, (px, py) in enumerate(piles):
        sx, sy = pile_shape(values, px, py).exterior.xy
        colour = PLOT["dark"]
        if ratio is not None:
            colour = "#2E5E36" if ratio[i] <= 0.9 else "#6B5418" if ratio[i] <= 1.0 else "#7A2620"
        ax.fill(sx, sy, facecolor=colour, edgecolor=PLOT["text"], lw=1.0, zorder=5)
        ax.annotate(str(i + 1), (px, py), ha="center", va="center", fontsize=9,
                    color=PLOT["text"], fontweight="bold", zorder=6)
        if reactions is not None:
            text = f"{reactions[i]:.0f}" + (f" / {low[i]:.0f}" if low[i] < 0 else "")
            ax.annotate(text, (px, py - values["dp"] / 2), xytext=(0, -9),
                        textcoords="offset points", ha="center", va="top", fontsize=8,
                        color=PLOT["red"] if low[i] < 0 else PLOT["text"], zorder=6)
    pad = 0.12 * max(Lx, Ly)
    ax.annotate(f"{Lx:,.0f}", (0, Ly / 2 + pad * 0.35), ha="center", color=PLOT["muted"],
                fontsize=8)
    ax.annotate(f"{Ly:,.0f}", (Lx / 2 + pad * 0.35, 0), ha="left", va="center", rotation=90,
                color=PLOT["muted"], fontsize=8)
    ax.set_aspect("equal")
    ax.set_xlim(-Lx / 2 - pad, Lx / 2 + pad)
    ax.set_ylim(-Ly / 2 - pad, Ly / 2 + pad)
    ax.grid(color=PLOT["line"], lw=0.4)
    spacing = min((math.dist(a, b) for i, a in enumerate(piles) for b in piles[i + 1:]),
                  default=0)
    lines = [f"Cap {Lx:,.0f} x {Ly:,.0f} mm, {len(piles)} piles of {values['dp']:g} mm, "
             f"least spacing {spacing:,.0f} mm ({spacing / values['dp']:.1f} dp)"]
    if cap is not None:
        lines.append(f"Thickness {cap.h:.0f} mm. Under each pile: its largest factored "
                     "reaction (kN), and its uplift in red. Pile colour: service load over the "
                     "compression capacity (green up to 0.9, amber to 1.0, red above). Dashed "
                     "red: two-way section at d/2; dash-dot green: one-way sections at d.")
    return "\n".join(lines)


def result_parts(cap: PileCap) -> list[tuple[str, str | None]]:
    """The results panel of the window."""
    parts = []
    for line in headline(cap):
        tag = "pass" if line.endswith("PASS") else "fail" if line.startswith("RESULT: FAIL") \
            else None
        parts.append((line + "\n", tag))
    parts.append(("\nBARS\n", "head"))
    for b in cap.bars:
        parts.append((f"{b.face} along {b.direction}: {b.count} - {b.diameter:g} mm @ "
                      f"{b.spacing:g}  ({b.anchorage.lower()})\n", None))
    parts.append(("\nCHECKS, HIGHEST RATIO FIRST\n", "head"))
    shown = sorted([c for c in cap.checks if c.governs and c.ratio is not None],
                   key=lambda c: -c.ratio if math.isfinite(c.ratio) else -1e9)
    for c in shown[:RESULT_ROWS]:
        parts.append((f"{c.ratio:6.3f}  " if math.isfinite(c.ratio) else "   inf  ",
                      "fail" if not c.passed else "pass"))
        parts.append((f"{c.group}: {c.name}\n", None))
    recs = [c for c in cap.checks if c.group == "Recommendation" and not c.passed]
    if recs:
        parts.append(("\nRECOMMENDATIONS NOT MET\n", "head"))
        for c in recs:
            parts.append((f"{c.name}: {c.capacity:,.0f} {c.unit} where {c.clause} "
                          f"recommends {c.demand:,.0f}\n", "muted"))
    if cap.notes:
        parts.append(("\nNOTES\n", "head"))
        for note in cap.notes:
            parts.append((note + "\n", "muted"))
    return parts


def window_summary(cap: PileCap) -> tuple[str, list[str]]:
    return (f"{cap.Lx:.0f} x {cap.Ly:.0f} x {cap.h:.0f} mm, {cap.n} piles: "
            f"{'PASS' if cap.passed else 'FAIL'}",
            [f"{c.group}: {c.name}" for c in cap.failing()])


from design.concrete.pile_cap_report import export_pdf  # noqa: E402,F401

if __name__ == "__main__":
    run()
