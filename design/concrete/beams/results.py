"""The beam results: order, column names and the results workbook."""

import re
import pandas as pd
from design.concrete.beams.deflection import DEFLECTION_COLUMNS
from utilities._calc_report import is_blank, number


_BEAM_MARK_PATTERN = re.compile(r"^(.*?)(BX|BY|GX|GY)-(\d+)(.*)$", re.IGNORECASE)


_BEAM_TYPE_PRIORITY = {"BX": 1, "BY": 2, "GX": 3, "GY": 4}


def beam_sort_key(name: object, face: object = "") -> tuple:
    """Order a beam mark by floor, type, numeric mark, suffix, and face."""
    name_text = str(name).upper()
    face_priority = 0 if "TOP" in str(face).upper() else 1
    match = _BEAM_MARK_PATTERN.match(name_text)
    if not match:
        return name_text, 5, 0, "", face_priority

    floor, beam_type, number, suffix = match.groups()
    return (
        floor,
        _BEAM_TYPE_PRIORITY[beam_type.upper()],
        int(number),
        suffix,
        face_priority,
    )


def sort_beam_rows(df: pd.DataFrame) -> pd.DataFrame:
    """Return rows sorted by beam mark and, when present, TOP before BOTTOM."""
    if "UniqueName" not in df.columns:
        return df.reset_index(drop=True)

    result = df.copy()
    if "Face" in result.columns:
        result["_sort_key"] = [
            beam_sort_key(name, face)
            for name, face in zip(result["UniqueName"], result["Face"])
        ]
    else:
        result["_sort_key"] = result["UniqueName"].map(beam_sort_key)

    return (
        result.sort_values("_sort_key", kind="stable")
        .drop(columns="_sort_key")
        .reset_index(drop=True)
    )


_BEAM_RESULT_LABELS = {
    **{
        f"n_{location}_L{layer}": f"n₍{location}, {layer}₎"
        for location in ("left", "mid", "right")
        for layer in (1, 2, 3)
    },
    "Mu_left": "Mᵤ, left",
    "Mu_mid": "Mᵤ, mid",
    "Mu_right": "Mᵤ, right",
    "Vu_left": "Vᵤ, left",
    "Vu_mid_2h": "Vᵤ, mid (2h)",
    "Vu_right": "Vᵤ, right",
    "Tu_left": "Tᵤ, left",
    "Tu_mid_2h": "Tᵤ, mid (2h)",
    "Tu_right": "Tᵤ, right",
    "ClearSpan_Ln": "Lₙ",
    "Spacing_2H": "s₂H (mm)",
    "Spacing_Mid": "sₘᵢd (mm)",
    "V_sway_max_kN": "Vₛway,max (kN)",
    "Ve_left_kN": "Vₑ, left (kN)",
    "Ve_right_kN": "Vₑ, right (kN)",
    "SMRF_Flexure_Ratio_Check": "SMRF moment strength ratio check",
    "SMRF_Rho_Check": "SMRF steel ratio check",
    "n_side_per_face_gov": "Side bars / face (governing)",
    "Stirrup_Legs": "Stirrup legs",
    "Vc_zero_left": "Concrete shear suppressed (left)",
    "Vc_zero_right": "Concrete shear suppressed (right)",
    "Anchorage_Check": "Stirrup anchorage check",
    "Alternating_Tie_Check": "Alternating tie check",
    "Torsion_Basis": "Torsion designed for",
    "Seismic_Rules": "Seismic rules",
    "Target_DCR_Flexure": "Target ratio, flexure",
    "Target_DCR_Shear": "Target ratio, shear and torsion",
    **DEFLECTION_COLUMNS,
    "Target_DCR_Deflection": "Target ratio, deflection",
    "Carried_Beam_Depth": "Deepest beam it carries (mm)",
    "Carrier_Depth_Check": "Depth against the beams it carries",
    "Design_Status": "Design status",
}


def display_beam_result_labels(df: pd.DataFrame) -> pd.DataFrame:
    """Return a copy with readable Excel headers while retaining internal keys."""
    mapping = {
        source: display
        for source, display in _BEAM_RESULT_LABELS.items()
        if source in df.columns
    }
    return df.rename(columns=mapping)


def restore_beam_result_labels(df: pd.DataFrame) -> pd.DataFrame:
    """Restore internal column keys from display labels read back from Excel."""
    inverse = {display: source for source, display in _BEAM_RESULT_LABELS.items()}
    mapping = {display: source for display, source in inverse.items() if display in df}
    return df.rename(columns=mapping)


def _beam_sort_key(name) -> tuple:
    """Floor, type (BX, BY, GX, GY), number and suffix of a beam mark."""
    name_str = str(name).upper()
    priority = next((i for i, t in enumerate(("BX", "BY", "GX", "GY"), start=1)
                     if t in name_str), 5)
    match = re.search(r"^(.*?)(BX|BY|GX|GY)-(\d+)(.*)$", name_str)
    if match:
        return (match.group(1), priority, int(match.group(3)), match.group(4))
    return (name_str, priority, 0, "")


_BEAM_GROUP_FILLS = ("BDD7EE", "E2EFDA", "FFF2CC", "FCE4D6", "DEEBF7")


_ETABS_COLUMN_ONLY_FIELDS = (
    "Diameter", "ReinfConfig", "IsSpiral", "IsDesigned", "Cover", "NumBars3Dir", "NumBars2Dir",
    "NumBarsCirc", "BarSizeLong", "BarSizeCorn", "BarSizeConf", "SpacingConf", "NumCBars3",
    "NumCBars2",
)


def beam_blank_reason(name: str, row: dict) -> str:
    """Why a cell of the beam results has no value."""
    if name in ("Defl_roof_mm", "Defl_roof_limit_mm"):
        return "N/A - not a roof beam"
    if name.startswith("Defl_"):
        check = str(row.get("Deflection_Check", "")).strip()
        return f"N/A - {check}" if check and check not in ("PASS", "FAIL") else "N/A"
    if name in ("Ve_left_kN", "Ve_right_kN", "V_sway_max_kN") or name.startswith("SMRF_"):
        return "N/A - no seismic design for this beam"
    if name == "Carried_Beam_Depth":
        return "N/A - carries no beam"
    return "N/A"


def write_beam_results_xlsx(results: pd.DataFrame, path: str) -> str:
    """The beam design results as a formatted Excel file (no Excel needed)."""
    from openpyxl import Workbook
    from openpyxl.styles import Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    from utilities._xlsx_values import cell_value, is_blank

    # ETABS column-section fields that come with FRAME DATA: never filled for a beam.
    unused = [c for c in _ETABS_COLUMN_ONLY_FIELDS
              if c in results.columns and results[c].map(is_blank).all()]
    results = results.drop(columns=unused)
    table = display_beam_result_labels(results)
    cols = table.columns.tolist()
    starts = [0]
    for label in ("n₍left, 1₎", "Side bars / face (governing)", "Stirrup legs",
                  "Stirrup anchorage check"):
        starts.append(cols.index(label) if label in cols else len(cols))
    book = Workbook()
    sheet = book.active
    sheet.title = "BEAM DESIGN"
    sheet.append(cols)
    internal = results.columns.tolist()
    for record, row in zip(results.to_dict("records"), table.astype(object).values.tolist()):
        sheet.append([
            cell_value(value, 3 if name in ("Defl_lambda", "Defl_ratio") else 2,
                       beam_blank_reason(name, record) if is_blank(value) else "")
            for name, value in zip(internal, row)
        ])
    thin, thick = Side(style="thin"), Side(style="medium")
    for index, name in enumerate(cols, start=1):
        group = max(g for g, start in enumerate(starts) if index - 1 >= start)
        cell = sheet.cell(row=1, column=index)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=_BEAM_GROUP_FILLS[min(group, 4)])
        cell.border = Border(top=thick, bottom=thick)
        width = max(len(str(name)), *(len(str(v)) for v in table.iloc[:, index - 1]
                                      .head(200).tolist())) if len(table) else len(str(name))
        sheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 8), 45)
    grey = PatternFill("solid", fgColor="F2F2F2")
    for first in range(2, len(table) + 2, 2):  # two rows per beam: TOP and BOTTOM
        for row in (first, first + 1):
            if row > len(table) + 1:
                continue
            for column in range(1, len(cols) + 1):
                cell = sheet.cell(row=row, column=column)
                if (first // 2) % 2 == 1:
                    cell.fill = grey
                if row == first + 1:
                    cell.border = Border(bottom=thin)
    if "UniqueName" in internal:
        from utilities._xlsx_values import mark_failures

        mark_failures(sheet, 2, internal.index("UniqueName") + 1)
    sheet.freeze_panes = "C2"
    book.save(path)
    return path
